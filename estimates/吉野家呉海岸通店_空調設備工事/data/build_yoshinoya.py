#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
吉野家 呉海岸通店 新設工事　空調設備工事 御見積書（宮地機工様式）生成
=====================================================================
様式: 三矢寮 見積（../../三矢寮_空調設備工事/data/build_miyaji.py）と同一構成・同一単価体系
  表紙(工事) ／ 内訳書(空調設備工事) ／ 法定福利費内訳明細書 ／ 御見積条件 ／ 数量拾い根拠(内部) ／ 単価マスタ(内部)
  A4縦、ＭＳ Ｐゴシック、公開列のみ印刷。内部列（原価・粗利）は印刷範囲外。
工事範囲（ご指示）: 機器吊込み・冷媒配管（ガス管保温20mm）・VPドレン・ドレン配管GW保温・室外機設置
金額は全て数式。原価（材料原価＋人工×原価日額）× 提出係数 → 提出単価（10円単位）。
使い方: python3 build_yoshinoya.py <出力ディレクトリ>
"""
import datetime
import json
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from helpers import (fnt, put, merge, set_border, grid, hline, fill_range, set_widths,  # noqa: E402
                     est_lines, text_width_pt, row_h, page_setup,
                     FILL_Y, FILL_G, FILL_L, C_GRAYTXT, FMT_AMT, FMT_QTY, PT_PER_UNIT)
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import PatternFill  # noqa: E402
from openpyxl.utils.indexed_list import IndexedList  # noqa: E402
from openpyxl.workbook.defined_name import DefinedName  # noqa: E402
from openpyxl.worksheet.pagebreak import Break  # noqa: E402
import pricing as PV  # noqa: E402
import items as IT  # noqa: E402

OUT_DIR = sys.argv[1] if len(sys.argv) > 1 else HERE

# =====================================================================
# 案件情報
# =====================================================================
CLIENT = '株式会社共栄店舗'          # 元請（工程表の作成者）。要確認
PROJECT_1 = '吉野家　呉海岸通店　新設工事'
PROJECT_2 = '空調設備工事'
PROJECT_FULL = '吉野家　呉海岸通店　新設工事　空調設備工事'
SCOPE = '機器吊込み・冷媒配管（ガス管保温20mm）・VPドレン・ドレン配管GW保温・室外機設置'
SITE = '広島県呉市海岸（国道31号沿い）'
TERM = '全体工期 2026年9月1日～2026年11月16日（空調設備工事は工程表による・天井配管 10月上旬予定）'
PAYMENT = '別途お打ち合わせ'
VALIDITY = '発行後3ヶ月'
EST_DATE = datetime.date(2026, 10, 7)
DATE_FMT = 'yyyy"年"m"月"d"日"'
COMPANY = '宮地機工株式会社'
REP = '代表取締役　濱本　義樹'
ZIP = '〒722-0051'
ADDR = '広島県尾道市東尾道9-9'
TEL = 'TEL.0848-20-2121'
FAX = 'FAX.0848-20-2126'
STAFF = '濱本　義武'
FILE = '御見積書_吉野家呉海岸通店新設工事_空調設備工事_20261007.xlsx'

SH_COVER = '表紙(工事)'
SH_DETAIL = '内訳書(空調設備工事)'
SH_WF = '法定福利費内訳明細書'
SH_COND = '御見積条件'

G = dict(coef_default=PV.COEF, overhead_rate=PV.OVERHEAD, rounding_unit=PV.ROUND_UNIT,
         labor_day_cost=PV.LABOR_DAY_COST, labor_day_list=PV.LABOR_DAY_LIST)
WELFARE_RATES = [
    ('健康保険料', 0.04890, '健康保険料率 9.780%（協会けんぽ 広島県 令和8年度）×1/2（労使折半）'),
    ('介護保険料', 0.00810, '介護保険料率 1.620%（全国一律 令和8年度）×1/2（労使折半）'),
    ('子ども・子育て支援金', 0.00115, '支援金率 0.230%（令和8年4月分〜）×1/2（労使折半）'),
    ('厚生年金保険料', 0.09150, '厚生年金保険料率 18.300%（平成29年9月以降固定）×1/2（労使折半）'),
    ('雇用保険料', 0.01050, '雇用保険料率（建設の事業）事業主負担分 1.050%'),
    ('子ども・子育て拠出金', 0.00360, '拠出金率 0.360%（事業主全額負担）'),
]
CIRC = '①②③④⑤⑥⑦⑧⑨'
FILL_RED = PatternFill('solid', fgColor='FDE9E9')
FMT_SIGNED = '#,##0_ ;[Red]-#,##0_ '

PUBLIC_REMARK = {'実測': '図面実測数量', '推測': '想定数量（別紙条件書）', '実測+推測': '図面実測＋想定（別紙）'}


def q(sheet):
    return "'" + sheet + "'"


NOTES = {
    'cover': [
        '本見積は換気空調機器表（F-01）・空調設備図（F-02-1）・平面図（A-08）に基づき当社にて数量を拾い出して積算しております。実施数量との差異は別途精算とさせていただきます。',
        '本見積の工事範囲は「' + SCOPE + '」です。空調機（AC-1・AC-2・RAC の室内機・室外機・リモコン）および吹出しパンカーは支給品のため含みません（吊込み・据付費のみ計上）。',
        '渡り操作配線・リモコン取付、試運転調整、AC-2 吹出しダクト・吹出しパンカー取付、換気設備、電源・接地等の電気工事、建築工事は含みません。',
        '仮設（足場・高所作業車等）、揚重機（クレーン等）、産業廃棄物（支給品の梱包材を含む）の処分は元請様にてご手配・ご負担願います。',
        '想定数量の項目は別紙「御見積条件」記載の前提にて計上しており、実数量確定後に精算させていただきます。',
        '本見積書に記載なき事項については別途とさせて頂きます。工事に係る電気・水道は無償支給願います。',
        '本見積金額には消費税は含まれておりません。有効期限内でも銅管・鋼材等の市況が著しく変動した場合は再見積をさせていただくことがあります。',
    ],
    'conditions': [
        '本見積は換気空調機器表 F-01（2025.05.25）・空調設備図 F-02-1（2026.05.25）・平面図 A-08（2026.07.29）に基づき当社にて拾い出した数量により積算しております。冷媒管・ドレン管の平面延長は空調設備図（1/50）から計測し、立上り・立下り・機器接続部は想定数量として計上しています（別紙【2】参照）。',
        '本見積の工事範囲は、ご指示の「' + SCOPE + '」です。空調設備図の特記で空調設備工事とされている項目のうち、渡り操作配線・リモコン取付、リモコン・スイッチ類の名称表示、試運転調整は本見積に含みません。必要な場合は別途お見積りいたします。',
        '空調機（AC-1 天井カセット形 SSRC140C×2台、AC-2 天井ビルトイン形 SSRB140C×1台、RAC 店長室用×1台、各リモコン）および吹出しパンカー PK-SB#10TK×4台は支給品のため含みません。本見積は空調機の搬入・吊込み・据付費のみを計上しています。',
        'AC-2（天井ビルトイン形）から吹出しパンカーまでのダクト・パンカー取付、換気設備（換気扇・送風機・フード・ダクト類）は本見積に含みません。',
        '空調設備図の特記により、室外機の基礎・鉄骨架台（2段積み・溶融亜鉛メッキ防錆処理）・防振ゴムは「室外機設置」に計上しています。架台は既製品相当とし、ダイキン製室外機（H1,430mm）に対応する段間寸法・背面と建物壁面の隙間200mm以上を確保します。外構の土間コンクリートに直接アンカー固定できる場合は基礎を減額いたします。',
        '冷媒管は被覆銅管とし、ガス管は保温厚20mm（耐熱ポリエチレンフォーム）、液管はメーカー標準被覆としています。屋外露出部は特記により断熱材巻きの上、屋外用化粧カバー（スリムダクト）仕上げとして計上しています。',
        '冷媒の追加充填は、各系統の配管長（最長約14m）が支給機のチャージレス配管長以内と想定し計上していません。必要となった場合は実費にて精算させていただきます。',
        'ドレン管は硬質塩化ビニル管（VP）とし、空調設備図の口径（25A・30A・40A）により計上しています。範囲は各機器から「空調機ドレン接続立上り（H=150）」までとし、以降は給排水工事です。図示のドレンホース用逆止弁 NDB-20-25 は本見積に計上しています。',
        'ドレン管の保温はグラスウール保温筒 厚20mm・アルミガラスクロス（ALGC）仕上げとして全長に計上しています。厨房内の露出部をステンレスラッキング等とする場合は別途協議とさせていただきます。',
        '冷媒管の外壁貫通穴明け（4箇所）は本見積に計上しています。構造材の補強、天井・壁の開口および仕上げ補修、天井点検口は建築工事とし含みません。',
        '仮設（足場・高所作業車・仮設電気・水道等）および産業廃棄物（支給品の梱包材を含む）の処分費は元請様にてご負担願います（当社発生材は指定集積場所への分別搬出まで）。',
        '工期は工程表（吉野家呉海岸通店新設工事 工程表）に基づく通常の日中作業を前提としております。夜間・休日作業、工程変更や他工事との輻輳による手待ち、工期延長に伴う経費増は別途協議とさせていただきます。',
        '本見積書に記載なき事項については別途とさせていただきます。工事に係る電気・水道は無償支給願います。',
        '本見積金額には消費税は含まれておりません。労務費に係る法定福利費（事業主負担分）は別紙「法定福利費内訳明細書」のとおり計上しています。',
    ],
    'checks': [
        '宛名は元請様を株式会社共栄店舗（工程表の作成者）と想定して記載しています。ご指定があれば修正します。',
        '工事名は工程表の「吉野家呉海岸通店新設工事」によりました（図面標題は「吉野家 呉海岸店 新設工事」）。施工場所は位置図（呉海岸）によります。地番をご教示ください。',
        '冷媒管・ドレン管は空調設備図 F-02-1 の配管線（R・D 記号）をベクトルデータから計測しています。立上り・立下りは室外機設置概要図（2段架台）と天井高（客席 CH2,600、厨房・店長室 CH2,500）から想定しています。',
        '室外機は室外機設置概要図により、架台①＝AC-1（下段）＋AC-1（上段）、架台②＝AC-2（下段）＋RAC（上段）の組合せとしています。',
        '特記の「天井内冷媒配管は厚20耐熱ポリスチレンフォーム断熱材巻き」は、冷媒管用の耐熱ポリエチレンフォーム保温材（厚20mm）と解釈しています。',
        '施工地域が「北日本吉野家・関東営業本部」以外のため、支給機はダイキン SSRC140C・SSRB140C（モリタニ・ダイキン手配）を前提としています。',
        'ドレン接続立上り（H=150）の立上り管および以降の排水管は給排水工事と解釈しています。',
    ],
}


# =====================================================================
# 内訳書（セクション別小計＋合計）
# =====================================================================
def to_lines():
    out = []
    for sec in IT.SECTIONS:
        for ln in sec['lines']:
            mat, md, ref, src = PV.MASTER[ln['key']]
            pub = ln['pub'] or PUBLIC_REMARK.get(ln['basis'], '')
            out.append(dict(ln, section=sec['key'], mat=mat, md=md, ref=ref, src=src, pub=pub))
    return out


def build_detail(ws):
    lines = to_lines()
    set_widths(ws, dict(A=4, B=27, C=24, D=4, E=8, F=5, G=10, H=13, I=15,
                        J=10, K=7, L=10, M=10, N=7, O=12, P=12, Q=9, R=12, S=8, T=11, U=8, V=40))
    merge(ws, 'A1:I1')
    put(ws, 'A1', '内　訳　書', size=16, bold=True, h='center')
    merge(ws, 'J1:V1')
    put(ws, 'J1', '原　価　内　訳　書（内部用・印刷範囲外）', size=11, bold=True, h='center', color=C_GRAYTXT, fill=FILL_G)
    fill_range(ws, 'J1:V1', FILL_G)
    ws.row_dimensions[1].height = 26
    put(ws, 'A2', '工事名：' + PROJECT_FULL)
    put(ws, 'I2', '内訳書 No.1', h='right')
    ws.row_dimensions[2].height = 18
    merge(ws, 'A3:B3')
    merge(ws, 'C3:D3')
    for ref, txt in [('A3', '名　　称'), ('C3', '規格・寸法'), ('E3', '数 量'), ('F3', '単位'),
                     ('G3', '単　価'), ('H3', '金　　額'), ('I3', '備　考')]:
        put(ws, ref, txt, h='center', fill=FILL_G)
    fill_range(ws, 'A3:I3', FILL_G)
    for ref, txt in [('J3', '材料原価'), ('K3', '人工/単位'), ('L3', '労務原価'), ('M3', '基準原価単価'), ('N3', '提出係数'),
                     ('O3', '基準原価'), ('P3', '労務費(提出)'), ('Q3', '数量根拠'), ('R3', '粗利'),
                     ('S3', '粗利率'), ('T3', '参考定価ベース'), ('U3', '当社/参考'), ('V3', '単価出典・メモ（内部）')]:
        put(ws, ref, txt, size=9, h='center', fill=FILL_G, color=C_GRAYTXT, shrink=True)
    ws.row_dimensions[3].height = 20
    merge(ws, 'A4:I4')
    put(ws, 'A4', '空調設備工事（' + SCOPE + '）', bold=True, shrink=True)
    ws.row_dimensions[4].height = 18
    r = 5
    sec_rows = {}
    wrapped = []
    for sec in IT.SECTIONS:
        merge(ws, f'A{r}:I{r}')
        put(ws, f'A{r}', sec['title'], size=10, bold=True, shrink=True)
        ws.row_dimensions[r].height = 18
        r += 1
        first = r
        for ln in [x for x in lines if x['section'] == sec['key']]:
            merge(ws, f'A{r}:B{r}')
            name_txt = '　' + ln['name']
            long_name = text_width_pt(name_txt, 10) > (31 - 0.8) * PT_PER_UNIT / 0.85
            put(ws, f'A{r}', name_txt, size=10, wrap=long_name, shrink=not long_name)
            merge(ws, f'C{r}:D{r}')
            spec_lines = est_lines(ln['spec'], 28 - 1, 9) if ln['spec'] else 1
            if spec_lines > 1:   # 長い規格は 9pt で折返し（縮小すると読めなくなるため）
                put(ws, f'C{r}', ln['spec'], size=9, wrap=True, indent=1)
            else:
                put(ws, f'C{r}', ln['spec'] or None, size=10, shrink=True, indent=1)
            put(ws, f'E{r}', ln['qty'], size=10, h='right', fmt=FMT_SIGNED)
            put(ws, f'F{r}', ln['unit'], size=10, h='center')
            put(ws, f'G{r}', f'=IF(M{r}="","",ROUND(M{r}*N{r},-1))', size=10, h='right', fmt=FMT_AMT)
            put(ws, f'H{r}', f'=IF(G{r}="","",ROUND(E{r}*G{r},0))', size=10, h='right', fmt=FMT_SIGNED)
            put(ws, f'I{r}', ln['pub'] or None, size=10, h='center', shrink=True)
            put(ws, f'J{r}', ln['mat'], size=10, h='right', fmt=FMT_AMT, fill=FILL_Y)
            put(ws, f'K{r}', ln['md'], size=10, h='right', fmt='0.000', fill=FILL_Y)
            put(ws, f'L{r}', f'=ROUND(K{r}*LABOR_DAY,0)', size=10, h='right', fmt=FMT_AMT)
            put(ws, f'M{r}', f'=J{r}+L{r}', size=10, h='right', fmt=FMT_AMT)
            put(ws, f'N{r}', '=COEF', size=10, h='center', fmt='0.00', fill=FILL_Y)
            put(ws, f'O{r}', f'=E{r}*M{r}', size=10, h='right', fmt=FMT_SIGNED)
            put(ws, f'P{r}', f'=ROUND(E{r}*L{r}*N{r},0)', size=10, h='right', fmt=FMT_SIGNED)
            put(ws, f'Q{r}', ln['basis'], size=9, h='center', shrink=True,
                fill=(FILL_RED if ln['basis'] in ('推測', '実測+推測') else None))
            put(ws, f'R{r}', f'=IF(H{r}="","",H{r}-O{r})', size=10, h='right', fmt=FMT_SIGNED)
            put(ws, f'S{r}', f'=IF(OR(H{r}="",H{r}=0),"",R{r}/H{r})', size=10, h='right', fmt='0.0%')
            put(ws, f'T{r}', (round(ln['ref']) if ln['ref'] else None), size=10, h='right', fmt=FMT_AMT)
            put(ws, f'U{r}', (f'=IF(OR(T{r}="",T{r}=0),"",G{r}/T{r})' if ln['ref'] else None), size=10, h='right', fmt='0%')
            put(ws, f'V{r}', ((ln['src'] + '｜' if ln['src'] else '') + (ln['memo'] or '')) or None, size=9)
            ws.row_dimensions[r].height = 18
            if long_name or spec_lines > 1:
                ws.row_dimensions[r].height = max(18, row_h(est_lines(name_txt, 31, 10), 10) if long_name else 0,
                                                  row_h(spec_lines, 9, pad=4) if spec_lines > 1 else 0)
                wrapped.append(r)
            r += 1
        last = r - 1
        # 小計
        merge(ws, f'A{r}:D{r}')
        put(ws, f'A{r}', '小　　計', size=10, bold=True, h='center')
        put(ws, f'H{r}', f'=SUM(H{first}:H{last})', size=10, bold=True, h='right', fmt=FMT_SIGNED)
        put(ws, f'K{r}', f'=SUMPRODUCT(E{first}:E{last},K{first}:K{last})', size=10, bold=True, h='right', fmt='#,##0.0')
        put(ws, f'O{r}', f'=SUM(O{first}:O{last})', size=10, bold=True, h='right', fmt=FMT_SIGNED)
        put(ws, f'P{r}', f'=SUM(P{first}:P{last})', size=10, bold=True, h='right', fmt=FMT_SIGNED)
        put(ws, f'R{r}', f'=H{r}-O{r}', size=10, bold=True, h='right', fmt=FMT_SIGNED)
        put(ws, f'S{r}', f'=IF(H{r}=0,"",R{r}/H{r})', size=10, bold=True, h='right', fmt='0.0%')
        ws.row_dimensions[r].height = 18
        hline(ws, f'A{r}:I{r}', 'top', 'thin')
        hline(ws, f'J{r}:V{r}', 'top', 'thin')
        sec_rows[sec['key']] = dict(first=first, last=last, sub=r)
        r += 1
    blank = r
    merge(ws, f'A{blank}:B{blank}')
    merge(ws, f'C{blank}:D{blank}')
    ws.row_dimensions[blank].height = 18
    tot = blank + 1
    subs = [v['sub'] for v in sec_rows.values()]
    merge(ws, f'A{tot}:D{tot}')
    put(ws, f'A{tot}', '【合　　計】', bold=True, h='center')
    put(ws, f'H{tot}', '=' + '+'.join(f'H{s}' for s in subs), bold=True, h='right', fmt=FMT_SIGNED)
    put(ws, f'K{tot}', '=' + '+'.join(f'K{s}' for s in subs), bold=True, h='right', fmt='#,##0.0')
    put(ws, f'O{tot}', '=' + '+'.join(f'O{s}' for s in subs), bold=True, h='right', fmt=FMT_SIGNED)
    put(ws, f'P{tot}', '=' + '+'.join(f'P{s}' for s in subs), bold=True, h='right', fmt=FMT_SIGNED)
    put(ws, f'R{tot}', f'=H{tot}-O{tot}', bold=True, h='right', fmt=FMT_SIGNED)
    put(ws, f'S{tot}', f'=IF(H{tot}=0,"",R{tot}/H{tot})', bold=True, h='right', fmt='0.0%')
    ws.row_dimensions[tot].height = 22
    grid(ws, f'A3:I{tot}', outer='medium', vert='thin', horiz='hair')
    grid(ws, f'J3:V{tot}', outer='thin', vert='thin', horiz='hair')
    for rng in ('A3:I3', 'J3:V3'):
        hline(ws, rng, 'bottom', 'thin')
    for v in sec_rows.values():   # 小計行の上に細線（grid で消えるため再設定）
        hline(ws, f'A{v["sub"]}:I{v["sub"]}', 'top', 'thin')
        hline(ws, f'J{v["sub"]}:V{v["sub"]}', 'top', 'thin')
    hline(ws, f'A{tot}:I{tot}', 'top', 'medium')
    hline(ws, f'J{tot}:V{tot}', 'top', 'medium')
    for c in range(1, 10):
        set_border(ws.cell(tot, c), bottom='medium')
    set_border(ws.cell(tot, 1), left='medium')
    set_border(ws.cell(tot, 9), right='medium')
    page_setup(ws, f'A1:I{tot}', fit_h=0, title_rows='1:3', footer='&P / &N')
    return dict(sec_rows=sec_rows, total=tot, lines=lines, wrapped=wrapped)


# =====================================================================
# 法定福利費内訳明細書
# =====================================================================
def build_welfare(ws, det):
    set_widths(ws, dict(A=2, B=25, C=9, D=9, E=14, F=16, G=31))
    merge(ws, 'B1:G1')
    put(ws, 'B1', '法 定 福 利 費 内 訳 明 細 書', size=14, bold=True, h='center')
    ws.row_dimensions[1].height = 26
    merge(ws, 'B2:G2')
    put(ws, 'B2', f'="（御見積書 No."&{q(SH_COVER)}!K3&" の添付資料）"', size=10, h='center')
    ws.row_dimensions[2].height = 16
    put(ws, 'B4', '見積年月日', fill=FILL_L, indent=1)
    merge(ws, 'C4:E4')
    put(ws, 'C4', f"={q(SH_COVER)}!J4", h='left', fmt=DATE_FMT, indent=1)
    put(ws, 'F4', '見積No.', fill=FILL_L, indent=1)
    put(ws, 'G4', f'=IF({q(SH_COVER)}!K3="","",{q(SH_COVER)}!K3)', h='left', indent=1)
    for r, lab, val in [(5, '発注者名', CLIENT + '　御中'), (6, '工事名', PROJECT_FULL), (7, '施工場所', SITE)]:
        put(ws, f'B{r}', lab, fill=FILL_L, indent=1)
        merge(ws, f'C{r}:G{r}')
        put(ws, f'C{r}', val, shrink=True, indent=1)
    put(ws, 'B8', '会社名', fill=FILL_L, indent=1)
    merge(ws, 'C8:E8')
    put(ws, 'C8', COMPANY, indent=1)
    put(ws, 'F8', '所在地', fill=FILL_L, indent=1)
    put(ws, 'G8', ADDR, indent=1)
    grid(ws, 'B4:G8', outer='thin', vert='thin', horiz='thin')
    for r in range(4, 9):
        ws.row_dimensions[r].height = 20
    ws.row_dimensions[3].height = 8
    ws.row_dimensions[9].height = 10
    put(ws, 'B10', '【1】　労務費の内訳（法定福利費の算定基礎）', bold=True)
    ws.row_dimensions[10].height = 20
    note1 = ('※ 本見積の労務費は、内訳書各項目の材工共単価のうち労務費相当額（自社施工分）の合計です。'
             '空調機は支給品のため含みません。')
    merge(ws, 'B11:G11')
    put(ws, 'B11', note1, size=10, wrap=True)
    ws.row_dimensions[11].height = row_h(est_lines(note1, 104, 10), 10)
    put(ws, 'B12', '項　　目', h='center', fill=FILL_G)
    merge(ws, 'C12:D12')
    put(ws, 'C12', '労務費（円）', h='center', fill=FILL_G)
    merge(ws, 'E12:G12')
    put(ws, 'E12', '摘　　要', h='center', fill=FILL_G)
    fill_range(ws, 'B12:G12', FILL_G)
    rows1 = []
    for i, sec in enumerate(IT.SECTIONS):
        sub = det['sec_rows'][sec['key']]['sub']
        rows1.append((13 + i, CIRC[i] + sec['cover'].split('（')[0], f"={q(SH_DETAIL)}!P{sub}",
                      f"内訳書 {sec['title'][:3]} 各項目の労務費相当額の合計"))
    tot_row = 13 + len(IT.SECTIONS)
    rows1.append((tot_row, '労 務 費 合 計', f'=SUM(C13:C{tot_row - 1})', '← 法定福利費の算定基礎額'))
    for r, lab, f, memo in rows1:
        bold = (r == tot_row)
        put(ws, f'B{r}', lab, bold=bold, h=('center' if bold else None), indent=(0 if bold else 1), shrink=True)
        merge(ws, f'C{r}:D{r}')
        put(ws, f'C{r}', f, bold=bold, h='right', fmt=FMT_AMT)
        merge(ws, f'E{r}:G{r}')
        put(ws, f'E{r}', memo, size=10, bold=bold, indent=1)
    for r in range(12, tot_row + 1):
        ws.row_dimensions[r].height = 20
    grid(ws, f'B12:G{tot_row}', outer='medium', vert='thin', horiz='thin')
    hline(ws, f'B{tot_row}:G{tot_row}', 'top', 'double')
    for c in range(2, 8):
        set_border(ws.cell(tot_row, c), bottom='medium')
    r = tot_row + 1
    ws.row_dimensions[r].height = 10
    r += 1
    put(ws, f'B{r}', '【2】　法定福利費の算定', bold=True)
    ws.row_dimensions[r].height = 20
    r += 1
    merge(ws, f'B{r}:G{r}')
    put(ws, f'B{r}', '計算式：　法定福利費 ＝ 労務費 × 事業主負担料率　（保険ごとに円未満切捨て）', size=10)
    ws.row_dimensions[r].height = 18
    r += 1
    hdr2 = r
    put(ws, f'B{r}', '保険の種類', h='center', fill=FILL_G)
    merge(ws, f'C{r}:D{r}')
    put(ws, f'C{r}', '対象労務費（円）', size=10, h='center', fill=FILL_G)
    put(ws, f'E{r}', '事業主負担料率', size=10, h='center', fill=FILL_G)
    put(ws, f'F{r}', '法定福利費（円）', size=10, h='center', fill=FILL_G)
    put(ws, f'G{r}', '算定根拠', h='center', fill=FILL_G)
    fill_range(ws, f'B{r}:G{r}', FILL_G)
    ws.row_dimensions[r].height = 22
    r += 1
    labor_cell = f'$C${tot_row}'
    for name, rate, basis in WELFARE_RATES:
        put(ws, f'B{r}', name, indent=1)
        merge(ws, f'C{r}:D{r}')
        put(ws, f'C{r}', f'={labor_cell}', h='right', fmt=FMT_AMT)
        put(ws, f'E{r}', rate, h='right', fmt='0.000%_ ')
        put(ws, f'F{r}', f'=ROUNDDOWN(C{r}*E{r},0)', h='right', fmt=FMT_AMT)
        put(ws, f'G{r}', basis, size=9, wrap=True)
        ws.row_dimensions[r].height = max(22, row_h(est_lines(basis, 31, 9), 9, pad=5))
        r += 1
    wf_first, wf_last = hdr2 + 1, r - 1
    tot = r
    put(ws, f'B{tot}', '合　　計', bold=True, h='center')
    merge(ws, f'C{tot}:D{tot}')
    put(ws, f'E{tot}', f'=SUM(E{wf_first}:E{wf_last})', bold=True, h='right', fmt='0.000%_ ')
    put(ws, f'F{tot}', f'=SUM(F{wf_first}:F{wf_last})', bold=True, h='right', fmt=FMT_AMT)
    ws.row_dimensions[tot].height = 24
    grid(ws, f'B{hdr2}:G{tot}', outer='medium', vert='thin', horiz='thin')
    hline(ws, f'B{tot}:G{tot}', 'top', 'double')
    for c in range(2, 8):
        set_border(ws.cell(tot, c), bottom='medium')
    ws.row_dimensions[tot + 1].height = 10
    notes = ['※ 料率は見積作成時点（令和8年度）の公表料率によります。料率が改定された場合は改定後の料率により精算させていただきます。',
             f'※ 法定福利費は御見積書の{CIRC[len(IT.SECTIONS)]}に計上しています。']
    r = tot + 2
    for t in notes:
        merge(ws, f'B{r}:G{r}')
        put(ws, f'B{r}', t, size=10, wrap=True)
        ws.row_dimensions[r].height = row_h(est_lines(t, 104, 10), 10)
        r += 1
    page_setup(ws, f'B1:G{r - 1}', fit_h=1)
    return dict(rate_cell=f'$E${tot}', amount_cell=f'F{tot}')


# =====================================================================
# 表紙(工事)
# =====================================================================
def build_cover(ws, det, wf):
    secs = IT.SECTIONS
    n = len(secs)
    FIRST = 22
    WF_ROW = FIRST + n
    OH_ROW = WF_ROW + 1
    DISC_ROW = OH_ROW + 2
    TOTAL_ROW = DISC_ROW + 1
    set_widths(ws, dict(A=6, B=8, C=3, D=15, E=9, F=9, G=7, H=5, I=7, J=7, K=14, L=12,
                        M=26, N=12, O=12, P=12, Q=9, R=12, S=12))
    merge(ws, 'A1:L2')
    put(ws, 'A1', '御　見　積　書', size=20, bold=True, h='center')
    ws.row_dimensions[1].height = 22
    ws.row_dimensions[2].height = 22
    put(ws, 'J3', 'No.', h='right')
    merge(ws, 'K3:L3')
    put(ws, 'K3', None, h='center')
    hline(ws, 'K3:L3', 'bottom', 'thin')
    ws.row_dimensions[3].height = 20
    merge(ws, 'J4:L4')
    put(ws, 'J4', EST_DATE, h='right', fmt=DATE_FMT)
    ws.row_dimensions[4].height = 20
    merge(ws, 'A5:F5')
    put(ws, 'A5', CLIENT, size=14, h='center')
    hline(ws, 'A5:F5', 'bottom', 'thin')
    put(ws, 'G5', '御中', size=12, h='left')
    ws.row_dimensions[5].height = 26
    ws.row_dimensions[6].height = 12
    put(ws, 'A7', '金額', size=12, h='center')
    merge(ws, 'B7:E7')
    put(ws, 'B7', f'="¥"&TEXT(K{TOTAL_ROW},"#,##0")&"-"', size=14, bold=True, h='center')
    hline(ws, 'B7:E7', 'bottom', 'double')
    put(ws, 'F7', '（税別）', h='left')
    put(ws, 'H7', COMPANY, size=12, bold=True)
    ws.row_dimensions[7].height = 28
    ws.row_dimensions[8].height = 10
    merge(ws, 'A9:F9')
    put(ws, 'A9', '上記の通り御見積申し上げます。', h='left')
    put(ws, 'H9', REP)
    put(ws, 'A10', '何卒ご用命賜わりますようお願い申し上げます。')
    put(ws, 'H11', ZIP)
    merge(ws, 'A12:B12')
    put(ws, 'A12', '工事名', h='distributed', indent=1)
    merge(ws, 'D12:G12')
    put(ws, 'D12', PROJECT_1, shrink=True)
    put(ws, 'H12', ADDR)
    merge(ws, 'D13:G13')
    put(ws, 'D13', PROJECT_2, shrink=True)
    put(ws, 'H13', TEL)
    put(ws, 'K13', FAX)
    merge(ws, 'A14:B14')
    put(ws, 'A14', '施工場所', h='distributed', indent=1)
    merge(ws, 'D14:G14')
    put(ws, 'D14', SITE, size=10, shrink=True)
    put(ws, 'H14', '担当者')
    put(ws, 'J14', STAFF)
    for r, lab, val in [(15, '工事期限', TERM), (16, '御支払条件', PAYMENT), (17, '有効期限', VALIDITY)]:
        merge(ws, f'A{r}:B{r}')
        put(ws, f'A{r}', lab, h='distributed', indent=1)
        merge(ws, f'D{r}:L{r}')
        put(ws, f'D{r}', val, shrink=True)
    for r in range(9, 18):
        ws.row_dimensions[r].height = 19
    ws.row_dimensions[14].height = 22
    ws.row_dimensions[18].height = 12
    merge(ws, 'B19:F19')
    merge(ws, 'I19:J19')
    for ref, txt, sz in [('A19', '項　目', 10), ('B19', '内　　　容', 11), ('G19', '数 量', 11), ('H19', '単位', 10),
                         ('I19', '単  価', 11), ('K19', '金　　額', 11), ('L19', '備　考', 11)]:
        put(ws, ref, txt, size=sz, h='center', fill=FILL_G, shrink=True)
    fill_range(ws, 'A19:L19', FILL_G)
    merge(ws, 'B20:L20')
    put(ws, 'B20', '【' + PROJECT_FULL + '】', indent=1, shrink=True)
    merge(ws, 'B21:L21')
    put(ws, 'B21', '［工事範囲：' + SCOPE + '　空調機・吹出しパンカーは支給品］', size=10, indent=1, shrink=True)
    for r in range(22, TOTAL_ROW + 1):
        merge(ws, f'B{r}:F{r}')
        merge(ws, f'I{r}:J{r}')
    rows = []
    for i, sec in enumerate(secs):
        sub = det['sec_rows'][sec['key']]['sub']
        rows.append((FIRST + i, CIRC[i], sec['cover'], f"={q(SH_DETAIL)}!H{sub}", '内訳書 ' + sec['title'][:3]))
    rows.append((WF_ROW, CIRC[n], '法定福利費（労務費に係る事業主負担分）', f"={q(SH_WF)}!{wf['amount_cell']}", '別紙明細書'))
    rows.append((OH_ROW, CIRC[n + 1], '諸経費（現場経費・資材運搬・工事諸経費）', f'=ROUND(SUM(K{FIRST}:K{WF_ROW - 1})*OVERHEAD_RATE,-3)',
                 f'="{CIRC[0]}〜{CIRC[n - 1]}計の"&TEXT(OVERHEAD_RATE,"0%")'))
    for r, no, name, fk, remark in rows:
        put(ws, f'A{r}', no, h='center')
        put(ws, f'B{r}', name, indent=1, shrink=True)
        put(ws, f'G{r}', 1, h='right', fmt=FMT_QTY)
        put(ws, f'H{r}', '式', h='center')
        put(ws, f'I{r}', f'=K{r}', h='right', fmt=FMT_SIGNED)
        put(ws, f'K{r}', fk, h='right', fmt=FMT_SIGNED)
        put(ws, f'L{r}', remark, size=10, h='center', shrink=True)
    put(ws, f'B{DISC_ROW}', '出精値引き', indent=1)
    put(ws, f'K{DISC_ROW}', f'=IF(TARGET="",FLOOR(SUM(K{FIRST}:K{OH_ROW}),ROUND_UNIT)-SUM(K{FIRST}:K{OH_ROW}),TARGET-SUM(K{FIRST}:K{OH_ROW}))',
        h='right', fmt=FMT_SIGNED)
    put(ws, f'L{DISC_ROW}', '端数調整', size=10, h='center')
    put(ws, f'B{TOTAL_ROW}', '　合　　計', bold=True, indent=1)
    put(ws, f'K{TOTAL_ROW}', f'=SUM(K{FIRST}:K{DISC_ROW})', bold=True, h='right', fmt=FMT_AMT)
    put(ws, f'L{TOTAL_ROW}', '（税別）', size=10, h='center')
    for r in range(19, TOTAL_ROW + 1):
        ws.row_dimensions[r].height = 24
    ws.row_dimensions[20].height = 22
    ws.row_dimensions[21].height = 20
    grid(ws, f'A19:L{TOTAL_ROW}', outer='medium', vert='thin', horiz='thin')
    hline(ws, f'A{TOTAL_ROW}:L{TOTAL_ROW}', 'top', 'double')
    for c in range(1, 13):
        set_border(ws.cell(TOTAL_ROW, c), bottom='medium')
    ws.row_dimensions[TOTAL_ROW + 1].height = 12
    r = TOTAL_ROW + 2
    for t in NOTES['cover'] + ['詳細条件は別紙「御見積条件」をご参照ください。']:
        put(ws, f'A{r}', '※', size=9.5, h='center', v='top')
        merge(ws, f'B{r}:L{r}')
        put(ws, f'B{r}', t, size=9.5, wrap=True, v='top')
        ws.row_dimensions[r].height = row_h(est_lines(t, 96, 9.5), 9.5, pad=3)
        r += 1
    last_note = r - 1

    # ---- 内部エリア（M〜Q、印刷範囲外）: 単価設定・原価・粗利 ----
    put(ws, 'M3', '【設定・内部用】（印刷範囲外）', bold=True, color=C_GRAYTXT, fill=FILL_G)
    put(ws, 'N3', None, fill=FILL_G)
    params = [
        (4, '提出係数（原価→定価ベース単価）', G['coef_default'], '0.00', True),
        (5, '諸経費率（現場経費・運搬・雑材・工事諸経費）', G['overhead_rate'], '0%', True),
        (6, '出精値引き 丸め単位（目標額空欄時）', G['rounding_unit'], '#,##0', True),
        (7, '法定福利費率（労務費比）', '=WELFARE_RATE', '0.000%', False),
        (8, '労務原価日額（円/人工）', G['labor_day_cost'], '#,##0', True),
        (9, '目標提出額（税別・空欄なら丸め）', None, '#,##0', True),
        (10, '参考：宮地機工 人工単価（定価ベース）', G['labor_day_list'], '#,##0', False),
    ]
    for rr, lab, val, fmt, is_input in params:
        put(ws, f'M{rr}', lab, size=10, color=C_GRAYTXT, fill=FILL_G, shrink=True)
        put(ws, f'N{rr}', val, h='right', fmt=fmt, fill=(FILL_Y if is_input else None))
    grid(ws, 'M3:N10', outer='thin', vert='thin', horiz='thin')
    put(ws, 'M11', '※ 原価＝材料原価＋人工×N8。提出単価＝原価×N4（1.22で参考見積の人工単価23,225に一致）。目標提出額 N9 を入れると出精値引きで合わせる', size=9, color=C_GRAYTXT)
    for ref, txt in [('M19', '内部集計（原価・粗利）'), ('N19', '基準原価'), ('O19', '労務費(提出)'),
                     ('P19', '粗利'), ('Q19', '粗利率')]:
        put(ws, ref, txt, size=10, h='center', color=C_GRAYTXT, fill=FILL_G, shrink=True)
    for i, sec in enumerate(secs):
        rr = FIRST + i
        sub = det['sec_rows'][sec['key']]['sub']
        put(ws, f'M{rr}', CIRC[i] + sec['cover'].split('（')[0], size=10, shrink=True)
        put(ws, f'N{rr}', f"={q(SH_DETAIL)}!O{sub}", h='right', fmt=FMT_SIGNED)
        put(ws, f'O{rr}', f"={q(SH_DETAIL)}!P{sub}", h='right', fmt=FMT_SIGNED)
        put(ws, f'P{rr}', f'=K{rr}-N{rr}', h='right', fmt=FMT_SIGNED)
        put(ws, f'Q{rr}', f'=IF(K{rr}=0,"",P{rr}/K{rr})', h='right', fmt='0.0%')
    put(ws, f'M{WF_ROW}', CIRC[n] + '法定福利費（原価＝提出値）', size=10, shrink=True)
    put(ws, f'N{WF_ROW}', f'=K{WF_ROW}', h='right', fmt=FMT_AMT)
    put(ws, f'M{OH_ROW}', CIRC[n + 1] + '諸経費（原価計上なし）', size=10, shrink=True)
    put(ws, f'N{OH_ROW}', 0, h='right', fmt=FMT_AMT)
    put(ws, f'M{DISC_ROW}', '出精値引き（原価なし・粗利から控除）', size=10, shrink=True)
    put(ws, f'N{DISC_ROW}', 0, h='right', fmt=FMT_AMT)
    put(ws, f'P{DISC_ROW}', f'=K{DISC_ROW}', h='right', fmt=FMT_SIGNED)
    put(ws, f'M{TOTAL_ROW}', '合計（原価合計／粗利／粗利率）', size=10, bold=True, shrink=True)
    put(ws, f'N{TOTAL_ROW}', f'=SUM(N{FIRST}:N{DISC_ROW})', bold=True, h='right', fmt=FMT_AMT)
    put(ws, f'O{TOTAL_ROW}', f'=SUM(O{FIRST}:O{DISC_ROW})', h='right', fmt=FMT_AMT)
    put(ws, f'P{TOTAL_ROW}', f'=K{TOTAL_ROW}-N{TOTAL_ROW}', bold=True, h='right', fmt=FMT_AMT)
    put(ws, f'Q{TOTAL_ROW}', f'=IF(K{TOTAL_ROW}=0,"",P{TOTAL_ROW}/K{TOTAL_ROW})', bold=True, h='right', fmt='0.0%')
    grid(ws, f'M19:Q{TOTAL_ROW}', outer='thin', vert='thin', horiz='hair')
    hline(ws, 'M19:Q19', 'bottom', 'thin')
    hline(ws, f'M{TOTAL_ROW}:Q{TOTAL_ROW}', 'top', 'thin')
    R0 = TOTAL_ROW + 2
    put(ws, f'M{R0}', '粗利率別 参考提出額（税別・万円未満切上げ）', size=10, color=C_GRAYTXT, fill=FILL_G, shrink=True)
    put(ws, f'N{R0}', '参考提出額', size=10, h='center', color=C_GRAYTXT, fill=FILL_G)
    put(ws, f'O{R0}', '必要提出係数', size=10, h='center', color=C_GRAYTXT, fill=FILL_G)
    for i, rate in enumerate([0.10, 0.15, 0.20, 0.25, 0.30]):
        rr = R0 + 1 + i
        put(ws, f'M{rr}', rate, h='center', fmt='"粗利率 "0%')
        put(ws, f'N{rr}', f'=ROUNDUP($N${TOTAL_ROW}/(1-M{rr}),-4)', h='right', fmt=FMT_AMT)
        put(ws, f'O{rr}', f'=IF(SUM($N${FIRST}:$N${WF_ROW - 1})=0,"",ROUND(($N${TOTAL_ROW}/(1-M{rr})-$N${WF_ROW}-$K${OH_ROW})/SUM($N${FIRST}:$N${WF_ROW - 1}),3))',
            h='right', fmt='0.000')
    grid(ws, f'M{R0}:O{R0 + 5}', outer='thin', vert='thin', horiz='thin')
    put(ws, f'M{R0 + 6}', '※ 必要提出係数は諸経費を現状額に固定した場合の目安', size=9, color=C_GRAYTXT)
    page_setup(ws, f'A1:L{last_note}', fit_h=1)
    return dict(first=FIRST, wf_row=WF_ROW, oh_row=OH_ROW, disc_row=DISC_ROW, total_row=TOTAL_ROW, last_note=last_note,
                margin_table=R0)


# =====================================================================
# 御見積条件
# =====================================================================
def build_conditions(ws, det):
    W = dict(A=5, B=10, C=25, D=20, E=16, F=19)
    set_widths(ws, W)
    BF = W['B'] + W['C'] + W['D'] + W['E'] + W['F']
    put(ws, 'A1', '御見積条件・注意事項', size=14, bold=True)
    ws.row_dimensions[1].height = 24
    put(ws, 'A2', PROJECT_FULL + '　御見積書 添付', size=10)
    ws.row_dimensions[2].height = 16
    ws.row_dimensions[3].height = 8
    r = 4
    put(ws, f'A{r}', '【1】御見積条件・注意事項', size=12, bold=True)
    ws.row_dimensions[r].height = 22
    r += 1
    for i, t in enumerate(NOTES['conditions'], 1):
        put(ws, f'A{r}', f'{i}.', h='right', v='top')
        merge(ws, f'B{r}:F{r}')
        put(ws, f'B{r}', t, wrap=True, v='top')
        ws.row_dimensions[r].height = row_h(est_lines(t, BF, 11), 11, pad=8)
        r += 1
    ws.row_dimensions[r].height = 10
    r += 1
    ws.row_breaks.append(Break(id=r - 1))
    put(ws, f'A{r}', '【2】想定数量（推測）項目の前提', size=12, bold=True)
    ws.row_dimensions[r].height = 22
    r += 1
    merge(ws, f'A{r}:F{r}')
    put(ws, f'A{r}', '図面に数量・経路の記載がなく当社にて想定した項目です。「図面実測＋想定」は平面延長を図面計測し、立上り等を想定加算したものです。実数量確定後に精算させていただきます。', size=10, wrap=True)
    ws.row_dimensions[r].height = row_h(2, 10, pad=4)
    r += 1
    hdr = r
    put(ws, f'A{r}', 'No.', size=10, h='center', fill=FILL_G)
    put(ws, f'B{r}', '内訳書', size=10, h='center', fill=FILL_G)
    put(ws, f'C{r}', '項　　目', size=10, h='center', fill=FILL_G)
    merge(ws, f'D{r}:E{r}')
    put(ws, f'D{r}', '前提・想定内容', size=10, h='center', fill=FILL_G)
    put(ws, f'F{r}', '数量（区分）', size=10, h='center', fill=FILL_G)
    fill_range(ws, f'A{r}:F{r}', FILL_G)
    ws.row_dimensions[r].height = 20
    r += 1
    k = 0
    for ln in det['lines']:
        if ln['basis'] not in ('推測', '実測+推測') or not ln['memo']:
            continue
        k += 1
        sec = next(s for s in IT.SECTIONS if s['key'] == ln['section'])
        item = ln['name'] + ((' ' + ln['spec']) if ln['spec'] else '')
        qtxt = f"{ln['qty']:,} {ln['unit']}（{'想定' if ln['basis'] == '推測' else '実測＋想定'}）"
        put(ws, f'A{r}', k, size=10, h='center', v='center')
        put(ws, f'B{r}', sec['title'][:3], size=10, h='center', v='center')
        put(ws, f'C{r}', item, size=10, wrap=True, v='center', indent=1)
        merge(ws, f'D{r}:E{r}')
        put(ws, f'D{r}', ln['memo'], size=10, wrap=True, v='center', indent=1)
        put(ws, f'F{r}', qtxt, size=10, wrap=True, v='center', h='center')
        nl = max(est_lines(item, W['C'] - 3, 10), est_lines(ln['memo'], W['D'] + W['E'] - 3, 10),
                 est_lines(qtxt, W['F'], 10))
        ws.row_dimensions[r].height = max(18, row_h(nl, 10, pad=4))
        r += 1
    grid(ws, f'A{hdr}:F{r - 1}', outer='medium', vert='thin', horiz='thin')
    ws.row_dimensions[r].height = 10
    r += 1
    put(ws, f'A{r}', '【3】図面読み取りに関する確認事項', size=12, bold=True)
    ws.row_dimensions[r].height = 22
    r += 1
    for i, t in enumerate(NOTES['checks'], 1):
        put(ws, f'A{r}', f'{i}.', h='right', v='top')
        merge(ws, f'B{r}:F{r}')
        put(ws, f'B{r}', t, wrap=True, v='top')
        ws.row_dimensions[r].height = row_h(est_lines(t, BF, 11), 11, pad=8)
        r += 1
    page_setup(ws, f'A1:F{r - 1}', fit_h=0)
    ws.print_title_rows = None
    return dict(last=r - 1, n_assump=k)


# =====================================================================
# 数量拾い根拠（内部用）
# =====================================================================
def build_basis(ws):
    set_widths(ws, dict(A=34, B=14, C=12, D=12, E=12, F=70))
    put(ws, 'A1', '数量拾い根拠（内部用）　※計測方法: 空調設備図 F-02-1（A3 1/50）のベクトル線分（黒・線幅0.72pt）を連結し系統別に延長を集計（data/takeoff_f02.py）。'
        '「実測」=図面線分の平面延長、「想定」=立上り・立下り等の図面に無い部分', size=10, bold=True)
    put(ws, 'A2', '縮尺確認: 通り芯 Y6-Y5 910mm＝51.6pt、Y5-Y4 1,820mm＝103.3pt → 1pt＝17.64mm。拾い根拠画像: takeoff_F02_冷媒ドレン.png', size=9)
    r = 4
    for title, headers, rows in IT.BASIS_TABLES:
        put(ws, f'A{r}', title, bold=True, fill=FILL_L)
        r += 1
        for i, h in enumerate(headers):
            put(ws, f'{chr(65 + i)}{r}', h, size=10, bold=True, h='center', fill=FILL_G)
        hdr = r
        r += 1
        for row in rows:
            for i, v in enumerate(row):
                put(ws, f'{chr(65 + i)}{r}', v, size=10, h=('left' if isinstance(v, str) else 'right'),
                    fmt=('#,##0.00' if isinstance(v, float) else None))
            r += 1
        grid(ws, f'A{hdr}:{chr(64 + len(headers))}{r - 1}', outer='thin', vert='thin', horiz='hair')
        r += 1
    page_setup(ws, f'A1:F{r}', fit_h=0)
    ws.page_setup.orientation = 'landscape'


def build_master(ws):
    set_widths(ws, dict(A=14, B=12, C=9, D=12, E=12, F=14, G=14, H=10, I=80))
    put(ws, 'A1', '単価マスタ（内部）　原価＝材料原価＋人工×労務原価日額、提出（定価ベース）＝原価×提出係数。参考列は宮地機工 参考見積8件（2025.11〜2026.10）から整理（三矢寮見積と共通）', size=10, bold=True)
    put(ws, 'A2', f'労務原価日額 {PV.LABOR_DAY_COST:,} 円/人工　参考見積の人工単価 {PV.LABOR_DAY_LIST:,} 円/人工（6330476-3 配管工費 40人工 929,000）　提出係数 {PV.COEF}　材料係数 ×{PV.MAT_FACTOR}（継手30%＋消耗15%＋支持金物40%）', size=9)
    hdr = ['キー', '材料原価', '人工', '労務原価', '原価', '提出(定価ベース)', '参考定価ベース', '当社/参考', '出典・根拠']
    for i, h in enumerate(hdr, 1):
        put(ws, f'{chr(64 + i)}4', h, size=10, bold=True, h='center', fill=FILL_G)
    r = 5
    for k, (mat, md, ref, src) in PV.MASTER.items():
        put(ws, f'A{r}', k, size=10)
        put(ws, f'B{r}', mat, size=10, h='right', fmt=FMT_AMT)
        put(ws, f'C{r}', md, size=10, h='right', fmt='0.000')
        put(ws, f'D{r}', f'=ROUND(C{r}*LABOR_DAY,0)', size=10, h='right', fmt=FMT_AMT)
        put(ws, f'E{r}', f'=IF(B{r}="","",B{r}+D{r})', size=10, h='right', fmt=FMT_AMT)
        put(ws, f'F{r}', f'=IF(E{r}="","",ROUND(E{r}*COEF,-1))', size=10, h='right', fmt=FMT_AMT)
        put(ws, f'G{r}', (round(ref) if ref else None), size=10, h='right', fmt=FMT_AMT)
        put(ws, f'H{r}', (f'=IF(OR(G{r}="",F{r}=""),"",F{r}/G{r})' if ref else None), size=10, h='right', fmt='0%')
        put(ws, f'I{r}', src, size=9)
        r += 1
    grid(ws, f'A4:I{r - 1}', outer='thin', vert='thin', horiz='hair')
    r += 1
    put(ws, f'A{r}', '参考見積の経費構造', size=10, bold=True, fill=FILL_L)
    for a, b in PV.REF_SOURCES:
        r += 1
        put(ws, f'A{r}', a, size=10, bold=True)
        merge(ws, f'B{r}:I{r}')
        put(ws, f'B{r}', b, size=9, wrap=True)
        ws.row_dimensions[r].height = 28
    page_setup(ws, f'A1:I{r}', fit_h=0)
    ws.page_setup.orientation = 'landscape'


def build_book():
    wb = Workbook()
    base = fnt(11)
    wb._fonts = IndexedList([base])
    wb._named_styles['Normal'].font = base
    wb.properties.creator = COMPANY
    wb.properties.title = '御見積書 ' + PROJECT_FULL
    ws_cover = wb.active
    ws_cover.title = SH_COVER
    ws_det = wb.create_sheet(SH_DETAIL)
    ws_wf = wb.create_sheet(SH_WF)
    ws_cd = wb.create_sheet(SH_COND)
    ws_bs = wb.create_sheet('数量拾い根拠(内部)')
    ws_pm = wb.create_sheet('単価マスタ(内部)')
    det = build_detail(ws_det)
    wf = build_welfare(ws_wf, det)
    cv = build_cover(ws_cover, det, wf)
    cd = build_conditions(ws_cd, det)
    build_basis(ws_bs)
    build_master(ws_pm)
    names = {
        'COEF': f"{q(SH_COVER)}!$N$4",
        'OVERHEAD_RATE': f"{q(SH_COVER)}!$N$5",
        'ROUND_UNIT': f"{q(SH_COVER)}!$N$6",
        'WELFARE_RATE': f"{q(SH_WF)}!{wf['rate_cell']}",
        'LABOR_DAY': f"{q(SH_COVER)}!$N$8",
        'TARGET': f"{q(SH_COVER)}!$N$9",
    }
    for k, ref in names.items():
        wb.defined_names[k] = DefinedName(k, attr_text=ref)
    wb.active = 0
    for ws in wb.worksheets:
        ws.sheet_view.tabSelected = (ws.title == SH_COVER)
    wb.calculation.fullCalcOnLoad = True
    os.makedirs(OUT_DIR, exist_ok=True)
    out = os.path.join(OUT_DIR, FILE)
    wb.save(out)
    return dict(out=out, sections={k: v for k, v in det['sec_rows'].items()}, detail_total=det['total'],
                cover=cv, conditions=cd, quantities=IT.Q)


if __name__ == '__main__':
    print(json.dumps(build_book(), ensure_ascii=False, indent=1))
