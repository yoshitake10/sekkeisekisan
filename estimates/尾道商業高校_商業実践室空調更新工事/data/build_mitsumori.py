# -*- coding: utf-8 -*-
"""尾道商業高校 商業実践室空調更新工事 御見積書
提出済み見積（宮地機工 6330515-2）と同じレイアウトの 表紙＋明細 を、学校の工事内訳書の区分
（1 空調設備機器 ／ 2 空調設備工事 ／ 3 アスベスト含有調査）で作る。

  python3 build_mitsumori.py <出力先ディレクトリ>

シート: 表紙 ／ 明細(1) ／ 明細(2)。1シート＝A4縦1枚（各シート 1×1 ページに合わせて印刷）。
金額の振り分け・按分は build_uchiwake.py と共通（ROWS・unit_prices_*）。
  御見積書_..._出精値引行あり.xlsx  明細は見積書の金額のまま。表紙で 小計 → 出精値引 → 合計（元の見積書と同じ形）
  御見積書_..._値引按分.xlsx       出精値引を全項目に按分した単価。表紙に値引行なし
どちらも合計 1,800,000 円（税別）。金額は数式（キャッシュ値は ../../_lib_est/recalc_cache.py で埋め込む）。
"""
import os
import subprocess
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils.indexed_list import IndexedList
from openpyxl.worksheet.properties import PageSetupProperties

import build_uchiwake as U

HERE = os.path.dirname(os.path.abspath(__file__))
RECALC = os.path.join(HERE, "..", "..", "_lib_est", "recalc_cache.py")

FONT = "ＭＳ 明朝"
NO = "6330515-2"
DATE = "2026年9月30日"
CLIENT = "広島県立尾道商業高等学校　殿"
TITLE = "商業実践室　空調更新工事"
TERMS = [("件名", TITLE), ("受渡場所", "御協議による"), ("受渡期間", "御協議による"),
         ("支払条件", "御協議による"), ("荷造運賃", "御協議による"), ("本書有効期間", "６０日")]
GREETING = ["つぎのとおりお見積いたしました。",
            "なにとぞご用命くださいますようお願い申しあげます。",
            "有効期間経過後のご註文は一応ご照会下さい"]
WELFARE = ["上記の工事費には、以下の法定福利費相当額を含んでいます。",
           "【法定福利相当額（Ａ）】29,016円　　　Ａ ＝ Ｙ × Ｚ",
           "　Ｙ：労務費相当額",
           "　Ｚ：法定福利費事業者負担率（合計値：　16.04／100）"]
FOOTNOTE = "※本見積は、2026/11月からの価格で作成しております。"
BASENAME = f"御見積書_{NO}_尾道商業高校_商業実践室空調更新工事"

SH_COVER, SH_D1, SH_D2 = "表紙", "明細(1)", "明細(2)"

# 表紙: 項 A ／ 品名 B:C ／ 数量 D ／ 金額 E:F ／ 備考 G:H（E〜H は印欄の4枠も兼ねる）
COVER_W = dict(A=5.5, B=9, C=33.5, D=10.5, E=9, F=9, G=9, H=9)
COVER_TABLE_ROWS = 25
# 明細: 品名 A ／ 数量 B(数)・C(単位) ／ 単価 D:E ／ 価格 F ／ 備考 G
DETAIL_W = dict(A=35, B=6, C=4.5, D=7.5, E=7.5, F=16, G=18)
DETAIL_ROWS = 45

FMT_AMT = "#,##0;-#,##0"
FMT_PAREN = "(#,##0)"
THIN = Side(style="thin")
HAIR = Side(style="hair")

# 空調機の仕様（見積書 p2 のまま）。(項目, 値, 単位)
SPEC = [("グリーン購入法適合", None, None),
        ("定格冷房標準能力", "25.0(11.3～28.0)", "ｋＷ"),
        ("定格暖房標準能力", "28.0(12.6～35.0)", "ｋＷ"),
        ("ＡＰＦ２０１５", "4.8", None),
        ("定格冷房消費電力", "9.75", "ｋＷ"),
        ("定格暖房消費電力", "8.57", "ｋＷ"),
        ("室内ファン風量強", "(28)×2", "ｍ３／ｍｉｎ"),
        ("室内ファン電動機出力", "(150×1)×2", "Ｗ"),
        ("室外ファン風量", "163", "ｍ３／ｍｉｎ"),
        ("室外ファン電動機出力", "(227+227)×1", "Ｗ"),
        ("圧縮機電動機出力", "5.90", "ｋＷ"),
        ("電源　３相　２００Ｖ６０Ｈｚ", None, None),
        ("外形寸法（Ｗ×Ｌ×Ｈ）・質量", None, None),
        ("　室内機　1590× 690× 235mm", "40kg　×2", None),
        ("　室外機　 940× 320×1430mm", "123kg　×1", None)]

# 様式の行 → 明細の品名・単位・備考（品名は学校の工事内訳書のまま）
NAME = {7: "SSRH280DD", 8: "BRC1G4", 9: "BRE50B2F",
        13: "搬入据付工事費", 14: "基礎工事費", 15: "リモコン配線工事費", 16: "冷媒配管工事費",
        17: "重機費（高所作業車）", 18: "既設機器撤去工事費", 19: "配管保温工事費（カラー鉄板ラッキング）",
        20: "電気工事費", 21: "試運転調整費", 22: "冷媒回収、破壊処理費", 23: "雑材消耗品",
        24: "資材運搬交通費", 25: "現場雑費",
        29: "検体調査", 30: "現場収集", 31: "雑材消耗品", 32: "資材運搬交通費", 33: "現場雑費"}
UNIT = {7: "台"}
NOTE = {7: "内機2台・外機1台", 16: "ドレン管・連絡線含む", 17: "レッカー等",
        19: "ラッキング切廻し含む", 20: "電源線脱着", 25: "工事諸経費含む"}
DESC = {8: "運転リモコン・液晶ワイヤード", 9: "センシングユニット（センサーキット）"}


def f(size=10, bold=False, underline=None):
    return Font(name=FONT, size=size, bold=bold, underline=underline)


def put(ws, ref, value, size=10, h=None, v="center", indent=0, fmt=None, underline=None, shrink=False):
    c = ws[ref]
    if isinstance(value, str) and indent:
        value = "　" * indent + value  # 字下げは全角スペース（Excel の インデント は LibreOffice で効かないため）
    c.value = None if value == "" else value
    c.font = f(size, underline=underline)
    c.alignment = Alignment(horizontal=h, vertical=v, shrink_to_fit=shrink)
    if fmt:
        c.number_format = fmt
    return c


def box(ws, rng, left=True, right=True, top=True, bottom=True):
    """範囲の外周だけ罫線を引く（結合セル用）"""
    from openpyxl.utils import range_boundaries
    c1, r1, c2, r2 = range_boundaries(rng)
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            cell = ws.cell(r, c)
            b = cell.border
            cell.border = Border(left=THIN if (left and c == c1) else b.left,
                                 right=THIN if (right and c == c2) else b.right,
                                 top=THIN if (top and r == r1) else b.top,
                                 bottom=THIN if (bottom and r == r2) else b.bottom)


def setup(ws, widths, last_col, last_row):
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.sheet_view.showGridLines = False
    ws.print_area = f"A1:{last_col}{last_row}"
    ps = ws.page_setup
    ps.paperSize = ws.PAPERSIZE_A4
    ps.orientation = "portrait"
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ps.fitToWidth = 1
    ps.fitToHeight = 1
    ws.print_options.horizontalCentered = True
    m = ws.page_margins
    m.left, m.right, m.top, m.bottom, m.header, m.footer = 0.45, 0.45, 0.5, 0.4, 0.2, 0.2


# ---------------------------------------------------------------------------
# 明細
# ---------------------------------------------------------------------------
def detail_lines(unit, page):
    """明細の行。kind: sec / text / item / desc / spec / sum / blank"""
    L = []
    if page == 1:
        L += [("sec", "1 空調設備機器"), ("text", "ACP-1"), ("text", "スカイエア"),
              ("text", "ＦＩＶＥ　ＳＴＡＲ　天井吊形"), ("item", 7)]
        L += [("spec",) + s for s in SPEC]
        L += [("item", 8), ("desc", DESC[8]), ("item", 9), ("desc", DESC[9]), ("blank",),
              ("sum", "機器合計", [7, 8, 9]), ("blank",),
              ("sec", "2 空調設備工事")]
        L += [("item", r) for r in range(13, 26)]
        L += [("sum", "－ 計 －", list(range(13, 26)))]
    else:
        L += [("sec", "3 アスベスト含有調査")]
        L += [("item", r) for r in range(29, 34)]
        L += [("sum", "－ 計 －", list(range(29, 34)))]
    assert len(L) <= DETAIL_ROWS, len(L)
    return L


def build_detail(ws, unit, page):
    setup(ws, DETAIL_W, "G", 5 + DETAIL_ROWS)
    for r, h in {1: 6, 2: 14, 3: 14, 4: 8, 5: 18}.items():
        ws.row_dimensions[r].height = h
    ws.merge_cells("F2:G2")
    ws.merge_cells("F3:G3")
    put(ws, "F2", f"PAGE No. {page + 1}")
    put(ws, "F3", f"見積番号　{NO}")
    # 見出し
    ws.merge_cells("B5:C5")
    ws.merge_cells("D5:E5")
    for ref, txt in [("A5", "品　　　　名"), ("B5", "数　量"), ("D5", "単 価(税別)"),
                     ("F5", "価 格(税別)"), ("G5", "備　　　考")]:
        put(ws, ref, txt, h="center")
    for rng in ("A5", "B5:C5", "D5:E5", "F5", "G5"):
        box(ws, rng)

    rows_of = {}
    lines = detail_lines(unit, page)
    first = 6
    for i in range(DETAIL_ROWS):
        r = first + i
        ws.row_dimensions[r].height = 15.75
        ln = lines[i] if i < len(lines) else ("blank",)
        k = ln[0]
        spec_like = k in ("spec", "desc")
        if not spec_like:
            ws.merge_cells(f"D{r}:E{r}")
        # 罫線: 横線は全行。縦線は 品名|数量・数量|単価 を仕様行では省く（見積書と同じ）
        for col in "ABCDEFG":
            ws[f"{col}{r}"].border = Border(
                top=HAIR, bottom=HAIR,
                left=THIN if col in ("A", "F", "G") or (col in ("B", "D") and not spec_like) else None,
                right=THIN if col == "G" else None)
        if k == "sec":
            put(ws, f"A{r}", ln[1])
        elif k == "text":
            put(ws, f"A{r}", ln[1])
        elif k == "desc":
            put(ws, f"A{r}", ln[1], indent=1)
        elif k == "spec":
            label, val, unit_label = ln[1:]
            put(ws, f"A{r}", label, indent=1)
            if val:
                ws.merge_cells(f"B{r}:D{r}")
                put(ws, f"B{r}", val, h="left")
            if unit_label:
                put(ws, f"E{r}", unit_label, h="left")
        elif k == "item":
            row = ln[1]
            qty = U.ROWS[row][0]
            rows_of[row] = r
            put(ws, f"A{r}", NAME[row], indent=0 if row <= 9 else 1, shrink=True)
            put(ws, f"B{r}", qty, h="right", fmt="0")
            put(ws, f"C{r}", " " + UNIT.get(row, "式"), h="left")
            if qty > 1 or row <= 9:
                put(ws, f"D{r}", unit[row], h="right", fmt=FMT_AMT)
                put(ws, f"F{r}", f"=B{r}*D{r}", h="right", fmt=FMT_AMT)
            else:
                put(ws, f"F{r}", unit[row], h="right", fmt=FMT_AMT)
            if row in NOTE:
                put(ws, f"G{r}", NOTE[row], size=9, shrink=True)
        elif k == "sum":
            label, members = ln[1], ln[2]
            put(ws, f"A{r}", label, indent=1)
            refs = ",".join(f"F{rows_of[m]}" for m in members)
            if len(members) > 3:
                refs = f"F{rows_of[members[0]]}:F{rows_of[members[-1]]}"
            put(ws, f"F{r}", f"=SUM({refs})", h="right", fmt=FMT_AMT)
            rows_of[label + str(page)] = r
    # 外枠
    box(ws, f"A{first}:G{first + DETAIL_ROWS - 1}")
    return rows_of


# ---------------------------------------------------------------------------
# 表紙
# ---------------------------------------------------------------------------
def build_cover(ws, rows1, rows2, with_discount):
    setup(ws, COVER_W, "H", 22 + COVER_TABLE_ROWS)
    heights = {1: 6, 2: 30, 19: 22, 20: 6, 21: 18}
    for r in range(3, 19):
        ws.row_dimensions[r].height = 14.25
    for r, h in heights.items():
        ws.row_dimensions[r].height = h

    ws.merge_cells("A2:H2")
    put(ws, "A2", "御見積書", size=20, h="center")
    ws.merge_cells("A3:D3")
    put(ws, "A3", CLIENT, size=11)
    ws.merge_cells("G3:H3")
    put(ws, "G3", "PAGE No. 1", h="right")
    right = {4: f"見積番号　{NO}", 5: DATE, 6: "宮地機工株式会社", 10: "〒722-0051",
             11: "広島県尾道市東尾道9-9", 13: "TEL 0848-20-2121", 14: "FAX 0848-20-2126"}
    for r, txt in right.items():
        ws.merge_cells(f"E{r}:H{r}")
        put(ws, f"E{r}", txt, size=11 if r == 6 else 10)
    for i, txt in enumerate(GREETING):
        ws.merge_cells(f"A{6 + i}:D{6 + i}")
        put(ws, f"A{6 + i}", txt)
    for i, (label, val) in enumerate(TERMS):
        r = 13 + i
        ws.merge_cells(f"A{r}:B{r}")
        ws.merge_cells(f"C{r}:D{r}")
        put(ws, f"A{r}", label)
        put(ws, f"C{r}", val)
    # 印欄（4枠）
    for col in "EFGH":
        ws.merge_cells(f"{col}15:{col}17")
        box(ws, f"{col}15:{col}17")

    # 表
    hdr = 21
    for rng in ("B21:C21", "E21:F21", "G21:H21"):
        ws.merge_cells(rng)
    for ref, txt in [("A21", "項"), ("B21", "品　　　　　名"), ("D21", "数　量"),
                     ("E21", "金　　額"), ("G21", "備　　　考")]:
        put(ws, ref, txt, h="center")

    d1 = lambda key: f"'{SH_D1}'!F{rows1[key]}"
    d2 = lambda key: f"'{SH_D2}'!F{rows2[key]}"
    T = [("", "空気調和工事", "", None, None),
         ("1", "空調設備機器", "一式", "=" + d1("機器合計1"), None),
         ("2", "空調設備工事", "一式", "=" + d1("－ 計 －1"), None),
         ("3", "アスベスト含有調査", "一式", "=" + d2("－ 計 －2"), None),
         None,
         ("", "機器本体", "", None, None),
         ("", "　スカイエア", "一式", None, "=" + d1(7)),
         ("", "別売付属品", "", None, None),
         ("", "　ダイキン別売品", "一式", None, f"={d1(8)}+{d1(9)}")]
    first = hdr + 1
    items = [first + 1, first + 2, first + 3]
    if with_discount:
        sub = first + len(T)
        T += [("", "　小計", "", "=SUM(" + ",".join(f"E{r}" for r in items) + ")", None),
              ("4", "出精値引", "", U.DISCOUNT, None),
              ("", "　合計", "", f"=E{sub}+E{sub + 1}", None)]
    else:
        T += [("", "　合計", "", "=SUM(" + ",".join(f"E{r}" for r in items) + ")", None)]
    total_row = first + len(T) - 1
    T += [None]
    note_from = first + len(T)
    T += [("note", t) for t in WELFARE]

    for i in range(COVER_TABLE_ROWS):
        r = first + i
        ws.row_dimensions[r].height = 15.75
        ln = T[i] if i < len(T) else None
        is_note = ln is not None and ln[0] == "note"
        if is_note:
            ws.merge_cells(f"B{r}:H{r}")
            put(ws, f"B{r}", ln[1])
        else:
            ws.merge_cells(f"B{r}:C{r}")
            ws.merge_cells(f"E{r}:F{r}")
            ws.merge_cells(f"G{r}:H{r}")
            if ln:
                no, name, qty, amt, note = ln
                put(ws, f"A{r}", no, h="center")
                put(ws, f"B{r}", name)
                put(ws, f"D{r}", qty, h="center")
                if amt is not None:
                    put(ws, f"E{r}", amt, h="right", fmt=FMT_AMT)
                if note is not None:
                    put(ws, f"G{r}", note, h="right", fmt=FMT_PAREN)
        for col in "ABCDEFGH":
            ws[f"{col}{r}"].border = Border(
                top=HAIR, bottom=HAIR,
                left=THIN if col in ("A", "B") or (col in ("D", "E", "G") and not is_note) else None,
                right=THIN if col == "H" else None)
    for rng in ("A21", "B21:C21", "D21", "E21:F21", "G21:H21"):
        box(ws, rng)
    box(ws, f"A{first}:H{first + COVER_TABLE_ROWS - 1}")
    # 合計の上に太めの区切り（元の見積書と同じく 合計 を強調）
    for col in "BCDEFGH":
        c = ws[f"{col}{total_row}"]
        c.border = Border(left=c.border.left, right=c.border.right, top=THIN, bottom=THIN)

    ws.merge_cells("B19:G19")
    put(ws, "B19", f'="お見積金額　　￥"&TEXT(E{total_row},"#,##0")&"-　(税別)"', size=13,
        h="center", underline="single")
    last = first + COVER_TABLE_ROWS
    ws.row_dimensions[last].height = 16
    ws.merge_cells(f"B{last}:H{last}")
    put(ws, f"B{last}", FOOTNOTE)
    ws.print_area = f"A1:H{last}"
    return total_row


def build(path, unit, with_discount):
    wb = Workbook()
    base = Font(name=FONT, size=10)
    wb._fonts = IndexedList([base])
    wb._named_styles["Normal"].font = base
    wb.properties.creator = "宮地機工株式会社"
    wb.properties.title = f"御見積書 {NO} 広島県立尾道商業高等学校 {TITLE}"
    ws_c = wb.active
    ws_c.title = SH_COVER
    ws_1 = wb.create_sheet(SH_D1)
    ws_2 = wb.create_sheet(SH_D2)
    rows1 = build_detail(ws_1, unit, 1)
    rows2 = build_detail(ws_2, unit, 2)
    build_cover(ws_c, rows1, rows2, with_discount)
    wb.save(path)
    subprocess.run([sys.executable, RECALC, path], check=True, stdout=subprocess.DEVNULL)
    vj = path + ".values.json"
    import json
    vals = json.load(open(vj, encoding="utf-8"))
    os.remove(vj)
    return vals


def main():
    out_dir = sys.argv[1] if len(sys.argv) > 1 else "."
    U.check_source()
    for name, unit, disc in [("出精値引行あり", U.unit_prices_gross(), True),
                             ("値引按分", U.unit_prices_prorated(), False)]:
        path = os.path.join(out_dir, f"{BASENAME}_{name}.xlsx")
        vals = build(path, unit, disc)
        cover = {k: v for k, v in vals.items() if k.startswith(SH_COVER)}
        print(name, path)
        for k, v in cover.items():
            print("  ", k, v)
        assert any(v == U.ESTIMATE_TOTAL for v in cover.values()), "合計が 1,800,000 になっていない"


if __name__ == "__main__":
    main()
