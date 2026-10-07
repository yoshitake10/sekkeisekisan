# -*- coding: utf-8 -*-
"""宮地機工様式 見積書 共通書式ヘルパー（参照セッションのビルダーから流用）"""
import math, unicodedata
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import range_boundaries
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.properties import PageSetupProperties
# =====================================================================
# 書式ヘルパー
# =====================================================================
FONT_NAME = 'ＭＳ Ｐゴシック'
C_YELLOW = 'FFF2CC'   # 入力セル
C_GRAY = 'D9D9D9'     # 見出し・内部列見出し
C_LABEL = 'F2F2F2'    # 帳票ラベル（法定福利費の見出し欄）
C_GRAYTXT = '7F7F7F'  # 内部用の灰文字
FILL_Y = PatternFill('solid', fgColor=C_YELLOW)
FILL_G = PatternFill('solid', fgColor=C_GRAY)
FILL_L = PatternFill('solid', fgColor=C_LABEL)
FMT_AMT = '#,##0_ '   # 金額（右端に半角1字分の余白）
FMT_QTY = '#,##0_ '   # 数量（数量は全件整数のため小数表示なし）

# LibreOffice(IPAゴシック代替)での較正値: 列幅1単位 = 5.5pt（Excel/ＭＳ Ｐゴシックは約6pt なので安全側）
PT_PER_UNIT = 5.5


def fnt(size=11, bold=False, color=None):
    return Font(name=FONT_NAME, size=size, bold=bold, color=color, charset=128, family=3)


def put(ws, ref, value=None, size=11, bold=False, h=None, v='center', wrap=False, shrink=False,
        fmt=None, fill=None, color=None, indent=0):
    c = ws[ref]
    if value is not None:
        c.value = value
    c.font = fnt(size, bold, color)
    if indent and h is None:
        h = 'left'  # インデントは左詰め指定時のみ有効（標準配置では Excel が無視する）
    c.alignment = Alignment(horizontal=h, vertical=v, wrap_text=wrap, shrink_to_fit=shrink, indent=indent)
    if fmt:
        c.number_format = fmt
    if fill is not None:
        c.fill = fill
    return c


def merge(ws, rng):
    ws.merge_cells(rng)
    return rng.split(':')[0]


def sd(style):
    return Side(style=style, color='000000') if style else Side()


def set_border(cell, left=None, right=None, top=None, bottom=None):
    """指定した辺だけ上書き（None は既存を維持）"""
    b = cell.border
    cell.border = Border(
        left=sd(left) if left is not None else b.left,
        right=sd(right) if right is not None else b.right,
        top=sd(top) if top is not None else b.top,
        bottom=sd(bottom) if bottom is not None else b.bottom,
    )


def grid(ws, rng, outer='medium', vert='thin', horiz='thin'):
    """範囲に罫線: 外枠 outer / 内側縦 vert / 内側横 horiz"""
    c1, r1, c2, r2 = range_boundaries(rng)
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            ws.cell(r, c).border = Border(
                left=sd(outer if c == c1 else vert),
                right=sd(outer if c == c2 else vert),
                top=sd(outer if r == r1 else horiz),
                bottom=sd(outer if r == r2 else horiz),
            )


def hline(ws, rng, pos, style):
    """範囲の上辺/下辺に線"""
    c1, r1, c2, r2 = range_boundaries(rng)
    for c in range(c1, c2 + 1):
        if pos == 'top':
            set_border(ws.cell(r1, c), top=style)
            if r1 > 1:
                set_border(ws.cell(r1 - 1, c), bottom=style)
        else:
            set_border(ws.cell(r2, c), bottom=style)
            set_border(ws.cell(r2 + 1, c), top=style)


def fill_range(ws, rng, fill):
    c1, r1, c2, r2 = range_boundaries(rng)
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            ws.cell(r, c).fill = fill


def set_widths(ws, widths):
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


KINSOKU_HEAD = set('、。，．・：；？！）」』】〕〉》］｝’”ー々ゝゞヽヾぁぃぅぇぉっゃゅょゎァィゥェォッャュョヮヵヶ％,.)]}:;!?%')
KINSOKU_TAIL = set('（「『【〔〈《［｛‘“([{')


def _cw(ch, pt):
    """文字幅(pt)。IPAゴシック基準: 全角・曖昧幅=1em、半角=0.5em"""
    return pt if unicodedata.east_asian_width(ch) in 'FWA' else pt / 2.0


def _is_word_char(ch):
    """英数字列（ギリシャ文字 φ 等を含む）は単語単位で折り返される"""
    if ch == ' ':
        return False
    if ord(ch) < 0x80:
        return True
    return ord(ch) < 0x2E80 and unicodedata.category(ch)[0] in 'LN'


def est_lines(text, width_units, pt, pad_units=0.8, safety=0.97):
    """折返し行数の見積り（貪欲折返しシミュレーション）。
    英数字列は単語単位、行頭禁則文字は前の字に、行末禁則文字は次の字に付けて分割不可とする（安全側）。"""
    cap = (width_units - pad_units) * PT_PER_UNIT * safety
    total = 0
    for para in str(text).split('\n'):
        toks, buf = [], ''
        for ch in para:
            if _is_word_char(ch):
                buf += ch
                continue
            if buf:
                toks.append(buf)
                buf = ''
            toks.append(ch)
        if buf:
            toks.append(buf)
        merged = []
        for t in toks:
            if merged and t != ' ' and merged[-1] != ' ' and (t[0] in KINSOKU_HEAD or merged[-1][-1] in KINSOKU_TAIL):
                merged[-1] += t
            else:
                merged.append(t)
        lines, cur = 1, 0.0
        for t in merged:
            w = sum(_cw(c, pt) for c in t)
            if t == ' ':
                if cur > 0 and cur + w <= cap:
                    cur += w
                continue
            if cur + w <= cap:
                cur += w
            elif w <= cap:
                lines += 1
                cur = w
            else:
                for c in t:
                    cw = _cw(c, pt)
                    if cur + cw > cap:
                        lines += 1
                        cur = 0.0
                    cur += cw
        total += lines
    return total


def text_width_pt(text, pt):
    return sum(_cw(c, pt) for c in str(text))


def row_h(lines, pt, pad=3.0):
    """行高(pt)。Excel(ＭＳ Ｐゴシック)の行送り ≒ 1.25em を基準（LibreOffice は 1.0em なので余裕側）"""
    return round(lines * pt * 1.25 + pad, 1)


def page_setup(ws, print_area, fit_h=0, title_rows=None, footer=None):
    ws.print_area = print_area
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = 'portrait'
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = fit_h
    cm = 1 / 2.54
    ws.page_margins = PageMargins(left=1.2 * cm, right=1.2 * cm, top=1.5 * cm, bottom=1.5 * cm,
                                  header=0.8 * cm, footer=0.6 * cm)
    ws.print_options.horizontalCentered = True
    ws.print_options.gridLines = False
    ws.sheet_view.showGridLines = False
    if title_rows:
        ws.print_title_rows = title_rows
    if footer:
        ws.oddFooter.center.text = footer
        ws.oddFooter.center.size = 9
        ws.oddFooter.center.font = FONT_NAME + ',標準'


