#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""提出用 PDF の作成（内部シートを除いた印刷用コピーを LibreOffice で PDF 化）
使い方: python3 make_pdf.py <御見積書.xlsx> <出力ディレクトリ>
内部シート（数量拾い根拠・単価マスタ）は原価情報を含むため PDF に含めない。内訳書の内部列は印刷範囲外。"""
import os
import shutil
import subprocess
import sys
import tempfile

from openpyxl import load_workbook

INTERNAL = ('数量拾い根拠(内部)', '単価マスタ(内部)', 'エスト比較(内部)', 'エストマスタ(参考)', '原価根拠(内部)')

src, out_dir = sys.argv[1], sys.argv[2]
tmpd = tempfile.mkdtemp()
wb = load_workbook(src)
for name in INTERNAL:
    if name in wb.sheetnames:
        wb.remove(wb[name])
base = os.path.splitext(os.path.basename(src))[0]
tmp_xlsx = os.path.join(tmpd, base + '.xlsx')
wb.save(tmp_xlsx)
profile = 'file://' + os.path.join(tmpd, 'lo_profile')
subprocess.run(['soffice', f'-env:UserInstallation={profile}', '--headless', '--convert-to', 'pdf',
                '--outdir', tmpd, tmp_xlsx], check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
os.makedirs(out_dir, exist_ok=True)
shutil.move(os.path.join(tmpd, base + '.pdf'), os.path.join(out_dir, base + '.pdf'))
shutil.rmtree(tmpd, ignore_errors=True)
print(os.path.join(out_dir, base + '.pdf'))
