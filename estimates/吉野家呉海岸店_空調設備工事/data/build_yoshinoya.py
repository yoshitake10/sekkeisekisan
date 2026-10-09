#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
吉野家 呉海岸店 空調設備工事 御見積書（宮地機工様式・エスト準拠）生成  v2
=========================================================================
エスト見積 6330560-1 と同じ並び・単価・経費率・工費算式で内訳書を作り、
原価（建設物価相当の材料＋人工×労務原価日額＋法定福利費）から粗利率20%の提出額を求めて出精値引で合わせる。
エストに単価がない項目は「要見積」（既定は別途見積）／「単価要確認」としてアラートを出す。

シート: 表紙(工事) ／ 内訳書(空調設備工事) ／ 法定福利費内訳明細書 ／ 御見積条件（以上が提出用）
        ／ エスト比較(内部) ／ エストマスタ(参考) ／ 数量拾い根拠(内部) ／ 原価根拠(内部)
使い方: python3 build_yoshinoya.py <出力ディレクトリ>
"""
import datetime
import json
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from helpers import (fnt, put, merge, set_border, grid, hline, fill_range, set_widths,  # noqa: E402
                     est_lines, text_width_pt, row_h, page_setup,
                     FILL_Y, FILL_G, FILL_L, C_GRAYTXT, FMT_AMT, FMT_QTY, PT_PER_UNIT)
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import PatternFill  # noqa: E402
from openpyxl.utils.indexed_list import IndexedList  # noqa: E402
from openpyxl.workbook.defined_name import DefinedName  # noqa: E402
from openpyxl.worksheet.pagebreak import Break  # noqa: E402
import pricing as PV  # noqa: E402
import items as IT  # noqa: E402
import lines as LN  # noqa: E402
import est_ref as ES  # noqa: E402

OUT_DIR = sys.argv[1] if len(sys.argv) > 1 else HERE

# =====================================================================
# 案件情報（エスト見積 6330560-1 に合わせる）
# =====================================================================
CLIENT = ES.EST_CLIENT
PROJECT_1 = '吉野家　呉海岸店'
PROJECT_2 = '空調設備工事'
PROJECT_FULL = '吉野家　呉海岸店　空調設備工事'
SITE = '広島県呉市海岸（国道31号沿い）'
TERM = '工程表による（全体工期 2026年9月1日～2026年11月16日）'
PAYMENT = '御協議による'
VALIDITY = '10日'
EST_DATE = datetime.date(2026, 10, 8)
DATE_FMT = 'yyyy"年"m"月"d"日"'
COMPANY = '宮地機工株式会社'
REP = '代表取締役　濱本　義樹'
ZIP = '〒722-0051'
ADDR = '広島県尾道市東尾道9-9'
TEL = 'TEL.0848-20-2121'
FAX = 'FAX.0848-20-2126'
STAFF = '濱本　義武'
FILE = '御見積書_吉野家呉海岸店_空調設備工事_20261008.xlsx'

SH_COVER = '表紙(工事)'
SH_DETAIL = '内訳書(空調設備工事)'
SH_WF = '法定福利費内訳明細書'
SH_COND = '御見積条件'
SH_CMP = 'エスト比較(内部)'
SH_MST = 'エストマスタ(参考)'
SH_BASIS = '数量拾い根拠(内部)'
SH_COST = '原価根拠(内部)'

WELFARE_RATES = [
    ('健康保険料', 0.04890, '健康保険料率 9.780%（協会けんぽ 広島県 令和8年度）×1/2（労使折半）'),
    ('介護保険料', 0.00810, '介護保険料率 1.620%（全国一律 令和8年度）×1/2（労使折半）'),
    ('子ども・子育て支援金', 0.00115, '支援金率 0.230%（令和8年4月分〜）×1/2（労使折半）'),
    ('厚生年金保険料', 0.09150, '厚生年金保険料率 18.300%（平成29年9月以降固定）×1/2（労使折半）'),
    ('雇用保険料', 0.01050, '雇用保険料率（建設の事業）事業主負担分 1.050%'),
    ('子ども・子育て拠出金', 0.00360, '拠出金率 0.360%（事業主全額負担）'),
]
CIRC = '①②③④⑤⑥⑦⑧⑨'
C_RED = 'C00000'
FILL_RED = PatternFill('solid', fgColor='F8CBAD')
FILL_PINK = PatternFill('solid', fgColor='FDE9E9')
FILL_AMB = PatternFill('solid', fgColor='FFE699')
FMT_SIGNED = '#,##0_ ;[Red]-#,##0_ '
FMT_TBD = '#,##0_ ;[Red]-#,##0_ ;"別途見積"_ '
FMT_SUPPLY = '#,##0_ ;[Red]-#,##0_ ;"支給品"_ '


def q(sheet):
    return "'" + sheet + "'"


NOTES = {
    'cover': [
        '本見積は2026年11月からの価格で作成しております。数量は換気空調機器表（F-01）・空調設備図（F-02-1）・平面図（A-08）より当社にて拾い出しており、実施数量との差異は別途精算とさせていただきます。',
        '空調機（AC-1・AC-2・RAC の室内機・室外機・リモコン）および吹出しパンカーは支給品のため含みません（搬入・吊込み・据付費を計上）。',
        '鉄骨架台（2段積み・溶融亜鉛メッキ）は製作品のため別途御見積とさせていただきます。',
        'リモコン配線、AC-2 吹出しダクト・吹出しパンカー取付、換気設備、電源・接地等の電気工事、建築工事、揚重機、足場・仮設、撤去・産業廃棄物処分は含みません。',
        '本見積書に記載なき事項については別途とさせて頂きます。工事に係る電気・水道は無償支給願います。本見積金額には消費税は含まれておりません。',
        '詳細条件は別紙「御見積条件」をご参照ください。',
    ],
    'conditions': [
        '本見積は2026年11月からの価格で作成しております。数量は換気空調機器表 F-01（2025.05.25）・空調設備図 F-02-1（2026.05.25）・平面図 A-08（2026.07.29）に基づき当社にて拾い出しています。冷媒管・ドレン管の平面延長は空調設備図（1/50）から計測し、立上り・立下り・機器接続部および内外連絡線は想定数量として計上しています（別紙【2】参照）。',
        '工事範囲: 空調機の搬入・吊込み・据付（AC-1 天井カセット形×2台、AC-2 天井ビルトイン形×1台、RAC 壁掛形×1台、室外機×4台）、室外機基礎（縁石）、冷媒配管（被覆銅管 液管10t・ガス管20t）・屋外化粧カバー、ドレン配管（VP）・保温（GW20mm＋アルミガラスクロス）、内外連絡線（VVF 2.0-3C）、外壁スリーブ、気密試験・真空引き、試運転調整。',
        '空調機（AC-1 SSRC140C×2台、AC-2 SSRB140C×1台、RAC×1台、各リモコン）および吹出しパンカー PK-SB#10TK×4台は支給品のため含みません。',
        '鉄骨架台（2段積み・溶融亜鉛メッキ防錆処理、防振ゴム共）は空調設備図の特記により空調設備工事ですが、製作品のため別途御見積とさせていただきます。架台はダイキン製室外機（H1,430mm）に対応する段間寸法とし、背面と建物壁面の隙間200mm以上を確保します。',
        'リモコン配線、AC-2（天井ビルトイン形）から吹出しパンカーまでのダクト・パンカー取付、換気設備（換気扇・送風機・フード・ダクト類）、電源・接地等の電気工事は含みません。',
        '冷媒の追加充填は、各系統の配管長（最長約14m）が支給機のチャージレス配管長以内と想定し計上していません（真空引き・ガス充填は基本料金）。必要となった場合は実費にて精算させていただきます。',
        'ドレン管は硬質塩化ビニル管（VP 25A・30A・40A）とし、範囲は各機器から「空調機ドレン接続立上り（H=150）」まで、以降は給排水工事です。図示のドレンホース用逆止弁 NDB-20-25 を計上しています。',
        '揚重機（クレーン等）、足場・高所作業車等の仮設、天井・壁の開口補強および仕上げ補修、天井点検口、既設機器の撤去、産業廃棄物（支給品の梱包材を含む）の処分は含みません。',
        '工期は工程表（吉野家呉海岸通店新設工事 工程表）に基づく通常の日中作業を前提としております。夜間・休日作業、工程変更や他工事との輻輳による手待ち、工期延長に伴う経費増は別途協議とさせていただきます。',
        '労務費に係る法定福利費（事業主負担分）は工事費に含んでおり、その額は御見積書表紙および別紙「法定福利費内訳明細書」のとおりです。',
        '本見積書に記載なき事項については別途とさせていただきます。工事に係る電気・水道は無償支給願います。本見積金額には消費税は含まれておりません。本書有効期間は10日です。',
    ],
    'checks': [
        '空調機器表の AC-1（天井カセット形 5馬力相当）は2台（室内機2台・室外機2台の2系統）、AC-2 は1系統として計上しています（空調設備図 F-02-1・室外機設置概要図による）。',
        '室外機は室外機設置概要図により、架台①＝AC-1（下段）＋AC-1（上段）、架台②＝AC-2（下段）＋RAC（上段）の2段積み2基としています。',
        '冷媒管・ドレン管は空調設備図 F-02-1 の配管線（R・D 記号）から計測し、立上り・立下りは室外機設置概要図と天井高（客席 CH2,600、厨房・店長室 CH2,500）から想定しています。',
        '特記の「天井内冷媒配管は厚20耐熱ポリスチレンフォーム断熱材巻き」は、冷媒管用の耐熱ポリエチレンフォーム保温材（ガス管20mm）と解釈しています。',
        'ドレン接続立上り（H=150）の立上り管および以降の排水管は給排水工事と解釈しています。',
        '施工場所は位置図（呉海岸）によります。地番をご教示ください。',
    ],
}


# =====================================================================
# 内訳書
# =====================================================================
COLS = dict(J='単価区分', K='アラート', L='材料原価/単位', M='人工/単位', N='労務原価/単位', O='原価単価', P='原価金額',
            Q='労務原価金額', R='労務費相当額(提出)', S='粗利(定価)', T='粗利率', U='率・歩掛', V='要見積 概算原価/単位',
            W='エスト数量', X='エスト単価', Y='エスト金額', Z='差額(当社-エスト)', AA='エストマスタ参考', AB='数量根拠', AC='出典・メモ（内部）')


def build_detail(ws):
    set_widths(ws, dict(A=4, B=27, C=24, D=4, E=8, F=5, G=10, H=13, I=15,
                        J=12, K=11, L=10, M=7, N=10, O=10, P=12, Q=11, R=12, S=11, T=8, U=8, V=11,
                        W=7, X=9, Y=11, Z=11, AA=40, AB=11, AC=70))
    merge(ws, 'A1:I1')
    put(ws, 'A1', '内　訳　書', size=16, bold=True, h='center')
    merge(ws, 'J1:AC1')
    put(ws, 'J1', '原　価・エスト比較（内部用・印刷範囲外）　黄セル＝入力、赤＝要見積、橙＝単価要確認・エストと数量差',
        size=11, bold=True, h='center', color=C_GRAYTXT, fill=FILL_G)
    fill_range(ws, 'J1:AC1', FILL_G)
    ws.row_dimensions[1].height = 26
    put(ws, 'A2', '工事名：' + PROJECT_FULL)
    put(ws, 'I2', '内訳書 No.1', h='right')
    ws.row_dimensions[2].height = 18
    merge(ws, 'A3:B3')
    merge(ws, 'C3:D3')
    for ref, txt in [('A3', '品　　名'), ('C3', '規格・寸法'), ('E3', '数 量'), ('F3', '単位'),
                     ('G3', '単　価'), ('H3', '金　　額'), ('I3', '備　考')]:
        put(ws, ref, txt, h='center', fill=FILL_G)
    fill_range(ws, 'A3:I3', FILL_G)
    for col, txt in COLS.items():
        put(ws, f'{col}3', txt, size=9, h='center', fill=FILL_G, color=C_GRAYTXT, shrink=True)
    ws.row_dimensions[3].height = 22
    merge(ws, 'A4:I4')
    put(ws, 'A4', '空気調和工事', bold=True)
    ws.row_dimensions[4].height = 18
    r = 5
    rows = {}          # key -> row
    sec_info = {}
    all_lines = []
    for sec in LN.SECTIONS:
        merge(ws, f'A{r}:I{r}')
        put(ws, f'A{r}', f"{sec['no']}　{sec['title']}", bold=True)
        ws.row_dimensions[r].height = 20
        r += 1
        first = r
        exp_rows = []
        for ln in sec['lines']:
            k = ln['kind']
            merge(ws, f'A{r}:B{r}')
            merge(ws, f'C{r}:D{r}')
            ws.row_dimensions[r].height = 18
            if k == 'head':
                put(ws, f'A{r}', '　' + ln['name'], size=10, bold=True)
                r += 1
                continue
            if k == 'exp':
                name, rate = ES.EST_EXP[sec['key']][ln['idx']]
                ln = dict(ln, key=f"{sec['key']}_exp{ln['idx']}", name=name, rate=rate)
            ln['_row'] = r
            ln['_sec'] = sec['key']
            rows[ln['key']] = r
            all_lines.append(ln)
            put(ws, f'A{r}', '　　' + ln['name'], size=10, shrink=True)
            # 共通: エスト比較列
            est = ln.get('est')
            if est:
                put(ws, f'W{r}', est[0], size=10, h='right', fmt='#,##0.##')
                put(ws, f'X{r}', est[1], size=10, h='right', fmt=FMT_AMT)
                put(ws, f'Y{r}', est[2], size=10, h='right', fmt=FMT_SIGNED)
                put(ws, f'Z{r}', f'=H{r}-Y{r}', size=10, h='right', fmt=FMT_SIGNED)
            put(ws, f'AA{r}', ln.get('master') or None, size=9)
            put(ws, f'AB{r}', ln.get('basis') or None, size=9, h='center', shrink=True,
                fill=(FILL_PINK if ln.get('basis') in ('推測', '実測+推測') else None))
            if k == 'item':
                spec = ln['spec']
                spec_lines = est_lines(spec, 28 - 1, 9) if spec else 1
                if spec_lines > 1:
                    put(ws, f'C{r}', spec, size=9, wrap=True, indent=1)
                    ws.row_dimensions[r].height = max(18, row_h(spec_lines, 9, pad=4))
                else:
                    put(ws, f'C{r}', spec or None, size=10, shrink=True, indent=1)
                put(ws, f'E{r}', ln['qty'], size=10, h='right', fmt='#,##0_ ')
                put(ws, f'F{r}', ln['unit'], size=10, h='center')
                put(ws, f'G{r}', ln['price'], size=10, h='right', fmt=FMT_AMT)
                put(ws, f'H{r}', f'=ROUND(E{r}*G{r},0)', size=10, h='right', fmt=FMT_SIGNED)
                pub = ln.get('pub') or {'実測+推測': '図面実測＋想定', '推測': '想定数量'}.get(ln.get('basis'), '')
                put(ws, f'I{r}', pub or None, size=10, h='center', shrink=True)
                put(ws, f'J{r}', ln['src'], size=9, h='center', shrink=True)
                alerts = []
                if ln['src'] == LN.SRC_OWN:
                    alerts.append('単価要確認')
                if est and est[0] != ln['qty']:
                    alerts.append('エストと数量差')
                put(ws, f'K{r}', '・'.join(alerts) or None, size=9, h='center', shrink=True,
                    fill=(FILL_AMB if alerts else None))
                mat = ln['mat']
                if isinstance(mat, tuple):
                    put(ws, f'L{r}', f'=ROUND(G{r}*MAT_RATIO*{mat[1]},0)', size=10, h='right', fmt=FMT_AMT)
                else:
                    put(ws, f'L{r}', mat, size=10, h='right', fmt=FMT_AMT, fill=FILL_Y)
                put(ws, f'M{r}', ln['md'], size=10, h='right', fmt='0.000', fill=FILL_Y)
                put(ws, f'N{r}', f'=ROUND(M{r}*LABOR_COST,0)', size=10, h='right', fmt=FMT_AMT)
                put(ws, f'O{r}', f'=L{r}+N{r}', size=10, h='right', fmt=FMT_AMT)
                put(ws, f'P{r}', f'=E{r}*O{r}', size=10, h='right', fmt=FMT_SIGNED)
                put(ws, f'Q{r}', f'=E{r}*N{r}', size=10, h='right', fmt=FMT_SIGNED)
                put(ws, f'R{r}', f'=IF(H{r}=0,0,IF(E{r}*M{r}*EST_LABOR_PIPE>H{r},H{r},E{r}*M{r}*EST_LABOR_PIPE))',
                    size=10, h='right', fmt=FMT_SIGNED)
                if ln.get('bug'):
                    put(ws, f'U{r}', ln['bug'], size=10, h='right', fmt='0.000', fill=FILL_Y)
                memo = ln.get('memo') or ''
                if isinstance(mat, tuple):
                    memo = f"材料原価＝エスト単価×材料原価率×管長割増{mat[1]}（建設物価相当）" + ('｜' + memo if memo else '')
                put(ws, f'AC{r}', memo or None, size=9)
            elif k == 'tbd':
                put(ws, f'C{r}', ln['spec'], size=9, wrap=True, indent=1, color=C_RED)
                ws.row_dimensions[r].height = max(18, row_h(est_lines(ln['spec'], 27, 9), 9, pad=4))
                put(ws, f'A{r}', '　　' + ln['name'], size=10, shrink=True, color=C_RED)
                put(ws, f'E{r}', ln['qty'], size=10, h='right', fmt='#,##0_ ')
                put(ws, f'F{r}', ln['unit'], size=10, h='center')
                put(ws, f'G{r}', f'=IF(INCLUDE_TBD=1,ROUNDUP(V{r}/(1-TARGET_MARGIN),-3),"")', size=10, h='right', fmt=FMT_AMT)
                put(ws, f'H{r}', f'=IF(INCLUDE_TBD=1,ROUND(E{r}*G{r},0),0)', size=10, h='right', fmt=FMT_TBD, color=C_RED)
                put(ws, f'I{r}', '=IF(INCLUDE_TBD=1,"概算（要見積）","別途見積")', size=10, h='center', shrink=True, color=C_RED)
                put(ws, f'J{r}', LN.SRC_TBD, size=9, h='center', fill=FILL_RED, bold=True)
                put(ws, f'K{r}', '要見積', size=9, h='center', fill=FILL_RED, bold=True, color=C_RED)
                put(ws, f'V{r}', ln['prov'], size=10, h='right', fmt=FMT_AMT, fill=FILL_Y)
                put(ws, f'L{r}', f'=V{r}', size=10, h='right', fmt=FMT_AMT)
                put(ws, f'M{r}', 0, size=10, h='right', fmt='0.000')
                put(ws, f'N{r}', 0, size=10, h='right', fmt=FMT_AMT)
                put(ws, f'O{r}', f'=L{r}+N{r}', size=10, h='right', fmt=FMT_AMT)
                put(ws, f'P{r}', f'=IF(INCLUDE_TBD=1,E{r}*O{r},0)', size=10, h='right', fmt=FMT_SIGNED)
                put(ws, f'Q{r}', 0, size=10, h='right', fmt=FMT_SIGNED)
                put(ws, f'R{r}', 0, size=10, h='right', fmt=FMT_SIGNED)
                put(ws, f'AC{r}', ln.get('memo') or None, size=9)
            elif k == 'pct':
                bases = [rows[b] for b in ln['base']]
                put(ws, f'E{r}', 1, size=10, h='right', fmt='#,##0_ ')
                put(ws, f'F{r}', '式', size=10, h='center')
                put(ws, f'H{r}', '=ROUND((' + '+'.join(f'H{b}' for b in bases) + f')*U{r},0)', size=10, h='right', fmt=FMT_SIGNED)
                put(ws, f'J{r}', LN.SRC_CALC, size=9, h='center', shrink=True)
                put(ws, f'U{r}', ln['rate'], size=10, h='right', fmt='0%', fill=FILL_Y)
                put(ws, f'P{r}', '=ROUND((' + '+'.join(f'P{b}' for b in bases) + f')*U{r},0)', size=10, h='right', fmt=FMT_SIGNED)
                put(ws, f'Q{r}', 0, size=10, h='right', fmt=FMT_SIGNED)
                put(ws, f'R{r}', 0, size=10, h='right', fmt=FMT_SIGNED)
                put(ws, f'AC{r}', '管金額×率（エスト・公共積算と同率）。原価も同率', size=9)
            elif k == 'labor':
                bases = [rows[b] for b in ln['base']]
                put(ws, f'E{r}', 1, size=10, h='right', fmt='#,##0_ ')
                put(ws, f'F{r}', '式', size=10, h='center')
                put(ws, f'M{r}', '=' + '+'.join(f'E{b}*U{b}' for b in bases), size=10, h='right', fmt='0.000')
                put(ws, f'H{r}', f'=ROUNDUP(M{r}*{ln["rate"]},-1)', size=10, h='right', fmt=FMT_SIGNED)
                put(ws, f'J{r}', LN.SRC_CALC, size=9, h='center', shrink=True)
                put(ws, f'L{r}', 0, size=10, h='right', fmt=FMT_AMT)
                put(ws, f'N{r}', f'=ROUND(M{r}*LABOR_COST,0)', size=10, h='right', fmt=FMT_AMT)
                put(ws, f'O{r}', f'=L{r}+N{r}', size=10, h='right', fmt=FMT_AMT)
                put(ws, f'P{r}', f'=E{r}*O{r}', size=10, h='right', fmt=FMT_SIGNED)
                put(ws, f'Q{r}', f'=E{r}*N{r}', size=10, h='right', fmt=FMT_SIGNED)
                put(ws, f'R{r}', f'=H{r}', size=10, h='right', fmt=FMT_SIGNED)
                put(ws, f'AC{r}', ln.get('memo') or None, size=9)
            elif k == 'exp':
                put(ws, f'E{r}', 1, size=10, h='right', fmt='#,##0_ ')
                put(ws, f'F{r}', '式', size=10, h='center')
                put(ws, f'J{r}', '経費', size=9, h='center')
                put(ws, f'U{r}', ln['rate'], size=10, h='right', fmt='0%', fill=FILL_Y)
                put(ws, f'P{r}', 0, size=10, h='right', fmt=FMT_SIGNED)
                put(ws, f'Q{r}', 0, size=10, h='right', fmt=FMT_SIGNED)
                put(ws, f'R{r}', 0, size=10, h='right', fmt=FMT_SIGNED)
                exp_rows.append(r)
            if k != 'exp':
                put(ws, f'S{r}', f'=IF(H{r}="","",H{r}-P{r})', size=10, h='right', fmt=FMT_SIGNED)
                put(ws, f'T{r}', f'=IF(OR(H{r}="",H{r}=0),"",S{r}/H{r})', size=10, h='right', fmt='0.0%')
            r += 1
        # 部門内経費（エスト規則）
        e1, e2, e3 = exp_rows
        base = f'SUM(H{first}:H{e1 - 1})'
        put(ws, f'H{e1}', f'=ROUNDUP({base}*U{e1},-2)', size=10, h='right', fmt=FMT_SIGNED)
        put(ws, f'H{e2}', f'=ROUNDUP(({base}+H{e1})*U{e2},-2)', size=10, h='right', fmt=FMT_SIGNED)
        s3 = f'({base}+H{e1}+H{e2})'
        put(ws, f'H{e3}', f'=ROUNDUP({s3}+ROUNDUP({s3}*U{e3},-2),-3)-{s3}', size=10, h='right', fmt=FMT_SIGNED)
        put(ws, f'AC{e1}', '基礎額×率（100円未満切上げ）', size=9)
        put(ws, f'AC{e2}', '（基礎額＋消耗品）×率（100円未満切上げ）', size=9)
        put(ws, f'AC{e3}', '（基礎額＋消耗品＋運搬費）×率を加え、計を1,000円未満切上げ（端数はここに吸収）', size=9)
        # 計
        tot = r
        merge(ws, f'A{tot}:D{tot}')
        put(ws, f'A{tot}', '－　計　－', bold=True, h='center')
        put(ws, f'H{tot}', f'=SUM(H{first}:H{e3})', bold=True, h='right', fmt=FMT_SIGNED)
        put(ws, f'M{tot}', f'=SUMPRODUCT(E{first}:E{e3},M{first}:M{e3})', bold=True, h='right', fmt='0.00')
        for col in 'PQR':
            put(ws, f'{col}{tot}', f'=SUM({col}{first}:{col}{e3})', bold=True, h='right', fmt=FMT_SIGNED)
        put(ws, f'S{tot}', f'=H{tot}-P{tot}', bold=True, h='right', fmt=FMT_SIGNED)
        put(ws, f'T{tot}', f'=IF(H{tot}=0,"",S{tot}/H{tot})', bold=True, h='right', fmt='0.0%')
        est_tot = [a[6] for a in ES.EST_A if a[0] == sec['title'] and a[1] == '－ 計 －'][0]
        put(ws, f'Y{tot}', est_tot, bold=True, h='right', fmt=FMT_SIGNED)
        put(ws, f'Z{tot}', f'=H{tot}-Y{tot}', bold=True, h='right', fmt=FMT_SIGNED)
        ws.row_dimensions[tot].height = 20
        hline(ws, f'A{tot}:I{tot}', 'top', 'thin')
        sec_info[sec['key']] = dict(first=first, exp=exp_rows, total=tot)
        r = tot + 1
        merge(ws, f'A{r}:B{r}')
        merge(ws, f'C{r}:D{r}')
        ws.row_dimensions[r].height = 10
        if sec['key'] == 'AC':
            ws.row_breaks.append(Break(id=r))   # 空調設備工事と配管設備工事でページを分ける
        r += 1
    # 工事諸経費
    oh = r
    merge(ws, f'A{oh}:B{oh}')
    merge(ws, f'C{oh}:D{oh}')
    put(ws, f'A{oh}', '８　工事諸経費', bold=True)
    put(ws, f'E{oh}', 1, size=10, h='right', fmt='#,##0_ ')
    put(ws, f'F{oh}', '式', size=10, h='center')
    ac_t, pp_t = sec_info['AC']['total'], sec_info['PIPE']['total']
    put(ws, f'H{oh}', f'=ROUNDUP((H{ac_t}+H{pp_t})*U{oh},-2)', bold=True, h='right', fmt=FMT_SIGNED)
    put(ws, f'J{oh}', '経費', size=9, h='center')
    put(ws, f'U{oh}', ES.EST_OVERHEAD, size=10, h='right', fmt='0%', fill=FILL_Y)
    put(ws, f'P{oh}', 0, size=10, h='right', fmt=FMT_SIGNED)
    put(ws, f'Y{oh}', 64300, size=10, h='right', fmt=FMT_SIGNED)
    put(ws, f'Z{oh}', f'=H{oh}-Y{oh}', size=10, h='right', fmt=FMT_SIGNED)
    put(ws, f'AC{oh}', '（空調設備工事＋配管設備工事）×5%（100円未満切上げ）', size=9)
    ws.row_dimensions[oh].height = 20
    r = oh + 1
    merge(ws, f'A{r}:B{r}')
    merge(ws, f'C{r}:D{r}')
    ws.row_dimensions[r].height = 10
    gt = r + 1
    merge(ws, f'A{gt}:D{gt}')
    put(ws, f'A{gt}', '【小　　計】', bold=True, h='center')
    put(ws, f'H{gt}', f'=H{ac_t}+H{pp_t}+H{oh}', bold=True, h='right', fmt=FMT_SIGNED)
    put(ws, f'I{gt}', '出精値引は表紙', size=9, h='center', shrink=True)
    put(ws, f'M{gt}', f'=M{ac_t}+M{pp_t}', bold=True, h='right', fmt='0.00')
    for col in 'PQR':
        put(ws, f'{col}{gt}', f'={col}{ac_t}+{col}{pp_t}', bold=True, h='right', fmt=FMT_SIGNED)
    put(ws, f'S{gt}', f'=H{gt}-P{gt}', bold=True, h='right', fmt=FMT_SIGNED)
    put(ws, f'T{gt}', f'=IF(H{gt}=0,"",S{gt}/H{gt})', bold=True, h='right', fmt='0.0%')
    put(ws, f'Y{gt}', ES.EST_SUBTOTAL, bold=True, h='right', fmt=FMT_SIGNED)
    put(ws, f'Z{gt}', f'=H{gt}-Y{gt}', bold=True, h='right', fmt=FMT_SIGNED)
    ws.row_dimensions[gt].height = 22
    grid(ws, f'A3:I{gt}', outer='medium', vert='thin', horiz='hair')
    grid(ws, f'J3:AC{gt}', outer='thin', vert='thin', horiz='hair')
    hline(ws, 'A3:I3', 'bottom', 'thin')
    hline(ws, 'J3:AC3', 'bottom', 'thin')
    for key in ('AC', 'PIPE'):
        t = sec_info[key]['total']
        hline(ws, f'A{t}:I{t}', 'top', 'thin')
        hline(ws, f'J{t}:AC{t}', 'top', 'thin')
    hline(ws, f'A{gt}:I{gt}', 'top', 'medium')
    hline(ws, f'J{gt}:AC{gt}', 'top', 'medium')
    for c in range(1, 10):
        set_border(ws.cell(gt, c), bottom='medium')
    set_border(ws.cell(gt, 1), left='medium')
    set_border(ws.cell(gt, 9), right='medium')
    ws.freeze_panes = 'A4'
    page_setup(ws, f'A1:I{gt}', fit_h=0, title_rows='1:3', footer='&P / &N')
    return dict(rows=rows, sec=sec_info, oh=oh, gt=gt, lines=all_lines)


# =====================================================================
# 法定福利費内訳明細書（工事費に内含）
# =====================================================================
def build_welfare(ws, det):
    set_widths(ws, dict(A=2, B=25, C=9, D=9, E=14, F=16, G=31))
    merge(ws, 'B1:G1')
    put(ws, 'B1', '法 定 福 利 費 内 訳 明 細 書', size=14, bold=True, h='center')
    ws.row_dimensions[1].height = 26
    merge(ws, 'B2:G2')
    put(ws, 'B2', f'="（御見積書 No."&{q(SH_COVER)}!K3&" の添付資料）"', size=10, h='center')
    put(ws, 'B4', '見積年月日', fill=FILL_L, indent=1)
    merge(ws, 'C4:E4')
    put(ws, 'C4', f"={q(SH_COVER)}!J4", h='left', fmt=DATE_FMT, indent=1)
    put(ws, 'F4', '見積No.', fill=FILL_L, indent=1)
    put(ws, 'G4', f'=IF({q(SH_COVER)}!K3="","",{q(SH_COVER)}!K3)', h='left', indent=1)
    for r, lab, val in [(5, '発注者名', CLIENT + '　御中'), (6, '工事名', PROJECT_FULL), (7, '施工場所', SITE)]:
        put(ws, f'B{r}', lab, fill=FILL_L, indent=1)
        merge(ws, f'C{r}:G{r}')
        put(ws, f'C{r}', val, shrink=True, indent=1)
    put(ws, 'B8', '会社名', fill=FILL_L, indent=1)
    merge(ws, 'C8:E8')
    put(ws, 'C8', COMPANY, indent=1)
    put(ws, 'F8', '所在地', fill=FILL_L, indent=1)
    put(ws, 'G8', ADDR, indent=1)
    grid(ws, 'B4:G8', outer='thin', vert='thin', horiz='thin')
    for r in range(4, 9):
        ws.row_dimensions[r].height = 20
    ws.row_dimensions[3].height = 8
    ws.row_dimensions[9].height = 10
    put(ws, 'B10', '【1】　労務費相当額（Ｙ）の内訳', bold=True)
    note1 = ('※ 労務費相当額は、内訳書各項目の金額のうち労務費に相当する額（配管工費・電線材料施工費、および据付・保温・試験等の人工相当額）'
             'に、出精値引き後の比率を乗じたものです。空調機は支給品のため含みません。')
    merge(ws, 'B11:G11')
    put(ws, 'B11', note1, size=10, wrap=True)
    ws.row_dimensions[11].height = row_h(est_lines(note1, 104, 10), 10)
    put(ws, 'B12', '項　　目', h='center', fill=FILL_G)
    merge(ws, 'C12:D12')
    put(ws, 'C12', '労務費相当額（円）', size=10, h='center', fill=FILL_G)
    merge(ws, 'E12:G12')
    put(ws, 'E12', '摘　　要', h='center', fill=FILL_G)
    fill_range(ws, 'B12:G12', FILL_G)
    ac_t, pp_t = det['sec']['AC']['total'], det['sec']['PIPE']['total']
    rows1 = [
        (13, '空調設備工事', f"={q(SH_DETAIL)}!R{ac_t}", '内訳書 ２ 空調設備工事（定価ベース）'),
        (14, '配管設備工事', f"={q(SH_DETAIL)}!R{pp_t}", '内訳書 ３ 配管設備工事（定価ベース）'),
        (15, '小　　計', '=C13+C14', '定価ベース'),
        (16, '労務費相当額（Ｙ）', f"=ROUNDDOWN(C15*{q(SH_COVER)}!K28/{q(SH_COVER)}!K26,0)", '小計 × 出精値引き後比率（合計÷小計）'),
    ]
    for r, lab, f, memo in rows1:
        bold = r == 16
        put(ws, f'B{r}', lab, bold=bold, indent=1)
        merge(ws, f'C{r}:D{r}')
        put(ws, f'C{r}', f, bold=bold, h='right', fmt=FMT_AMT)
        merge(ws, f'E{r}:G{r}')
        put(ws, f'E{r}', memo, size=10, indent=1)
        ws.row_dimensions[r].height = 20
    grid(ws, 'B12:G16', outer='medium', vert='thin', horiz='thin')
    hline(ws, 'B16:G16', 'top', 'double')
    r = 18
    put(ws, f'B{r}', '【2】　法定福利費相当額（Ａ）の算定', bold=True)
    r += 1
    merge(ws, f'B{r}:G{r}')
    put(ws, f'B{r}', '計算式：　Ａ ＝ Ｙ × Ｚ（Ｚ：法定福利費事業者負担率、保険ごとに円未満切捨て）', size=10)
    r += 1
    hdr2 = r
    for ref, txt in [('B', '保険の種類'), ('E', '事業主負担料率'), ('F', '法定福利費（円）'), ('G', '算定根拠')]:
        put(ws, f'{ref}{r}', txt, size=10, h='center', fill=FILL_G)
    merge(ws, f'C{r}:D{r}')
    put(ws, f'C{r}', '労務費相当額（円）', size=10, h='center', fill=FILL_G)
    fill_range(ws, f'B{r}:G{r}', FILL_G)
    ws.row_dimensions[r].height = 22
    r += 1
    for name, rate, basis in WELFARE_RATES:
        put(ws, f'B{r}', name, indent=1)
        merge(ws, f'C{r}:D{r}')
        put(ws, f'C{r}', '=$C$16', h='right', fmt=FMT_AMT)
        put(ws, f'E{r}', rate, h='right', fmt='0.000%_ ')
        put(ws, f'F{r}', f'=ROUNDDOWN(C{r}*E{r},0)', h='right', fmt=FMT_AMT)
        put(ws, f'G{r}', basis, size=9, wrap=True)
        ws.row_dimensions[r].height = max(22, row_h(est_lines(basis, 31, 9), 9, pad=5))
        r += 1
    wf_first, wf_last = hdr2 + 1, r - 1
    tot = r
    put(ws, f'B{tot}', '合計（Ｚ ／ Ａ）', bold=True, h='center')
    merge(ws, f'C{tot}:D{tot}')
    put(ws, f'E{tot}', f'=SUM(E{wf_first}:E{wf_last})', bold=True, h='right', fmt='0.000%_ ')
    put(ws, f'F{tot}', f'=SUM(F{wf_first}:F{wf_last})', bold=True, h='right', fmt=FMT_AMT)
    ws.row_dimensions[tot].height = 24
    grid(ws, f'B{hdr2}:G{tot}', outer='medium', vert='thin', horiz='thin')
    hline(ws, f'B{tot}:G{tot}', 'top', 'double')
    r = tot + 2
    for t in ['※ 法定福利費相当額（Ａ）は見積金額（工事費）に含まれています。別途加算するものではありません。',
              '※ 料率は見積作成時点（令和8年度）の公表料率によります。料率が改定された場合は改定後の料率により精算させていただきます。']:
        merge(ws, f'B{r}:G{r}')
        put(ws, f'B{r}', t, size=10, wrap=True)
        ws.row_dimensions[r].height = row_h(est_lines(t, 104, 10), 10)
        r += 1
    page_setup(ws, f'B1:G{r - 1}', fit_h=1)
    return dict(rate_cell=f'$E${tot}', amount_cell=f'$F${tot}', y_cell='$C$16')


# =====================================================================
# 表紙(工事)
# =====================================================================
def build_cover(ws, det, wf):
    FIRST = 22
    SUB, DISC, TOTAL = 26, 27, 28
    set_widths(ws, dict(A=6, B=8, C=3, D=15, E=9, F=9, G=7, H=5, I=7, J=7, K=14, L=12,
                        M=44, N=13, O=13, P=13, Q=10))
    merge(ws, 'A1:L2')
    put(ws, 'A1', '御　見　積　書', size=20, bold=True, h='center')
    ws.row_dimensions[1].height = 22
    ws.row_dimensions[2].height = 22
    put(ws, 'J3', 'No.', h='right')
    merge(ws, 'K3:L3')
    put(ws, 'K3', None, h='center')
    hline(ws, 'K3:L3', 'bottom', 'thin')
    merge(ws, 'J4:L4')
    put(ws, 'J4', EST_DATE, h='right', fmt=DATE_FMT)
    merge(ws, 'A5:F5')
    put(ws, 'A5', CLIENT, size=14, h='center')
    hline(ws, 'A5:F5', 'bottom', 'thin')
    put(ws, 'G5', '御中', size=12, h='left')
    ws.row_dimensions[5].height = 26
    ws.row_dimensions[6].height = 12
    put(ws, 'A7', '金額', size=12, h='center')
    merge(ws, 'B7:E7')
    put(ws, 'B7', f'="¥"&TEXT(K{TOTAL},"#,##0")&"-"', size=14, bold=True, h='center')
    hline(ws, 'B7:E7', 'bottom', 'double')
    put(ws, 'F7', '（税別）', h='left')
    put(ws, 'H7', COMPANY, size=12, bold=True)
    ws.row_dimensions[7].height = 28
    ws.row_dimensions[8].height = 10
    merge(ws, 'A9:F9')
    put(ws, 'A9', 'つぎのとおりお見積いたしました。', h='left')
    put(ws, 'H9', REP)
    put(ws, 'A10', 'なにとぞご用命くださいますようお願い申しあげます。')
    put(ws, 'H11', ZIP)
    merge(ws, 'A12:B12')
    put(ws, 'A12', '件名', h='distributed', indent=1)
    merge(ws, 'D12:G12')
    put(ws, 'D12', PROJECT_1, shrink=True)
    put(ws, 'H12', ADDR)
    merge(ws, 'D13:G13')
    put(ws, 'D13', PROJECT_2, shrink=True)
    put(ws, 'H13', TEL)
    put(ws, 'K13', FAX)
    merge(ws, 'A14:B14')
    put(ws, 'A14', '施工場所', h='distributed', indent=1)
    merge(ws, 'D14:G14')
    put(ws, 'D14', SITE, size=10, shrink=True)
    put(ws, 'H14', '担当者')
    put(ws, 'J14', STAFF)
    for r, lab, val in [(15, '工事期限', TERM), (16, '支払条件', PAYMENT), (17, '本書有効期間', VALIDITY)]:
        merge(ws, f'A{r}:B{r}')
        put(ws, f'A{r}', lab, h='distributed', indent=1, shrink=True)
        merge(ws, f'D{r}:L{r}')
        put(ws, f'D{r}', val, shrink=True)
    for r in range(9, 18):
        ws.row_dimensions[r].height = 19
    ws.row_dimensions[18].height = 12
    merge(ws, 'B19:F19')
    merge(ws, 'I19:J19')
    for ref, txt, sz in [('A19', '項', 10), ('B19', '品　　　名', 11), ('G19', '数 量', 11), ('H19', '単位', 10),
                         ('I19', '単  価', 11), ('K19', '金　　額', 11), ('L19', '備　考', 11)]:
        put(ws, ref, txt, size=sz, h='center', fill=FILL_G, shrink=True)
    fill_range(ws, 'A19:L19', FILL_G)
    merge(ws, 'B20:L20')
    put(ws, 'B20', '【' + PROJECT_FULL + '】　空気調和工事', indent=1, shrink=True)
    merge(ws, 'B21:L21')
    put(ws, 'B21', '=IF(INCLUDE_TBD=1,"［空調機・リモコン・吹出しパンカーは支給品。鉄骨架台は概算計上（詳細見積後に精算）］",'
        '"［空調機・リモコン・吹出しパンカーは支給品。鉄骨架台は別途御見積］")', size=10, indent=1, shrink=True)
    for r in range(22, TOTAL + 1):
        merge(ws, f'B{r}:F{r}')
        merge(ws, f'I{r}:J{r}')
    ac_t, pp_t, oh = det['sec']['AC']['total'], det['sec']['PIPE']['total'], det['oh']
    rows = [
        (22, CIRC[0], '空調設備機器', 0, '支給品', FMT_SUPPLY),
        (23, CIRC[1], LN.SECTIONS[0]['cover'], f"={q(SH_DETAIL)}!H{ac_t}", '内訳書 ２', FMT_SIGNED),
        (24, CIRC[2], LN.SECTIONS[1]['cover'], f"={q(SH_DETAIL)}!H{pp_t}", '内訳書 ３', FMT_SIGNED),
        (25, CIRC[3], '工事諸経費', f"={q(SH_DETAIL)}!H{oh}", '内訳書 ８', FMT_SIGNED),
    ]
    for r, no, name, fk, remark, fmt in rows:
        put(ws, f'A{r}', no, h='center')
        put(ws, f'B{r}', name, indent=1, shrink=True)
        put(ws, f'G{r}', 1, h='right', fmt=FMT_QTY)
        put(ws, f'H{r}', '式', h='center')
        put(ws, f'I{r}', f'=K{r}', h='right', fmt=fmt)
        put(ws, f'K{r}', fk, h='right', fmt=fmt)
        put(ws, f'L{r}', remark, size=10, h='center', shrink=True)
    put(ws, f'B{SUB}', '　小　　計', indent=1)
    put(ws, f'K{SUB}', f'=SUM(K{FIRST}:K{SUB - 1})', h='right', fmt=FMT_SIGNED)
    put(ws, f'A{DISC}', CIRC[4], h='center')
    put(ws, f'B{DISC}', '出精値引', indent=1)
    put(ws, f'K{DISC}', f'=IF(SUBMIT<K{SUB},SUBMIT-K{SUB},0)', h='right', fmt=FMT_SIGNED)
    put(ws, f'B{TOTAL}', '　合　　計', bold=True, indent=1)
    put(ws, f'K{TOTAL}', f'=K{SUB}+K{DISC}', bold=True, h='right', fmt=FMT_AMT)
    put(ws, f'L{TOTAL}', '（税別）', size=10, h='center')
    for r in range(19, TOTAL + 1):
        ws.row_dimensions[r].height = 24
    ws.row_dimensions[20].height = 22
    ws.row_dimensions[21].height = 20
    grid(ws, f'A19:L{TOTAL}', outer='medium', vert='thin', horiz='thin')
    hline(ws, f'A{SUB}:L{SUB}', 'top', 'medium')
    hline(ws, f'A{TOTAL}:L{TOTAL}', 'top', 'double')
    for c in range(1, 13):
        set_border(ws.cell(TOTAL, c), bottom='medium')
    # 法定福利費相当額（内含）
    r = TOTAL + 1
    ws.row_dimensions[r].height = 8
    r += 1
    merge(ws, f'A{r}:L{r}')
    put(ws, f'A{r}', '上記の工事費には、以下の法定福利費相当額を含んでいます。', size=10, indent=1)
    r += 1
    merge(ws, f'A{r}:L{r}')
    put(ws, f'A{r}', f'="【法定福利費相当額（Ａ）】"&TEXT({q(SH_WF)}!{wf["amount_cell"]},"#,##0")&"円　　Ａ＝Ｙ×Ｚ　Ｙ：労務費相当額 "'
        f'&TEXT({q(SH_WF)}!{wf["y_cell"]},"#,##0")&"円　Ｚ：法定福利費事業者負担率（合計値："&TEXT({q(SH_WF)}!{wf["rate_cell"]}*100,"0.000")&"／100）"',
        size=10, indent=1, shrink=True)
    grid(ws, f'A{r - 1}:L{r}', outer='thin', vert=None, horiz=None)
    r += 2
    for t in NOTES['cover']:
        put(ws, f'A{r}', '※', size=9.5, h='center', v='top')
        merge(ws, f'B{r}:L{r}')
        put(ws, f'B{r}', t, size=9.5, wrap=True, v='top')
        ws.row_dimensions[r].height = row_h(est_lines(t, 96, 9.5), 9.5, pad=3)
        r += 1
    last_note = r - 1

    # ---- 内部エリア（M〜Q、印刷範囲外）
    put(ws, 'M3', '【設定・内部用】（印刷範囲外）黄セル＝入力', bold=True, color=C_GRAYTXT, fill=FILL_G)
    put(ws, 'N3', None, fill=FILL_G)
    params = [
        (4, '目標粗利率', PV.TARGET_MARGIN, '0%', True),
        (5, '提出額の丸め単位（切上げ）', PV.ROUND_UNIT, '#,##0', True),
        (6, '労務原価日額（円/人工、当社）', PV.LABOR_COST, '#,##0', True),
        (7, '材料原価率（エスト材料単価＝建設物価相当に対する比率）', PV.MAT_RATIO, '0.00', True),
        (8, 'エスト労務単価 配管工費（円/人）', PV.EST_LABOR_PIPE, '#,##0', True),
        (9, 'エスト労務単価 電線材料施工費（円/人）', PV.EST_LABOR_ELEC, '#,##0', True),
        (10, '法定福利費率（事業主負担）', '=WELFARE_RATE', '0.000%', False),
        (11, '要見積（概算）を提出額に含める（1=含める／0=別途見積）', PV.INCLUDE_TBD, '0', True),
        (12, '目標提出額（税別・空欄なら粗利率から自動）', None, '#,##0', True),
        (13, '提出額（自動計算）', '=IF(TARGET="",ROUNDUP(COST_TOTAL/(1-TARGET_MARGIN)/ROUND_UNIT,0)*ROUND_UNIT,TARGET)', '#,##0', False),
    ]
    for rr, lab, val, fmt, is_input in params:
        put(ws, f'M{rr}', lab, size=10, color=C_GRAYTXT, fill=FILL_G, shrink=True)
        put(ws, f'N{rr}', val, h='right', fmt=fmt, fill=(FILL_Y if is_input else None), bold=(rr == 13))
    grid(ws, 'M3:N13', outer='thin', vert='thin', horiz='thin')
    put(ws, 'M14', '※ 提出額＝原価合計÷（1－目標粗利率）を丸め単位で切上げ。小計（エスト単価の定価ベース）との差を出精値引に計上', size=9, color=C_GRAYTXT)
    put(ws, 'M15', '※ 原価＝材料（エスト材料単価×材料原価率＝建設物価相当）＋人工×労務原価日額＋法定福利費（事業主負担）', size=9, color=C_GRAYTXT)
    for ref, txt in [('M19', '内部集計'), ('N19', '定価ベース'), ('O19', '原価'), ('P19', '粗利'), ('Q19', '粗利率')]:
        put(ws, ref, txt, size=10, h='center', color=C_GRAYTXT, fill=FILL_G, shrink=True)
    gt = det['gt']
    summ = [
        (23, '②空調設備工事', f'=K23', f"={q(SH_DETAIL)}!P{ac_t}"),
        (24, '③配管設備工事', f'=K24', f"={q(SH_DETAIL)}!P{pp_t}"),
        (25, '④工事諸経費（原価なし）', f'=K25', 0),
        (26, '法定福利費（事業主負担・原価）', None, f"=ROUND({q(SH_DETAIL)}!Q{gt}*WELFARE_RATE,0)"),
        (27, '原価合計', f'=K{SUB}', '=SUM(O23:O26)'),
    ]
    for rr, lab, nv, ov in summ:
        put(ws, f'M{rr}', lab, size=10, bold=(rr == 27), shrink=True)
        put(ws, f'N{rr}', nv, h='right', fmt=FMT_SIGNED)
        put(ws, f'O{rr}', ov, h='right', fmt=FMT_SIGNED, bold=(rr == 27))
        if nv is not None and rr != 27:
            put(ws, f'P{rr}', f'=N{rr}-O{rr}', h='right', fmt=FMT_SIGNED)
            put(ws, f'Q{rr}', f'=IF(N{rr}=0,"",P{rr}/N{rr})', h='right', fmt='0.0%')
    put(ws, f'M{TOTAL}', '提出額（合計）／粗利／粗利率', size=10, bold=True, shrink=True)
    put(ws, f'N{TOTAL}', f'=K{TOTAL}', bold=True, h='right', fmt=FMT_AMT)
    put(ws, f'O{TOTAL}', '=O27', h='right', fmt=FMT_AMT)
    put(ws, f'P{TOTAL}', f'=N{TOTAL}-O{TOTAL}', bold=True, h='right', fmt=FMT_SIGNED)
    put(ws, f'Q{TOTAL}', f'=IF(N{TOTAL}=0,"",P{TOTAL}/N{TOTAL})', bold=True, h='right', fmt='0.0%')
    grid(ws, f'M19:Q{TOTAL}', outer='thin', vert='thin', horiz='hair')
    # アラート
    rows_tbd = [ln['_row'] for ln in det['lines'] if ln['kind'] == 'tbd']
    rows_own = [ln for ln in det['lines'] if ln['kind'] == 'item' and ln['src'] == LN.SRC_OWN]
    rows_qd = [ln for ln in det['lines'] if ln.get('est') and ln['kind'] == 'item' and ln['est'][0] != ln['qty']]
    A0 = 31
    put(ws, f'M{A0}', '⚠ アラート（提出前に確認）', bold=True, color=C_RED, fill=FILL_RED)
    for c in 'NOPQ':
        put(ws, f'{c}{A0}', None, fill=FILL_RED)
    tbd_cost = '+'.join(f"{q(SH_DETAIL)}!E{r}*{q(SH_DETAIL)}!V{r}" for r in rows_tbd)
    alerts = [
        (A0 + 1, f'要見積 {len(rows_tbd)}件：架台制作及び取付費（概算原価の合計）', f'={tbd_cost}',
         '=IF(INCLUDE_TBD=1,"概算を提出額に含めています","別途見積（提出額に含めていません）")'),
        (A0 + 2, '要見積を概算で含めた場合の提出額（目安）',
         f'=IF(INCLUDE_TBD=1,SUBMIT,ROUNDUP((COST_TOTAL+N{A0 + 1})/(1-TARGET_MARGIN)/ROUND_UNIT,0)*ROUND_UNIT)', 'N11 を 1 にすると反映'),
        (A0 + 3, f'単価要確認 {len(rows_own)}件：' + '・'.join(ln['name'] for ln in rows_own), None, 'エスト単価なし。当社単価で仮計上'),
        (A0 + 4, f'エストと数量が異なる行 {len(rows_qd)}件（内訳書 K列・エスト比較シート）', None,
         'SA室外機3台（エスト2台）ほか'),
        (A0 + 5, '定価ベース（エスト単価）で目標粗利に届くか', None,
         f'=IF(SUBMIT>K{SUB},"⚠ 値引きなしでも粗利率 "&TEXT((K{SUB}-COST_TOTAL)/K{SUB},"0.0%")&"（目標未達）","OK（出精値引で調整）")'),
    ]
    for rr, lab, nv, ov in alerts:
        put(ws, f'M{rr}', lab, size=10, shrink=True, color=C_RED)
        put(ws, f'N{rr}', nv, h='right', fmt=FMT_AMT)
        merge(ws, f'O{rr}:Q{rr}')
        put(ws, f'O{rr}', ov, size=10, shrink=True)
    grid(ws, f'M{A0}:Q{A0 + 5}', outer='thin', vert='thin', horiz='hair')
    # 参考
    R0 = A0 + 7
    put(ws, f'M{R0}', '参考（感度・エストとの比較）', bold=True, color=C_GRAYTXT, fill=FILL_G)
    for c in 'NOPQ':
        put(ws, f'{c}{R0}', None, fill=FILL_G)
    put(ws, f'N{R0}', '金額', size=10, h='center', color=C_GRAYTXT, fill=FILL_G)
    put(ws, f'O{R0}', '原価', size=10, h='center', color=C_GRAYTXT, fill=FILL_G)
    put(ws, f'Q{R0}', '粗利率', size=10, h='center', color=C_GRAYTXT, fill=FILL_G)
    labor_md = f"{q(SH_DETAIL)}!M{gt}"
    refs = [
        (R0 + 1, 'エスト見積 6330560-1 の提出額（当社数量・当社原価で見た粗利率）', ES.EST_TOTAL, '=COST_TOTAL',
         f'=(N{R0 + 1}-O{R0 + 1})/N{R0 + 1}'),
        (R0 + 2, '労務を公共工事設計労務単価（広島 配管工 24,700）で計算した場合', f'=K{TOTAL}',
         f'=COST_TOTAL+{labor_md}*(N{R0 + 2 + 4}-LABOR_COST)*(1+WELFARE_RATE)', f'=(N{R0 + 2}-O{R0 + 2})/N{R0 + 2}'),
        (R0 + 3, '出精値引率（変更増減の算定に使用：増減額×(1－この率)）', f'=IF(K{SUB}=0,"",-K{DISC}/K{SUB})', None, None),
        (R0 + 4, '人工合計（内訳書）', f'={labor_md}', None, None),
    ]
    for rr, lab, nv, ov, qv in refs:
        put(ws, f'M{rr}', lab, size=10, shrink=True)
        put(ws, f'N{rr}', nv, h='right', fmt=('0.0%' if rr == R0 + 3 else ('0.00' if rr == R0 + 4 else FMT_AMT)))
        put(ws, f'O{rr}', ov, h='right', fmt=FMT_AMT)
        put(ws, f'Q{rr}', qv, h='right', fmt='0.0%')
    put(ws, f'M{R0 + 6}', '公共工事設計労務単価 広島 配管工（R8.3）', size=10, color=C_GRAYTXT)
    put(ws, f'N{R0 + 6}', PV.PUBLIC_LABOR_R8_HIROSHIMA['配管工'], h='right', fmt=FMT_AMT, fill=FILL_Y)
    grid(ws, f'M{R0}:Q{R0 + 4}', outer='thin', vert='thin', horiz='hair')
    T0 = R0 + 8
    put(ws, f'M{T0}', '粗利率別 参考提出額（税別・万円未満切上げ）', size=10, color=C_GRAYTXT, fill=FILL_G, shrink=True)
    put(ws, f'N{T0}', '参考提出額', size=10, h='center', color=C_GRAYTXT, fill=FILL_G)
    put(ws, f'O{T0}', '出精値引率', size=10, h='center', color=C_GRAYTXT, fill=FILL_G)
    for i, rate in enumerate([0.10, 0.15, 0.20, 0.25, 0.30]):
        rr = T0 + 1 + i
        put(ws, f'M{rr}', rate, h='center', fmt='"粗利率 "0%')
        put(ws, f'N{rr}', f'=ROUNDUP(COST_TOTAL/(1-M{rr}),-4)', h='right', fmt=FMT_AMT)
        put(ws, f'O{rr}', f'=IF(K{SUB}=0,"",1-N{rr}/K{SUB})', h='right', fmt='0.0%')
    grid(ws, f'M{T0}:O{T0 + 5}', outer='thin', vert='thin', horiz='thin')
    page_setup(ws, f'A1:L{last_note}', fit_h=1)
    return dict(sub=SUB, disc=DISC, total=TOTAL, last_note=last_note, cost_cell='$O$27', alerts=A0)


# =====================================================================
# 御見積条件
# =====================================================================
def build_conditions(ws, det):
    W = dict(A=5, B=10, C=25, D=20, E=16, F=19)
    set_widths(ws, W)
    BF = W['B'] + W['C'] + W['D'] + W['E'] + W['F']
    put(ws, 'A1', '御見積条件・注意事項', size=14, bold=True)
    ws.row_dimensions[1].height = 24
    put(ws, 'A2', PROJECT_FULL + '　御見積書 添付', size=10)
    ws.row_dimensions[3].height = 8
    r = 4
    put(ws, f'A{r}', '【1】御見積条件・注意事項', size=12, bold=True)
    ws.row_dimensions[r].height = 22
    r += 1
    for i, t in enumerate(NOTES['conditions'], 1):
        put(ws, f'A{r}', f'{i}.', h='right', v='top')
        merge(ws, f'B{r}:F{r}')
        put(ws, f'B{r}', t, wrap=True, v='top')
        ws.row_dimensions[r].height = row_h(est_lines(t, BF, 11), 11, pad=8)
        r += 1
    ws.row_dimensions[r].height = 10
    r += 1
    ws.row_breaks.append(Break(id=r - 1))
    put(ws, f'A{r}', '【2】想定数量（推測）項目の前提', size=12, bold=True)
    ws.row_dimensions[r].height = 22
    r += 1
    merge(ws, f'A{r}:F{r}')
    put(ws, f'A{r}', '図面に数量・経路の記載がなく当社にて想定した項目です。「図面実測＋想定」は平面延長を図面計測し、立上り等を想定加算したものです。実数量確定後に精算させていただきます。', size=10, wrap=True)
    ws.row_dimensions[r].height = row_h(2, 10, pad=4)
    r += 1
    hdr = r
    for ref, txt in [('A', 'No.'), ('B', '内訳書'), ('C', '項　　目'), ('D', '前提・想定内容'), ('F', '数量（区分）')]:
        put(ws, f'{ref}{r}', txt, size=10, h='center', fill=FILL_G)
    merge(ws, f'D{r}:E{r}')
    fill_range(ws, f'A{r}:F{r}', FILL_G)
    ws.row_dimensions[r].height = 20
    r += 1
    k = 0
    for ln in det['lines']:
        if ln['kind'] != 'item' or ln.get('basis') not in ('推測', '実測+推測') or not ln.get('memo'):
            continue
        k += 1
        sec = '２' if ln['_sec'] == 'AC' else '３'
        item = ln['name'] + (' ' + ln['spec'] if ln.get('spec') else '')
        qtxt = f"{ln['qty']:,} {ln['unit']}（{'想定' if ln['basis'] == '推測' else '実測＋想定'}）"
        memo = ln['memo'].replace('（エスト見積も 56m）', '')
        put(ws, f'A{r}', k, size=10, h='center')
        put(ws, f'B{r}', sec, size=10, h='center')
        put(ws, f'C{r}', item, size=10, wrap=True, indent=1)
        merge(ws, f'D{r}:E{r}')
        put(ws, f'D{r}', memo, size=10, wrap=True, indent=1)
        put(ws, f'F{r}', qtxt, size=10, wrap=True, h='center')
        nl = max(est_lines(item, W['C'] - 3, 10), est_lines(memo, W['D'] + W['E'] - 3, 10), est_lines(qtxt, W['F'], 10))
        ws.row_dimensions[r].height = max(18, row_h(nl, 10, pad=4))
        r += 1
    grid(ws, f'A{hdr}:F{r - 1}', outer='medium', vert='thin', horiz='thin')
    ws.row_dimensions[r].height = 10
    r += 1
    put(ws, f'A{r}', '【3】図面読み取りに関する確認事項', size=12, bold=True)
    ws.row_dimensions[r].height = 22
    r += 1
    for i, t in enumerate(NOTES['checks'], 1):
        put(ws, f'A{r}', f'{i}.', h='right', v='top')
        merge(ws, f'B{r}:F{r}')
        put(ws, f'B{r}', t, wrap=True, v='top')
        ws.row_dimensions[r].height = row_h(est_lines(t, BF, 11), 11, pad=8)
        r += 1
    page_setup(ws, f'A1:F{r - 1}', fit_h=0)
    return dict(last=r - 1, n_assump=k)


# =====================================================================
# エスト比較（内部）
# =====================================================================
REASONS = {
    'sa_out': '図面は AC-1×2・AC-2 の3台（エストは2台）',
    'tr_sa': '同上（3台）', 'lk_sa': '同上（3系統）', 'vc_sa': '同上（3系統）',
    'cov_rac': '図面実測 2.11m→3m（エスト 5m）', 'cov_pac': '図面実測 7.90m→8m（エスト 10m）',
    'chk': 'エストに単価・行なし（当社単価で追加）', 'frame1': 'エストは 0（単価なし）→ 要見積', 'frame2': '同上',
}


def build_compare(ws, det):
    set_widths(ws, dict(A=12, B=34, C=30, D=8, E=10, F=12, G=8, H=10, I=12, J=11, K=44))
    put(ws, 'A1', 'エスト見積 6330560-1 との比較（内部用）　単価はエストに合わせ、数量は図面拾い。差額の主因を右列に記載', size=11, bold=True)
    hdr = ['区分', '品名', '規格', 'エスト数量', 'エスト単価', 'エスト金額', '当社数量', '当社単価', '当社金額', '差額', '差異の理由']
    for i, h in enumerate(hdr):
        put(ws, f'{chr(65 + i)}3', h, size=10, bold=True, h='center', fill=FILL_G)
    r = 4
    for ln in det['lines']:
        est = ln.get('est')
        sec = '空調設備工事' if ln['_sec'] == 'AC' else '配管設備工事'
        dr = ln['_row']
        put(ws, f'A{r}', sec, size=10)
        put(ws, f'B{r}', ln['name'], size=10)
        put(ws, f'C{r}', ln.get('spec') or None, size=9, shrink=True)
        if est:
            put(ws, f'D{r}', est[0], size=10, h='right', fmt='#,##0.##')
            put(ws, f'E{r}', est[1], size=10, h='right', fmt=FMT_AMT)
            put(ws, f'F{r}', est[2], size=10, h='right', fmt=FMT_SIGNED)
        put(ws, f'G{r}', f"={q(SH_DETAIL)}!E{dr}", size=10, h='right', fmt='#,##0.##')
        put(ws, f'H{r}', f"={q(SH_DETAIL)}!G{dr}" if ln['kind'] in ('item', 'tbd') else None, size=10, h='right', fmt=FMT_AMT)
        put(ws, f'I{r}', f"={q(SH_DETAIL)}!H{dr}", size=10, h='right', fmt=FMT_SIGNED)
        put(ws, f'J{r}', f'=I{r}-F{r}', size=10, h='right', fmt=FMT_SIGNED)
        reason = REASONS.get(ln.get('key'), '')
        if not reason and ln['kind'] in ('pct', 'labor'):
            reason = '数量同一のため同額（率・算式はエストと同じ）'
        if not reason and ln['kind'] == 'exp':
            reason = '基礎額の差による（率はエストと同じ）'
        put(ws, f'K{r}', reason or None, size=9, fill=(FILL_AMB if ln.get('key') in REASONS else None))
        r += 1
    for sec, name, qty, price, amt, reason in LN.EST_ONLY:
        put(ws, f'A{r}', sec, size=10)
        put(ws, f'B{r}', name, size=10)
        put(ws, f'D{r}', qty, size=10, h='right')
        put(ws, f'E{r}', price, size=10, h='right', fmt=FMT_AMT)
        put(ws, f'F{r}', amt, size=10, h='right', fmt=FMT_SIGNED)
        put(ws, f'I{r}', 0, size=10, h='right', fmt=FMT_SIGNED)
        put(ws, f'J{r}', f'=I{r}-F{r}', size=10, h='right', fmt=FMT_SIGNED)
        put(ws, f'K{r}', reason, size=9, fill=FILL_AMB)
        r += 1
    grid(ws, f'A3:K{r - 1}', outer='thin', vert='thin', horiz='hair')
    r += 1
    put(ws, f'A{r}', '集計', bold=True, fill=FILL_L)
    r += 1
    for i, h in enumerate(['項目', '', '', '', '', 'エスト', '', '', '当社', '差額']):
        if h:
            put(ws, f'{chr(65 + i)}{r}', h, size=10, bold=True, h='center', fill=FILL_G)
    r += 1
    s0 = r
    ac_t, pp_t, oh, gt = det['sec']['AC']['total'], det['sec']['PIPE']['total'], det['oh'], det['gt']
    for lab, ev, ov in [('空調設備工事 計', 174000, f"={q(SH_DETAIL)}!H{ac_t}"),
                        ('配管設備工事 計', 1112000, f"={q(SH_DETAIL)}!H{pp_t}"),
                        ('工事諸経費', 64300, f"={q(SH_DETAIL)}!H{oh}"),
                        ('小計', ES.EST_SUBTOTAL, f"={q(SH_DETAIL)}!H{gt}"),
                        ('出精値引', ES.EST_DISCOUNT, f"={q(SH_COVER)}!K27"),
                        ('合計（税別）', ES.EST_TOTAL, f"={q(SH_COVER)}!K28"),
                        ('原価合計（当社原価モデル）', None, f"={q(SH_COVER)}!O27"),
                        ('粗利率', f'=IF(I{s0 + 6}="","",(F{s0 + 5}-I{s0 + 6})/F{s0 + 5})', f"={q(SH_COVER)}!Q28")]:
        put(ws, f'A{r}', lab, size=10, bold=lab.startswith('合計'))
        is_rate = lab == '粗利率'
        put(ws, f'F{r}', ev, size=10, h='right', fmt=('0.0%' if is_rate else FMT_SIGNED))
        put(ws, f'I{r}', ov, size=10, h='right', fmt=('0.0%' if is_rate else FMT_SIGNED))
        if not is_rate and ev is not None:
            put(ws, f'J{r}', f'=I{r}-F{r}', size=10, h='right', fmt=FMT_SIGNED)
        r += 1
    put(ws, f'K{s0 + 7}', 'エスト側の粗利率はエスト提出額を当社原価（エストと数量が異なる分を含む）で見た目安', size=9)
    grid(ws, f'A{s0 - 1}:J{r - 1}', outer='thin', vert='thin', horiz='hair')
    r += 1
    put(ws, f'A{r}', 'エストの計算規則（見積金額から逆算し、エスト数量で各計が一致することを data/simulate.py で確認）', size=10, bold=True)
    for t in ['冷媒管: 継手類30%・消耗品15%・支持金物40%（管金額比）。塩ビ管: 20%・10%・25%',
              '配管工費＝国交省 公共建築工事標準単価積算基準 R8 の歩掛×33,750円/人。電線材料施工費＝0.017人/m×33,600円/人',
              '部門内経費: 空調 3%→4%→8%、配管 3%→5%→10% を順に積上げ（100円未満切上げ）、計を1,000円未満切上げ',
              '工事諸経費＝（空調＋配管）×5%。提出額は出精値引で調整。法定福利費は工事費に内含して表示']:
        r += 1
        put(ws, f'A{r}', '・' + t, size=9)
    page_setup(ws, f'A1:K{r}', fit_h=0)
    ws.page_setup.orientation = 'landscape'


def build_master(ws):
    set_widths(ws, dict(A=30, B=34, C=70))
    put(ws, 'A1', 'エスト単価マスタ（参考）　内訳書はエスト見積 6330560-1 の単価を優先し、このマスタの値は内訳書 AA 列に併記', size=11, bold=True)
    put(ws, 'A2', '見積とマスタで単価が異なる行（据付・基礎）はエスト見積の値を採用。マスタを採用する場合は内訳書 G 列を差し替え', size=9)
    for i, h in enumerate(['区分', '種別', '単価（円）']):
        put(ws, f'{chr(65 + i)}4', h, size=10, bold=True, h='center', fill=FILL_G)
    r = 5
    for a, b, c in ES.EST_B:
        put(ws, f'A{r}', a, size=10)
        put(ws, f'B{r}', b, size=10)
        put(ws, f'C{r}', c, size=10)
        r += 1
    grid(ws, f'A4:C{r - 1}', outer='thin', vert='thin', horiz='hair')
    page_setup(ws, f'A1:C{r}', fit_h=0)
    ws.page_setup.orientation = 'landscape'


def build_basis(ws):
    set_widths(ws, dict(A=34, B=14, C=12, D=12, E=12, F=70))
    put(ws, 'A1', '数量拾い根拠（内部用）　※空調設備図 F-02-1（A3 1/50）のベクトル線分（黒・線幅0.72pt）を連結し系統別に延長を集計（data/takeoff_f02.py）。'
        '「実測」=図面線分の平面延長、「想定」=立上り・立下り等の図面に無い部分', size=10, bold=True)
    put(ws, 'A2', '縮尺確認: 通り芯 Y6-Y5 910mm＝51.6pt、Y5-Y4 1,820mm＝103.3pt → 1pt＝17.64mm。拾い根拠画像: takeoff_F02_冷媒ドレン.png', size=9)
    r = 4
    for title, headers, rows in IT.BASIS_TABLES:
        put(ws, f'A{r}', title, bold=True, fill=FILL_L)
        r += 1
        for i, h in enumerate(headers):
            put(ws, f'{chr(65 + i)}{r}', h, size=10, bold=True, h='center', fill=FILL_G)
        hdr = r
        r += 1
        for row in rows:
            for i, v in enumerate(row):
                put(ws, f'{chr(65 + i)}{r}', v, size=10, h=('left' if isinstance(v, str) else 'right'),
                    fmt=('#,##0.00' if isinstance(v, float) else None))
            r += 1
        grid(ws, f'A{hdr}:{chr(64 + len(headers))}{r - 1}', outer='thin', vert='thin', horiz='hair')
        r += 1
    page_setup(ws, f'A1:F{r}', fit_h=0)
    ws.page_setup.orientation = 'landscape'


def build_cost(ws):
    set_widths(ws, dict(A=40, B=16, C=90))
    put(ws, 'A1', '原価根拠（内部用）　材料＝建設物価相当、労務＝人工×労務原価日額、法定福利費＝労務原価×事業主負担率', size=11, bold=True)
    r = 3
    put(ws, f'A{r}', '出典', bold=True, fill=FILL_L)
    r += 1
    for a, b in PV.SOURCES:
        put(ws, f'A{r}', a, size=10, bold=True)
        merge(ws, f'B{r}:C{r}')
        put(ws, f'B{r}', b, size=9, wrap=True)
        ws.row_dimensions[r].height = 30
        r += 1
    r += 1
    put(ws, f'A{r}', '公共工事設計労務単価（令和8年3月適用・広島県、円/日）', bold=True, fill=FILL_L)
    r += 1
    for k, v in PV.PUBLIC_LABOR_R8_HIROSHIMA.items():
        put(ws, f'A{r}', k, size=10)
        put(ws, f'B{r}', v, size=10, h='right', fmt=FMT_AMT)
        r += 1
    put(ws, f'A{r}', '※ 設計労務単価は事業主負担の法定福利費等を含まない。当社の労務原価日額（表紙 N6）との比較・感度確認に使用', size=9)
    r += 2
    put(ws, f'A{r}', '国交省 公共建築工事標準単価積算基準（令和8年改定）配管工 歩掛（人/m）', bold=True, fill=FILL_L)
    r += 1
    for k, v in PV.MLIT_BUGAKARI.items():
        put(ws, f'A{r}', k, size=10)
        put(ws, f'B{r}', v, size=10, h='right', fmt='0.000')
        r += 1
    put(ws, f'A{r}', '※ 被覆銅管は液管・ガス管それぞれの外径の歩掛を合算。継手・接合材 管単価×0.30、支持金物 ×0.40、雑材料 ×0.15、管長 1.05（塩ビ 1.10）', size=9)
    r += 2
    put(ws, f'A{r}', 'メーカー公表価格（参考）', bold=True, fill=FILL_L)
    r += 1
    for k, v in PV.LIST_PRICE_REF.items():
        put(ws, f'A{r}', k, size=10)
        put(ws, f'B{r}', v, size=10, h='right', fmt=FMT_AMT)
        put(ws, f'C{r}', f'1mあたり {v / 20:,.0f} 円。エスト材料単価はこの約6割（3分5分 5,138円/m＝58.6%、2分3分 3,388円/m＝61.3%）＝実勢価格水準', size=9)
        r += 1
    r += 1
    put(ws, f'A{r}', '原価の前提（内訳書 L・M 列の黄セル）', bold=True, fill=FILL_L)
    for t in ['据付の人工: 天井カセット 0.9、天井ビルトイン 1.2、RA室内 0.4、SA室外 下段0.8・上段1.3（揚重）、RA室外 上段0.5 人工/台',
              '縁石: 材料3,000＋0.15人工/台。試運転: SA 材料450＋0.45人工、RA 0.1人工。気密: SA 2,000＋0.45、RA 1,000＋0.2。真空引き: SA 500＋0.15、RA 300＋0.1',
              '化粧カバー: 材料 小径1,800・大径2,400円/m＋0.10人工/m。GW保温: 25A 900＋0.06、30A 1,000＋0.065、40A 1,100＋0.07（市場単価は非公表のため当社想定）',
              'スリーブ: 4箇所×（材料1,500＋0.25人工）。架台: 要見積（概算原価 300,000／250,000、鉄工所またはメーカー見積で差し替え）',
              '材料原価率を下げると（実際の仕入が建設物価より安い場合）提出額は下がる。材料原価率0.8で提出額は約1,180,000円（粗利20%）']:
        r += 1
        merge(ws, f'A{r}:C{r}')
        put(ws, f'A{r}', '・' + t, size=9, wrap=True)
        ws.row_dimensions[r].height = 28
    page_setup(ws, f'A1:C{r}', fit_h=0)
    ws.page_setup.orientation = 'landscape'


def build_book():
    wb = Workbook()
    base = fnt(11)
    wb._fonts = IndexedList([base])
    wb._named_styles['Normal'].font = base
    wb.properties.creator = COMPANY
    wb.properties.title = '御見積書 ' + PROJECT_FULL
    ws_cover = wb.active
    ws_cover.title = SH_COVER
    ws_det = wb.create_sheet(SH_DETAIL)
    ws_wf = wb.create_sheet(SH_WF)
    ws_cd = wb.create_sheet(SH_COND)
    ws_cmp = wb.create_sheet(SH_CMP)
    ws_mst = wb.create_sheet(SH_MST)
    ws_bs = wb.create_sheet(SH_BASIS)
    ws_cost = wb.create_sheet(SH_COST)
    det = build_detail(ws_det)
    wf = build_welfare(ws_wf, det)
    cv = build_cover(ws_cover, det, wf)
    cd = build_conditions(ws_cd, det)
    build_compare(ws_cmp, det)
    build_master(ws_mst)
    build_basis(ws_bs)
    build_cost(ws_cost)
    names = {
        'TARGET_MARGIN': f"{q(SH_COVER)}!$N$4",
        'ROUND_UNIT': f"{q(SH_COVER)}!$N$5",
        'LABOR_COST': f"{q(SH_COVER)}!$N$6",
        'MAT_RATIO': f"{q(SH_COVER)}!$N$7",
        'EST_LABOR_PIPE': f"{q(SH_COVER)}!$N$8",
        'EST_LABOR_ELEC': f"{q(SH_COVER)}!$N$9",
        'WELFARE_RATE': f"{q(SH_WF)}!{wf['rate_cell']}",
        'INCLUDE_TBD': f"{q(SH_COVER)}!$N$11",
        'TARGET': f"{q(SH_COVER)}!$N$12",
        'SUBMIT': f"{q(SH_COVER)}!$N$13",
        'COST_TOTAL': f"{q(SH_COVER)}!{cv['cost_cell']}",
    }
    for k, ref in names.items():
        wb.defined_names[k] = DefinedName(k, attr_text=ref)
    wb.active = 0
    for ws in wb.worksheets:
        ws.sheet_view.tabSelected = (ws.title == SH_COVER)
    wb.calculation.fullCalcOnLoad = True
    os.makedirs(OUT_DIR, exist_ok=True)
    out = os.path.join(OUT_DIR, FILE)
    wb.save(out)
    return dict(out=out, detail=dict(sec=det['sec'], oh=det['oh'], gt=det['gt']), cover=cv, conditions=cd, quantities=IT.Q)


if __name__ == '__main__':
    print(json.dumps(build_book(), ensure_ascii=False, indent=1))
