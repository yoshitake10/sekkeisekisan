# -*- coding: utf-8 -*-
"""
エスト準拠 御見積書ビルダー（宮地機工様式・共通）
==================================================
物件ごとの job モジュール（定数・明細・注記）を受け取り、ワークブックを作る。
シート: 表紙(工事) ／ 内訳書(空調設備工事) ／ 法定福利費内訳明細書 ／ 御見積条件（提出用）
        ／ 単価出典(内部) ／ エストマスタ(参考) ／ 数量拾い根拠(内部) ／ 原価根拠(内部)
job モジュールに必要な属性:
  CLIENT, PROJECT_1, PROJECT_2, PROJECT_FULL, SITE, TERM, PAYMENT, VALIDITY, EST_DATE, FILE,
  NOTES = dict(cover=[...], conditions=[...], checks=[...]), SECTIONS（明細）, BASIS_TABLES,
  COVER_SCOPE = dict(include='...', exclude='...')（表紙21行目の注記。要見積の計上有無で切替）,
  EQUIP_ROW = ('空調設備機器', '別途')（表紙①の名称と表示）, COST_NOTES（原価の前提）,
  任意: BASIS_NOTE（数量拾い根拠の注記2行）, ALERT_NOTES（表紙内部欄の追加アラート）,
        WORK_TITLE（表紙20行目・内訳書4行目の工事種別。既定 '空気調和工事'）, OH_NO（工事諸経費の番号。既定 '８'）,
        WF_EQUIP_TXT（法定福利費明細の「○○は含みません」。既定 '空調機器'）,
        EQUIP_ROW = None で表紙の静的な機器行を出さない（機器を明細の部門として計上する場合）
部門（SECTIONS[*]）: key, no, title, lines。任意: exp_key（部門内経費の率のキー。既定 key）,
  oh（工事諸経費の対象。既定 True。機器部門は False）, break_after（この部門の後で改ページ。
  どの部門にも指定がなければ先頭部門の後で改ページ）。部門内経費の行（kind='exp'）が無い部門は計＝明細の合計。
明細（SECTIONS[*]['lines']）の kind:
  head / item（cat=単価カタログのキー、または price・src・ref を直接指定）/ tbd（要見積、prov=概算原価）
  / pct（rate・base）/ labor（base・rate='EST_LABOR_PIPE' または 'EST_LABOR_ELEC'）/ exp（idx=0,1,2）
  item の mat: 数値（材料原価/単位）または ('EST', 割増)＝単価×材料原価率×割増、('EQ', 比)＝機器の提出単価×材料原価率×比
"""
import math
import os

from openpyxl import Workbook
from openpyxl.styles import PatternFill
from openpyxl.utils.indexed_list import IndexedList
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.pagebreak import Break

from helpers import (fnt, put, merge, set_border, grid, hline, fill_range, set_widths,
                     est_lines, row_h, page_setup, FILL_Y, FILL_G, FILL_L, C_GRAYTXT, FMT_AMT, FMT_QTY)
import est_rules as R

SH_COVER = '表紙(工事)'
SH_DETAIL = '内訳書(空調設備工事)'
SH_WF = '法定福利費内訳明細書'
SH_COND = '御見積条件'
SH_SRC = '単価出典(内部)'
SH_MST = 'エストマスタ(参考)'
SH_BASIS = '数量拾い根拠(内部)'
SH_COST = '原価根拠(内部)'

CIRC = '①②③④⑤⑥⑦⑧⑨'
C_RED = 'C00000'
FILL_RED = PatternFill('solid', fgColor='F8CBAD')
FILL_PINK = PatternFill('solid', fgColor='FDE9E9')
FILL_AMB = PatternFill('solid', fgColor='FFE699')
FMT_SIGNED = '#,##0_ ;[Red]-#,##0_ '
FMT_TBD = '#,##0_ ;[Red]-#,##0_ ;"別途見積"_ '
ALERT_OF = {R.SRC_OWN: '単価要確認', R.SRC_DERIVED: '単価要確認（推定）'}

COLS = dict(J='単価区分', K='アラート', L='材料原価/単位', M='人工/単位', N='労務原価/単位', O='原価単価', P='原価金額',
            Q='労務原価金額', R='労務費相当額(提出)', S='粗利(定価)', T='粗利率', U='率・歩掛', V='要見積 概算原価/単位',
            W='単価の参照元', X='エストマスタ参考', Y='数量根拠', Z='出典・メモ（内部）')


def q(sheet):
    return "'" + sheet + "'"


def resolve(ln):
    """item 行の単価・単価区分・参照元をカタログから解決"""
    if ln['kind'] != 'item':
        return ln
    ln = dict(ln)
    if ln.get('cat'):
        price, src, ref = R.CAT[ln['cat']]
        ln.setdefault('price', price)
        ln.setdefault('src', src)
        ln.setdefault('ref', ref)
        if 'bug' not in ln and ln['cat'] in R.BUG:
            ln['bug'] = R.BUG[ln['cat']]
    if ln.get('src') == R.SRC_OWN and ln.get('price') is None:
        ln['price'] = R.own_price(ln['mat'], ln['md'])
    return ln


def sections(job):
    out = []
    for sec in job.SECTIONS:
        out.append(dict(sec, lines=[resolve(ln) for ln in sec['lines']]))
    return out


def work_title(job):
    return getattr(job, 'WORK_TITLE', '空気調和工事')


def cover_rows(job, secs):
    """表紙の金額行の並び: [('equip', None) | ('sec', sec) ..., ('oh', None)] と 小計・値引・合計の行番号"""
    rows = []
    if getattr(job, 'EQUIP_ROW', None):
        rows.append(('equip', None))
    rows += [('sec', s) for s in secs]
    rows.append(('oh', None))
    first = 22
    sub = first + len(rows)
    return dict(rows=rows, first=first, sub=sub, disc=sub + 1, total=sub + 2)


# =====================================================================
# 内訳書
# =====================================================================
def build_detail(ws, job):
    secs = sections(job)
    set_widths(ws, dict(A=4, B=27, C=24, D=4, E=8, F=5, G=10, H=13, I=15,
                        J=13, K=13, L=10, M=7, N=10, O=10, P=12, Q=11, R=12, S=11, T=8, U=8, V=11,
                        W=46, X=36, Y=11, Z=70))
    merge(ws, 'A1:I1')
    put(ws, 'A1', '内　訳　書', size=16, bold=True, h='center')
    merge(ws, 'J1:Z1')
    put(ws, 'J1', '原　価・単価出典（内部用・印刷範囲外）　黄セル＝入力、赤＝要見積、橙＝単価要確認',
        size=11, bold=True, h='center', color=C_GRAYTXT, fill=FILL_G)
    fill_range(ws, 'J1:Z1', FILL_G)
    ws.row_dimensions[1].height = 26
    put(ws, 'A2', '工事名：' + job.PROJECT_FULL)
    put(ws, 'I2', '内訳書 No.1', h='right')
    merge(ws, 'A3:B3')
    merge(ws, 'C3:D3')
    for ref, txt in [('A3', '品　　名'), ('C3', '規格・寸法'), ('E3', '数 量'), ('F3', '単位'),
                     ('G3', '単　価'), ('H3', '金　　額'), ('I3', '備　考')]:
        put(ws, ref, txt, h='center', fill=FILL_G)
    fill_range(ws, 'A3:I3', FILL_G)
    for col, txt in COLS.items():
        put(ws, f'{col}3', txt, size=9, h='center', fill=FILL_G, color=C_GRAYTXT, shrink=True)
    ws.row_dimensions[3].height = 22
    merge(ws, 'A4:I4')
    put(ws, 'A4', work_title(job), bold=True)
    r = 5
    rows, sec_info, all_lines = {}, {}, []
    for si, sec in enumerate(secs):
        merge(ws, f'A{r}:I{r}')
        put(ws, f'A{r}', f"{sec['no']}　{sec['title']}", bold=True)
        ws.row_dimensions[r].height = 20
        r += 1
        first = r
        exp_rows = []
        for ln in sec['lines']:
            k = ln['kind']
            merge(ws, f'A{r}:B{r}')
            merge(ws, f'C{r}:D{r}')
            ws.row_dimensions[r].height = 18
            if k == 'head':
                put(ws, f'A{r}', '　' + ln['name'], size=10, bold=True)
                r += 1
                continue
            if k == 'exp':
                name, rate = R.EST_EXP[sec.get('exp_key', sec['key'])][ln['idx']]
                ln = dict(ln, key=f"{sec['key']}_exp{ln['idx']}", name=name, rate=rate)
            ln['_row'], ln['_sec'] = r, sec['key']
            rows[ln['key']] = r
            all_lines.append(ln)
            put(ws, f'A{r}', '　　' + ln['name'], size=10, shrink=True)
            put(ws, f'X{r}', ln.get('master') or None, size=9)
            put(ws, f'Y{r}', ln.get('basis') or None, size=9, h='center', shrink=True,
                fill=(FILL_PINK if ln.get('basis') in ('推測', '実測+推測') else None))
            if k in ('item', 'tbd'):
                spec = ln.get('spec') or ''
                nl = est_lines(spec, 27, 9) if spec else 1
                color = C_RED if k == 'tbd' else None
                if nl > 1:
                    put(ws, f'C{r}', spec, size=9, wrap=True, indent=1, color=color)
                    ws.row_dimensions[r].height = max(18, row_h(nl, 9, pad=4))
                else:
                    put(ws, f'C{r}', spec or None, size=10, shrink=True, indent=1, color=color)
                put(ws, f'E{r}', ln['qty'], size=10, h='right', fmt='#,##0_ ')
                put(ws, f'F{r}', ln['unit'], size=10, h='center')
            if k == 'item':
                put(ws, f'G{r}', ln['price'], size=10, h='right', fmt=FMT_AMT)
                put(ws, f'H{r}', f'=ROUND(E{r}*G{r},0)', size=10, h='right', fmt=FMT_SIGNED)
                pub = ln.get('pub') or {'実測+推測': '図面実測＋想定', '推測': '想定数量'}.get(ln.get('basis'), '')
                put(ws, f'I{r}', pub or None, size=10, h='center', shrink=True)
                put(ws, f'J{r}', ln['src'], size=9, h='center', shrink=True)
                alert = ALERT_OF.get(ln['src'])
                put(ws, f'K{r}', alert, size=9, h='center', shrink=True, fill=(FILL_AMB if alert else None))
                mat = ln['mat']
                if isinstance(mat, tuple):
                    put(ws, f'L{r}', f'=ROUND(G{r}*MAT_RATIO*{mat[1]},0)', size=10, h='right', fmt=FMT_AMT)
                else:
                    put(ws, f'L{r}', mat, size=10, h='right', fmt=FMT_AMT, fill=FILL_Y)
                put(ws, f'M{r}', ln.get('md', 0), size=10, h='right', fmt='0.000', fill=FILL_Y)
                put(ws, f'N{r}', f'=ROUND(M{r}*LABOR_COST,0)', size=10, h='right', fmt=FMT_AMT)
                put(ws, f'O{r}', f'=L{r}+N{r}', size=10, h='right', fmt=FMT_AMT)
                put(ws, f'P{r}', f'=E{r}*O{r}', size=10, h='right', fmt=FMT_SIGNED)
                put(ws, f'Q{r}', f'=E{r}*N{r}', size=10, h='right', fmt=FMT_SIGNED)
                put(ws, f'R{r}', f'=IF(H{r}=0,0,IF(E{r}*M{r}*EST_LABOR_PIPE>H{r},H{r},E{r}*M{r}*EST_LABOR_PIPE))',
                    size=10, h='right', fmt=FMT_SIGNED)
                if ln.get('bug'):
                    put(ws, f'U{r}', ln['bug'], size=10, h='right', fmt='0.000', fill=FILL_Y)
                put(ws, f'W{r}', ln.get('ref') or None, size=9)
                memo = ln.get('memo') or ''
                if isinstance(mat, tuple):
                    pre = (f'機器原価＝提出単価×材料原価率×{mat[1]}（仕入率÷提出率）' if mat[0] == 'EQ'
                           else f'材料原価＝単価×材料原価率×管長割増{mat[1]}（建設物価相当）')
                    memo = pre + ('｜' + memo if memo else '')
                put(ws, f'Z{r}', memo or None, size=9)
            elif k == 'tbd':
                put(ws, f'A{r}', '　　' + ln['name'], size=10, shrink=True, color=C_RED)
                put(ws, f'G{r}', f'=IF(INCLUDE_TBD=1,ROUNDUP(V{r}/(1-TARGET_MARGIN),-3),"")', size=10, h='right', fmt=FMT_AMT)
                put(ws, f'H{r}', f'=IF(INCLUDE_TBD=1,ROUND(E{r}*G{r},0),0)', size=10, h='right', fmt=FMT_TBD, color=C_RED)
                put(ws, f'I{r}', '=IF(INCLUDE_TBD=1,"概算（要見積）","別途見積")', size=10, h='center', shrink=True, color=C_RED)
                put(ws, f'J{r}', R.SRC_TBD, size=9, h='center', fill=FILL_RED, bold=True)
                put(ws, f'K{r}', '要見積', size=9, h='center', fill=FILL_RED, bold=True, color=C_RED)
                put(ws, f'V{r}', ln['prov'], size=10, h='right', fmt=FMT_AMT, fill=FILL_Y)
                put(ws, f'L{r}', f'=V{r}', size=10, h='right', fmt=FMT_AMT)
                put(ws, f'M{r}', 0, size=10, h='right', fmt='0.000')
                put(ws, f'N{r}', 0, size=10, h='right', fmt=FMT_AMT)
                put(ws, f'O{r}', f'=L{r}+N{r}', size=10, h='right', fmt=FMT_AMT)
                put(ws, f'P{r}', f'=IF(INCLUDE_TBD=1,E{r}*O{r},0)', size=10, h='right', fmt=FMT_SIGNED)
                put(ws, f'Q{r}', 0, size=10, h='right', fmt=FMT_SIGNED)
                put(ws, f'R{r}', 0, size=10, h='right', fmt=FMT_SIGNED)
                put(ws, f'W{r}', '単価情報なし（メーカー・鉄工所の見積を取得）', size=9)
                put(ws, f'Z{r}', ln.get('memo') or None, size=9)
            elif k == 'pct':
                bases = [rows[b] for b in ln['base']]
                put(ws, f'E{r}', 1, size=10, h='right', fmt='#,##0_ ')
                put(ws, f'F{r}', '式', size=10, h='center')
                put(ws, f'H{r}', '=ROUND((' + '+'.join(f'H{b}' for b in bases) + f')*U{r},0)', size=10, h='right', fmt=FMT_SIGNED)
                put(ws, f'J{r}', R.SRC_CALC, size=9, h='center', shrink=True)
                put(ws, f'U{r}', ln['rate'], size=10, h='right', fmt='0%', fill=FILL_Y)
                put(ws, f'P{r}', '=ROUND((' + '+'.join(f'P{b}' for b in bases) + f')*U{r},0)', size=10, h='right', fmt=FMT_SIGNED)
                put(ws, f'Q{r}', 0, size=10, h='right', fmt=FMT_SIGNED)
                put(ws, f'R{r}', 0, size=10, h='right', fmt=FMT_SIGNED)
                put(ws, f'W{r}', '管金額×率（エスト・公共積算と同率）。原価も同率', size=9)
            elif k == 'labor':
                bases = [rows[b] for b in ln['base']]
                put(ws, f'E{r}', 1, size=10, h='right', fmt='#,##0_ ')
                put(ws, f'F{r}', '式', size=10, h='center')
                put(ws, f'M{r}', '=' + '+'.join(f'E{b}*U{b}' for b in bases), size=10, h='right', fmt='0.000')
                put(ws, f'H{r}', f'=ROUNDUP(M{r}*{ln["rate"]},-1)', size=10, h='right', fmt=FMT_SIGNED)
                put(ws, f'J{r}', R.SRC_CALC, size=9, h='center', shrink=True)
                put(ws, f'L{r}', 0, size=10, h='right', fmt=FMT_AMT)
                put(ws, f'N{r}', f'=ROUND(M{r}*LABOR_COST,0)', size=10, h='right', fmt=FMT_AMT)
                put(ws, f'O{r}', f'=L{r}+N{r}', size=10, h='right', fmt=FMT_AMT)
                put(ws, f'P{r}', f'=E{r}*O{r}', size=10, h='right', fmt=FMT_SIGNED)
                put(ws, f'Q{r}', f'=E{r}*N{r}', size=10, h='right', fmt=FMT_SIGNED)
                put(ws, f'R{r}', f'=H{r}', size=10, h='right', fmt=FMT_SIGNED)
                put(ws, f'W{r}', ln.get('memo') or None, size=9)
            elif k == 'exp':
                put(ws, f'E{r}', 1, size=10, h='right', fmt='#,##0_ ')
                put(ws, f'F{r}', '式', size=10, h='center')
                put(ws, f'J{r}', '経費', size=9, h='center')
                put(ws, f'U{r}', ln['rate'], size=10, h='right', fmt='0%', fill=FILL_Y)
                for col in 'PQR':
                    put(ws, f'{col}{r}', 0, size=10, h='right', fmt=FMT_SIGNED)
                exp_rows.append(r)
            if k != 'exp':
                put(ws, f'S{r}', f'=IF(H{r}="","",H{r}-P{r})', size=10, h='right', fmt=FMT_SIGNED)
                put(ws, f'T{r}', f'=IF(OR(H{r}="",H{r}=0),"",S{r}/H{r})', size=10, h='right', fmt='0.0%')
            r += 1
        if exp_rows:
            e1, e2, e3 = exp_rows
            base = f'SUM(H{first}:H{e1 - 1})'
            put(ws, f'H{e1}', f'=ROUNDUP({base}*U{e1},-2)', size=10, h='right', fmt=FMT_SIGNED)
            put(ws, f'H{e2}', f'=ROUNDUP(({base}+H{e1})*U{e2},-2)', size=10, h='right', fmt=FMT_SIGNED)
            s3 = f'({base}+H{e1}+H{e2})'
            put(ws, f'H{e3}', f'=ROUNDUP({s3}+ROUNDUP({s3}*U{e3},-2),-3)-{s3}', size=10, h='right', fmt=FMT_SIGNED)
            put(ws, f'W{e1}', '基礎額×率（100円未満切上げ）', size=9)
            put(ws, f'W{e2}', '（基礎額＋消耗品）×率（100円未満切上げ）', size=9)
            put(ws, f'W{e3}', '（基礎額＋消耗品＋運搬費）×率を加え、計を1,000円未満切上げ', size=9)
        else:
            e3 = r - 1
        tot = r
        merge(ws, f'A{tot}:D{tot}')
        put(ws, f'A{tot}', '－　計　－', bold=True, h='center')
        put(ws, f'H{tot}', f'=SUM(H{first}:H{e3})', bold=True, h='right', fmt=FMT_SIGNED)
        put(ws, f'M{tot}', f'=SUMPRODUCT(E{first}:E{e3},M{first}:M{e3})', bold=True, h='right', fmt='0.00')
        for col in 'PQR':
            put(ws, f'{col}{tot}', f'=SUM({col}{first}:{col}{e3})', bold=True, h='right', fmt=FMT_SIGNED)
        put(ws, f'S{tot}', f'=H{tot}-P{tot}', bold=True, h='right', fmt=FMT_SIGNED)
        put(ws, f'T{tot}', f'=IF(H{tot}=0,"",S{tot}/H{tot})', bold=True, h='right', fmt='0.0%')
        ws.row_dimensions[tot].height = 20
        sec_info[sec['key']] = dict(first=first, exp=exp_rows, total=tot)
        r = tot + 1
        merge(ws, f'A{r}:B{r}')
        merge(ws, f'C{r}:D{r}')
        ws.row_dimensions[r].height = 10
        if (sec.get('break_after', False) if any('break_after' in x for x in secs) else si == 0):
            ws.row_breaks.append(Break(id=r))
        r += 1
    oh = r
    merge(ws, f'A{oh}:B{oh}')
    merge(ws, f'C{oh}:D{oh}')
    put(ws, f'A{oh}', getattr(job, 'OH_NO', '８') + '　工事諸経費', bold=True)
    put(ws, f'E{oh}', 1, size=10, h='right', fmt='#,##0_ ')
    put(ws, f'F{oh}', '式', size=10, h='center')
    oh_secs = [s for s in secs if s.get('oh', True)]
    oh_sum = '+'.join(f"H{sec_info[s['key']]['total']}" for s in oh_secs)
    put(ws, f'H{oh}', f'=ROUNDUP(({oh_sum})*U{oh},-2)', bold=True, h='right', fmt=FMT_SIGNED)
    put(ws, f'J{oh}', '経費', size=9, h='center')
    put(ws, f'U{oh}', R.EST_OVERHEAD, size=10, h='right', fmt='0%', fill=FILL_Y)
    put(ws, f'P{oh}', 0, size=10, h='right', fmt=FMT_SIGNED)
    put(ws, f'W{oh}', '（' + '＋'.join(s['title'] for s in oh_secs) + '）×5%（100円未満切上げ）', size=9)
    ws.row_dimensions[oh].height = 20
    r = oh + 1
    merge(ws, f'A{r}:B{r}')
    merge(ws, f'C{r}:D{r}')
    ws.row_dimensions[r].height = 10
    gt = r + 1
    merge(ws, f'A{gt}:D{gt}')
    put(ws, f'A{gt}', '【小　　計】', bold=True, h='center')
    tots = [sec_info[s['key']]['total'] for s in secs]
    put(ws, f'H{gt}', '=' + '+'.join(f'H{t}' for t in tots) + f'+H{oh}', bold=True, h='right', fmt=FMT_SIGNED)
    put(ws, f'I{gt}', '出精値引は表紙', size=9, h='center', shrink=True)
    put(ws, f'M{gt}', '=' + '+'.join(f'M{t}' for t in tots), bold=True, h='right', fmt='0.00')
    for col in 'PQR':
        put(ws, f'{col}{gt}', '=' + '+'.join(f'{col}{t}' for t in tots), bold=True, h='right', fmt=FMT_SIGNED)
    put(ws, f'S{gt}', f'=H{gt}-P{gt}', bold=True, h='right', fmt=FMT_SIGNED)
    put(ws, f'T{gt}', f'=IF(H{gt}=0,"",S{gt}/H{gt})', bold=True, h='right', fmt='0.0%')
    ws.row_dimensions[gt].height = 22
    grid(ws, f'A3:I{gt}', outer='medium', vert='thin', horiz='hair')
    grid(ws, f'J3:Z{gt}', outer='thin', vert='thin', horiz='hair')
    hline(ws, 'A3:I3', 'bottom', 'thin')
    hline(ws, 'J3:Z3', 'bottom', 'thin')
    for t in tots:
        hline(ws, f'A{t}:I{t}', 'top', 'thin')
        hline(ws, f'J{t}:Z{t}', 'top', 'thin')
    hline(ws, f'A{gt}:I{gt}', 'top', 'medium')
    hline(ws, f'J{gt}:Z{gt}', 'top', 'medium')
    for c in range(1, 10):
        set_border(ws.cell(gt, c), bottom='medium')
    set_border(ws.cell(gt, 1), left='medium')
    set_border(ws.cell(gt, 9), right='medium')
    ws.freeze_panes = 'A4'
    page_setup(ws, f'A1:I{gt}', fit_h=0, title_rows='1:3', footer='&P / &N')
    return dict(rows=rows, sec=sec_info, oh=oh, gt=gt, lines=all_lines, secs=secs)


# =====================================================================
# 法定福利費内訳明細書（工事費に内含）
# =====================================================================
def build_welfare(ws, job, det, cl):
    set_widths(ws, dict(A=2, B=25, C=9, D=9, E=14, F=16, G=31))
    merge(ws, 'B1:G1')
    put(ws, 'B1', '法 定 福 利 費 内 訳 明 細 書', size=14, bold=True, h='center')
    ws.row_dimensions[1].height = 26
    merge(ws, 'B2:G2')
    put(ws, 'B2', f'="（御見積書 No."&{q(SH_COVER)}!K3&" の添付資料）"', size=10, h='center')
    put(ws, 'B4', '見積年月日', fill=FILL_L, indent=1)
    merge(ws, 'C4:E4')
    put(ws, 'C4', f"={q(SH_COVER)}!J4", h='left', fmt=R.DATE_FMT, indent=1)
    put(ws, 'F4', '見積No.', fill=FILL_L, indent=1)
    put(ws, 'G4', f'=IF({q(SH_COVER)}!K3="","",{q(SH_COVER)}!K3)', h='left', indent=1)
    for r, lab, val in [(5, '発注者名', job.CLIENT + '　御中'), (6, '工事名', job.PROJECT_FULL), (7, '施工場所', job.SITE)]:
        put(ws, f'B{r}', lab, fill=FILL_L, indent=1)
        merge(ws, f'C{r}:G{r}')
        put(ws, f'C{r}', val, shrink=True, indent=1)
    put(ws, 'B8', '会社名', fill=FILL_L, indent=1)
    merge(ws, 'C8:E8')
    put(ws, 'C8', R.COMPANY, indent=1)
    put(ws, 'F8', '所在地', fill=FILL_L, indent=1)
    put(ws, 'G8', R.ADDR, indent=1)
    grid(ws, 'B4:G8', outer='thin', vert='thin', horiz='thin')
    for r in range(4, 9):
        ws.row_dimensions[r].height = 20
    put(ws, 'B10', '【1】　労務費相当額（Ｙ）の内訳', bold=True)
    note1 = ('※ 労務費相当額は、内訳書各項目の金額のうち労務費に相当する額（配管工費・電線材料施工費、および据付・保温・試験等の人工相当額）'
             f"に、出精値引き後の比率を乗じたものです。{getattr(job, 'WF_EQUIP_TXT', '空調機器')}は含みません。")
    merge(ws, 'B11:G11')
    put(ws, 'B11', note1, size=10, wrap=True)
    ws.row_dimensions[11].height = row_h(est_lines(note1, 104, 10), 10)
    put(ws, 'B12', '項　　目', h='center', fill=FILL_G)
    merge(ws, 'C12:D12')
    put(ws, 'C12', '労務費相当額（円）', size=10, h='center', fill=FILL_G)
    merge(ws, 'E12:G12')
    put(ws, 'E12', '摘　　要', h='center', fill=FILL_G)
    fill_range(ws, 'B12:G12', FILL_G)
    wsecs = [s for s in det['secs'] if s.get('oh', True)]
    items = [(13 + i, s['title'], f"={q(SH_DETAIL)}!R{det['sec'][s['key']]['total']}",
              f"内訳書 {s['no']} {s['title']}（定価ベース）") for i, s in enumerate(wsecs)]
    rs, ry = 13 + len(wsecs), 14 + len(wsecs)
    items += [(rs, '小　　計', '=' + '+'.join(f'C{x[0]}' for x in items), '定価ベース'),
              (ry, '労務費相当額（Ｙ）', f"=ROUNDDOWN(C{rs}*{q(SH_COVER)}!K{cl['total']}/{q(SH_COVER)}!K{cl['sub']},0)",
               '小計 × 出精値引き後比率（合計÷小計）')]
    for r, lab, f, memo in items:
        bold = r == ry
        put(ws, f'B{r}', lab, bold=bold, indent=1)
        merge(ws, f'C{r}:D{r}')
        put(ws, f'C{r}', f, bold=bold, h='right', fmt=FMT_AMT)
        merge(ws, f'E{r}:G{r}')
        put(ws, f'E{r}', memo, size=10, indent=1)
        ws.row_dimensions[r].height = 20
    grid(ws, f'B12:G{ry}', outer='medium', vert='thin', horiz='thin')
    hline(ws, f'B{ry}:G{ry}', 'top', 'double')
    r = ry + 2
    put(ws, f'B{r}', '【2】　法定福利費相当額（Ａ）の算定', bold=True)
    r += 1
    merge(ws, f'B{r}:G{r}')
    put(ws, f'B{r}', '計算式：　Ａ ＝ Ｙ × Ｚ（Ｚ：法定福利費事業者負担率、保険ごとに円未満切捨て）', size=10)
    r += 1
    hdr2 = r
    for ref, txt in [('B', '保険の種類'), ('E', '事業主負担料率'), ('F', '法定福利費（円）'), ('G', '算定根拠')]:
        put(ws, f'{ref}{r}', txt, size=10, h='center', fill=FILL_G)
    merge(ws, f'C{r}:D{r}')
    put(ws, f'C{r}', '労務費相当額（円）', size=10, h='center', fill=FILL_G)
    fill_range(ws, f'B{r}:G{r}', FILL_G)
    ws.row_dimensions[r].height = 22
    r += 1
    for name, rate, basis in R.WELFARE_RATES:
        put(ws, f'B{r}', name, indent=1)
        merge(ws, f'C{r}:D{r}')
        put(ws, f'C{r}', f'=$C${ry}', h='right', fmt=FMT_AMT)
        put(ws, f'E{r}', rate, h='right', fmt='0.000%_ ')
        put(ws, f'F{r}', f'=ROUNDDOWN(C{r}*E{r},0)', h='right', fmt=FMT_AMT)
        put(ws, f'G{r}', basis, size=9, wrap=True)
        ws.row_dimensions[r].height = max(22, row_h(est_lines(basis, 31, 9), 9, pad=5))
        r += 1
    tot = r
    put(ws, f'B{tot}', '合計（Ｚ ／ Ａ）', bold=True, h='center')
    merge(ws, f'C{tot}:D{tot}')
    put(ws, f'E{tot}', f'=SUM(E{hdr2 + 1}:E{tot - 1})', bold=True, h='right', fmt='0.000%_ ')
    put(ws, f'F{tot}', f'=SUM(F{hdr2 + 1}:F{tot - 1})', bold=True, h='right', fmt=FMT_AMT)
    ws.row_dimensions[tot].height = 24
    grid(ws, f'B{hdr2}:G{tot}', outer='medium', vert='thin', horiz='thin')
    hline(ws, f'B{tot}:G{tot}', 'top', 'double')
    r = tot + 2
    for t in ['※ 法定福利費相当額（Ａ）は見積金額（工事費）に含まれています。別途加算するものではありません。',
              '※ 料率は見積作成時点（令和8年度）の公表料率によります。料率が改定された場合は改定後の料率により精算させていただきます。']:
        merge(ws, f'B{r}:G{r}')
        put(ws, f'B{r}', t, size=10, wrap=True)
        ws.row_dimensions[r].height = row_h(est_lines(t, 104, 10), 10)
        r += 1
    page_setup(ws, f'B1:G{r - 1}', fit_h=1)
    return dict(rate_cell=f'$E${tot}', amount_cell=f'$F${tot}', y_cell=f'$C${ry}')


# =====================================================================
# 表紙(工事)
# =====================================================================
def build_cover(ws, job, det, wf, cl):
    FIRST, SUB, DISC, TOTAL = cl['first'], cl['sub'], cl['disc'], cl['total']
    set_widths(ws, dict(A=6, B=8, C=3, D=15, E=9, F=9, G=7, H=5, I=7, J=7, K=14, L=12,
                        M=46, N=13, O=13, P=13, Q=10))
    merge(ws, 'A1:L2')
    put(ws, 'A1', '御　見　積　書', size=20, bold=True, h='center')
    ws.row_dimensions[1].height = 22
    ws.row_dimensions[2].height = 22
    put(ws, 'J3', 'No.', h='right')
    merge(ws, 'K3:L3')
    put(ws, 'K3', None, h='center')
    hline(ws, 'K3:L3', 'bottom', 'thin')
    merge(ws, 'J4:L4')
    put(ws, 'J4', job.EST_DATE, h='right', fmt=R.DATE_FMT)
    merge(ws, 'A5:F5')
    put(ws, 'A5', job.CLIENT, size=14, h='center')
    hline(ws, 'A5:F5', 'bottom', 'thin')
    put(ws, 'G5', '御中', size=12, h='left')
    ws.row_dimensions[5].height = 26
    ws.row_dimensions[6].height = 12
    put(ws, 'A7', '金額', size=12, h='center')
    merge(ws, 'B7:E7')
    put(ws, 'B7', f'="¥"&TEXT(K{TOTAL},"#,##0")&"-"', size=14, bold=True, h='center')
    hline(ws, 'B7:E7', 'bottom', 'double')
    put(ws, 'F7', '（税別）', h='left')
    put(ws, 'H7', R.COMPANY, size=12, bold=True)
    ws.row_dimensions[7].height = 28
    ws.row_dimensions[8].height = 10
    merge(ws, 'A9:F9')
    put(ws, 'A9', 'つぎのとおりお見積いたしました。', h='left')
    put(ws, 'H9', R.REP)
    put(ws, 'A10', 'なにとぞご用命くださいますようお願い申しあげます。')
    put(ws, 'H11', R.ZIP)
    merge(ws, 'A12:B12')
    put(ws, 'A12', '件名', h='distributed', indent=1)
    merge(ws, 'D12:G12')
    put(ws, 'D12', job.PROJECT_1, shrink=True)
    put(ws, 'H12', R.ADDR)
    merge(ws, 'D13:G13')
    put(ws, 'D13', job.PROJECT_2, shrink=True)
    put(ws, 'H13', R.TEL)
    put(ws, 'K13', R.FAX)
    merge(ws, 'A14:B14')
    put(ws, 'A14', '施工場所', h='distributed', indent=1)
    merge(ws, 'D14:G14')
    put(ws, 'D14', job.SITE, size=10, shrink=True)
    put(ws, 'H14', '担当者')
    put(ws, 'J14', R.STAFF)
    for r, lab, val in [(15, '工事期限', job.TERM), (16, '支払条件', job.PAYMENT), (17, '本書有効期間', job.VALIDITY)]:
        merge(ws, f'A{r}:B{r}')
        put(ws, f'A{r}', lab, h='distributed', indent=1, shrink=True)
        merge(ws, f'D{r}:L{r}')
        put(ws, f'D{r}', val, shrink=True)
    for r in range(9, 18):
        ws.row_dimensions[r].height = 19
    ws.row_dimensions[18].height = 12
    merge(ws, 'B19:F19')
    merge(ws, 'I19:J19')
    for ref, txt, sz in [('A19', '項', 10), ('B19', '品　　　名', 11), ('G19', '数 量', 11), ('H19', '単位', 10),
                         ('I19', '単  価', 11), ('K19', '金　　額', 11), ('L19', '備　考', 11)]:
        put(ws, ref, txt, size=sz, h='center', fill=FILL_G, shrink=True)
    fill_range(ws, 'A19:L19', FILL_G)
    merge(ws, 'B20:L20')
    put(ws, 'B20', '【' + job.PROJECT_FULL + '】　' + work_title(job), indent=1, shrink=True)
    merge(ws, 'B21:L21')
    put(ws, 'B21', f'=IF(INCLUDE_TBD=1,"{job.COVER_SCOPE["include"]}","{job.COVER_SCOPE["exclude"]}")', size=10, indent=1, shrink=True)
    for r in range(22, TOTAL + 1):
        merge(ws, f'B{r}:F{r}')
        merge(ws, f'I{r}:J{r}')
    oh = det['oh']
    spec = []
    for i, (kind, sec) in enumerate(cl['rows']):
        r = FIRST + i
        if kind == 'equip':
            eq_name, eq_show = job.EQUIP_ROW
            spec.append((r, CIRC[i], eq_name, 0, eq_show, f'#,##0_ ;[Red]-#,##0_ ;"{eq_show}"_ '))
        elif kind == 'sec':
            spec.append((r, CIRC[i], sec['title'], f"={q(SH_DETAIL)}!H{det['sec'][sec['key']]['total']}",
                         '内訳書 ' + sec['no'], FMT_SIGNED))
        else:
            spec.append((r, CIRC[i], '工事諸経費', f"={q(SH_DETAIL)}!H{oh}", '内訳書 ' + getattr(job, 'OH_NO', '８'),
                         FMT_SIGNED))
    for r, no, name, fk, remark, fmt in spec:
        put(ws, f'A{r}', no, h='center')
        put(ws, f'B{r}', name, indent=1, shrink=True)
        put(ws, f'G{r}', 1, h='right', fmt=FMT_QTY)
        put(ws, f'H{r}', '式', h='center')
        put(ws, f'I{r}', f'=K{r}', h='right', fmt=fmt)
        put(ws, f'K{r}', fk, h='right', fmt=fmt)
        put(ws, f'L{r}', remark, size=10, h='center', shrink=True)
    put(ws, f'B{SUB}', '　小　　計', indent=1)
    put(ws, f'K{SUB}', f'=SUM(K{FIRST}:K{SUB - 1})', h='right', fmt=FMT_SIGNED)
    put(ws, f'A{DISC}', CIRC[len(cl['rows'])], h='center')
    put(ws, f'B{DISC}', '出精値引', indent=1)
    put(ws, f'K{DISC}', f'=IF(SUBMIT<K{SUB},SUBMIT-K{SUB},0)', h='right', fmt=FMT_SIGNED)
    put(ws, f'B{TOTAL}', '　合　　計', bold=True, indent=1)
    put(ws, f'K{TOTAL}', f'=K{SUB}+K{DISC}', bold=True, h='right', fmt=FMT_AMT)
    put(ws, f'L{TOTAL}', '（税別）', size=10, h='center')
    for r in range(19, TOTAL + 1):
        ws.row_dimensions[r].height = 24
    ws.row_dimensions[20].height = 22
    ws.row_dimensions[21].height = 20
    grid(ws, f'A19:L{TOTAL}', outer='medium', vert='thin', horiz='thin')
    hline(ws, f'A{SUB}:L{SUB}', 'top', 'medium')
    hline(ws, f'A{TOTAL}:L{TOTAL}', 'top', 'double')
    for c in range(1, 13):
        set_border(ws.cell(TOTAL, c), bottom='medium')
    r = TOTAL + 1
    ws.row_dimensions[r].height = 8
    r += 1
    merge(ws, f'A{r}:L{r}')
    put(ws, f'A{r}', '上記の工事費には、以下の法定福利費相当額を含んでいます。', size=10, indent=1)
    r += 1
    merge(ws, f'A{r}:L{r}')
    put(ws, f'A{r}', f'="【法定福利費相当額（Ａ）】"&TEXT({q(SH_WF)}!{wf["amount_cell"]},"#,##0")&"円　　Ａ＝Ｙ×Ｚ　Ｙ：労務費相当額 "'
        f'&TEXT({q(SH_WF)}!{wf["y_cell"]},"#,##0")&"円　Ｚ：法定福利費事業者負担率（合計値："&TEXT({q(SH_WF)}!{wf["rate_cell"]}*100,"0.000")&"／100）"',
        size=10, indent=1, shrink=True)
    grid(ws, f'A{r - 1}:L{r}', outer='thin', vert=None, horiz=None)
    r += 2
    for t in job.NOTES['cover']:
        put(ws, f'A{r}', '※', size=9.5, h='center', v='top')
        merge(ws, f'B{r}:L{r}')
        put(ws, f'B{r}', t, size=9.5, wrap=True, v='top')
        ws.row_dimensions[r].height = row_h(est_lines(t, 96, 9.5), 9.5, pad=3)
        r += 1
    last_note = r - 1
    # ---- 内部エリア
    put(ws, 'M3', '【設定・内部用】（印刷範囲外）黄セル＝入力', bold=True, color=C_GRAYTXT, fill=FILL_G)
    put(ws, 'N3', None, fill=FILL_G)
    for rr, lab, val, fmt, is_input in [
            (4, '目標粗利率', R.TARGET_MARGIN, '0%', True),
            (5, '提出額の丸め単位（切上げ）', R.ROUND_UNIT, '#,##0', True),
            (6, '労務原価日額（円/人工、当社）', R.LABOR_COST, '#,##0', True),
            (7, '材料原価率（エスト材料単価＝建設物価相当に対する比率）', R.MAT_RATIO, '0.00', True),
            (8, 'エスト労務単価 配管工費（円/人）', R.EST_LABOR_PIPE, '#,##0', True),
            (9, 'エスト労務単価 電線材料施工費（円/人）', R.EST_LABOR_ELEC, '#,##0', True),
            (10, '法定福利費率（事業主負担）', '=WELFARE_RATE', '0.000%', False),
            (11, '要見積（概算）を提出額に含める（1=含める／0=別途見積）', R.INCLUDE_TBD, '0', True),
            (12, '目標提出額（税別・空欄なら粗利率から自動）', None, '#,##0', True),
            (13, '提出額（自動計算）', '=IF(TARGET="",ROUNDUP(COST_TOTAL/(1-TARGET_MARGIN)/ROUND_UNIT,0)*ROUND_UNIT,TARGET)', '#,##0', False)]:
        put(ws, f'M{rr}', lab, size=10, color=C_GRAYTXT, fill=FILL_G, shrink=True)
        put(ws, f'N{rr}', val, h='right', fmt=fmt, fill=(FILL_Y if is_input else None), bold=(rr == 13))
    grid(ws, 'M3:N13', outer='thin', vert='thin', horiz='thin')
    put(ws, 'M14', '※ 提出額＝原価合計÷（1－目標粗利率）を丸め単位で切上げ。小計（エスト単価の定価ベース）との差を出精値引に計上', size=9, color=C_GRAYTXT)
    put(ws, 'M15', '※ 原価＝材料（エスト材料単価×材料原価率＝建設物価相当）＋人工×労務原価日額＋法定福利費（事業主負担）', size=9, color=C_GRAYTXT)
    for ref, txt in [('M19', '内部集計'), ('N19', '定価ベース'), ('O19', '原価'), ('P19', '粗利'), ('Q19', '粗利率')]:
        put(ws, ref, txt, size=10, h='center', color=C_GRAYTXT, fill=FILL_G, shrink=True)
    gt = det['gt']
    internal = []
    for i, (kind, sec) in enumerate(cl['rows']):
        rr = FIRST + i
        if kind == 'sec':
            internal.append((rr, CIRC[i] + sec['title'], f'=K{rr}', f"={q(SH_DETAIL)}!P{det['sec'][sec['key']]['total']}"))
        elif kind == 'oh':
            internal.append((rr, CIRC[i] + '工事諸経費（原価なし）', f'=K{rr}', 0))
    first_int = internal[0][0]
    internal += [(SUB, '法定福利費（事業主負担・原価）', None, f"=ROUND({q(SH_DETAIL)}!Q{gt}*WELFARE_RATE,0)"),
                 (DISC, '原価合計', f'=K{SUB}', f'=SUM(O{first_int}:O{SUB})')]
    for rr, lab, nv, ov in internal:
        put(ws, f'M{rr}', lab, size=10, bold=(rr == DISC), shrink=True)
        put(ws, f'N{rr}', nv, h='right', fmt=FMT_SIGNED)
        put(ws, f'O{rr}', ov, h='right', fmt=FMT_SIGNED, bold=(rr == DISC))
        if nv is not None and rr != DISC:
            put(ws, f'P{rr}', f'=N{rr}-O{rr}', h='right', fmt=FMT_SIGNED)
            put(ws, f'Q{rr}', f'=IF(N{rr}=0,"",P{rr}/N{rr})', h='right', fmt='0.0%')
    put(ws, f'M{TOTAL}', '提出額（合計）／粗利／粗利率', size=10, bold=True, shrink=True)
    put(ws, f'N{TOTAL}', f'=K{TOTAL}', bold=True, h='right', fmt=FMT_AMT)
    put(ws, f'O{TOTAL}', f'=O{DISC}', h='right', fmt=FMT_AMT)
    put(ws, f'P{TOTAL}', f'=N{TOTAL}-O{TOTAL}', bold=True, h='right', fmt=FMT_SIGNED)
    put(ws, f'Q{TOTAL}', f'=IF(N{TOTAL}=0,"",P{TOTAL}/N{TOTAL})', bold=True, h='right', fmt='0.0%')
    grid(ws, f'M19:Q{TOTAL}', outer='thin', vert='thin', horiz='hair')
    # アラート
    tbd = [ln for ln in det['lines'] if ln['kind'] == 'tbd']
    own = [ln for ln in det['lines'] if ln['kind'] == 'item' and ln['src'] in ALERT_OF]
    A0 = TOTAL + 3
    put(ws, f'M{A0}', '⚠ アラート（提出前に確認）', bold=True, color=C_RED, fill=FILL_RED)
    for c in 'NOPQ':
        put(ws, f'{c}{A0}', None, fill=FILL_RED)
    tbd_cost = '+'.join(f"{q(SH_DETAIL)}!E{ln['_row']}*{q(SH_DETAIL)}!V{ln['_row']}" for ln in tbd) or '0'
    alerts = [
        (f'要見積 {len(tbd)}件：' + '・'.join(sorted(set(ln['name'] for ln in tbd))) + '（概算原価の合計）', f'={tbd_cost}',
         '=IF(INCLUDE_TBD=1,"概算を提出額に含めています","別途見積（提出額に含めていません）")'),
        ('要見積を概算で含めた場合の提出額（目安）',
         f'=IF(INCLUDE_TBD=1,SUBMIT,ROUNDUP((COST_TOTAL+N{A0 + 1})/(1-TARGET_MARGIN)/ROUND_UNIT,0)*ROUND_UNIT)', 'N11 を 1 にすると反映'),
        (f'単価要確認 {len(own)}件：' + '・'.join(ln['name'] for ln in own), None, 'エストに単価なし。当社単価・推定単価で仮計上（単価出典シート）'),
        ('定価ベース（エスト単価）で目標粗利に届くか', None,
         f'=IF(SUBMIT>K{SUB},"⚠ 値引きなしでも粗利率 "&TEXT((K{SUB}-COST_TOTAL)/K{SUB},"0.0%")&"（目標未達）","OK（出精値引で調整）")'),
    ] + [(t, None, None) for t in getattr(job, 'ALERT_NOTES', [])]
    for i, (lab, nv, ov) in enumerate(alerts):
        rr = A0 + 1 + i
        put(ws, f'M{rr}', lab, size=10, shrink=True, color=C_RED)
        put(ws, f'N{rr}', nv, h='right', fmt=FMT_AMT)
        merge(ws, f'O{rr}:Q{rr}')
        put(ws, f'O{rr}', ov, size=10, shrink=True)
    A1 = A0 + len(alerts)
    grid(ws, f'M{A0}:Q{A1}', outer='thin', vert='thin', horiz='hair')
    R0 = A1 + 2
    put(ws, f'M{R0}', '参考（感度）', bold=True, color=C_GRAYTXT, fill=FILL_G)
    for c, t in zip('NOPQ', ['金額', '原価', '', '粗利率']):
        put(ws, f'{c}{R0}', t or None, size=10, h='center', color=C_GRAYTXT, fill=FILL_G)
    labor_md = f"{q(SH_DETAIL)}!M{gt}"
    put(ws, f'M{R0 + 1}', '労務を公共工事設計労務単価（広島 配管工）で計算した場合', size=10, shrink=True)
    put(ws, f'N{R0 + 1}', f'=K{TOTAL}', h='right', fmt=FMT_AMT)
    put(ws, f'O{R0 + 1}', f'=COST_TOTAL+{labor_md}*(N{R0 + 4}-LABOR_COST)*(1+WELFARE_RATE)', h='right', fmt=FMT_AMT)
    put(ws, f'Q{R0 + 1}', f'=(N{R0 + 1}-O{R0 + 1})/N{R0 + 1}', h='right', fmt='0.0%')
    put(ws, f'M{R0 + 2}', '出精値引率（変更増減の算定に使用：増減額×(1－この率)）', size=10, shrink=True)
    put(ws, f'N{R0 + 2}', f'=IF(K{SUB}=0,"",-K{DISC}/K{SUB})', h='right', fmt='0.0%')
    put(ws, f'M{R0 + 3}', '人工合計（内訳書）', size=10)
    put(ws, f'N{R0 + 3}', f'={labor_md}', h='right', fmt='0.00')
    put(ws, f'M{R0 + 4}', '公共工事設計労務単価 広島 配管工（R8.3）', size=10, color=C_GRAYTXT)
    put(ws, f'N{R0 + 4}', R.PUBLIC_LABOR_R8_HIROSHIMA['配管工'], h='right', fmt=FMT_AMT, fill=FILL_Y)
    grid(ws, f'M{R0}:Q{R0 + 4}', outer='thin', vert='thin', horiz='hair')
    T0 = R0 + 6
    put(ws, f'M{T0}', '粗利率別 参考提出額（税別・万円未満切上げ）', size=10, color=C_GRAYTXT, fill=FILL_G, shrink=True)
    put(ws, f'N{T0}', '参考提出額', size=10, h='center', color=C_GRAYTXT, fill=FILL_G)
    put(ws, f'O{T0}', '出精値引率', size=10, h='center', color=C_GRAYTXT, fill=FILL_G)
    for i, rate in enumerate([0.10, 0.15, 0.20, 0.25, 0.30]):
        rr = T0 + 1 + i
        put(ws, f'M{rr}', rate, h='center', fmt='"粗利率 "0%')
        put(ws, f'N{rr}', f'=ROUNDUP(COST_TOTAL/(1-M{rr}),-4)', h='right', fmt=FMT_AMT)
        put(ws, f'O{rr}', f'=IF(K{SUB}=0,"",1-N{rr}/K{SUB})', h='right', fmt='0.0%')
    grid(ws, f'M{T0}:O{T0 + 5}', outer='thin', vert='thin', horiz='thin')
    page_setup(ws, f'A1:L{last_note}', fit_h=1)
    return dict(sub=SUB, disc=DISC, total=TOTAL, cost_cell=f'$O${DISC}')


# =====================================================================
# 御見積条件
# =====================================================================
def build_conditions(ws, job, det):
    W = dict(A=5, B=10, C=25, D=20, E=16, F=19)
    set_widths(ws, W)
    BF = W['B'] + W['C'] + W['D'] + W['E'] + W['F']
    put(ws, 'A1', '御見積条件・注意事項', size=14, bold=True)
    ws.row_dimensions[1].height = 24
    put(ws, 'A2', job.PROJECT_FULL + '　御見積書 添付', size=10)
    r = 4
    put(ws, f'A{r}', '【1】御見積条件・注意事項', size=12, bold=True)
    ws.row_dimensions[r].height = 22
    r += 1
    for i, t in enumerate(job.NOTES['conditions'], 1):
        put(ws, f'A{r}', f'{i}.', h='right', v='top')
        merge(ws, f'B{r}:F{r}')
        put(ws, f'B{r}', t, wrap=True, v='top')
        ws.row_dimensions[r].height = row_h(est_lines(t, BF, 11), 11, pad=8)
        r += 1
    r += 1
    ws.row_breaks.append(Break(id=r - 1))
    put(ws, f'A{r}', '【2】想定数量（推測）項目の前提', size=12, bold=True)
    ws.row_dimensions[r].height = 22
    r += 1
    merge(ws, f'A{r}:F{r}')
    put(ws, f'A{r}', '図面に数量・経路の記載がなく当社にて想定した項目です。「図面実測＋想定」は平面延長を図面計測し、立上り等を想定加算したものです。実数量確定後に精算させていただきます。', size=10, wrap=True)
    ws.row_dimensions[r].height = row_h(2, 10, pad=4)
    r += 1
    hdr = r
    for ref, txt in [('A', 'No.'), ('B', '内訳書'), ('C', '項　　目'), ('D', '前提・想定内容'), ('F', '数量（区分）')]:
        put(ws, f'{ref}{r}', txt, size=10, h='center', fill=FILL_G)
    merge(ws, f'D{r}:E{r}')
    fill_range(ws, f'A{r}:F{r}', FILL_G)
    r += 1
    k = 0
    for ln in det['lines']:
        if ln['kind'] != 'item' or ln.get('basis') not in ('推測', '実測+推測') or not ln.get('qmemo'):
            continue
        k += 1
        item = ln['name'] + (' ' + ln['spec'] if ln.get('spec') else '')
        qtxt = f"{ln['qty']:,} {ln['unit']}（{'想定' if ln['basis'] == '推測' else '実測＋想定'}）"
        put(ws, f'A{r}', k, size=10, h='center')
        put(ws, f'B{r}', {s['key']: s['no'] for s in det['secs']}[ln['_sec']], size=10, h='center')
        put(ws, f'C{r}', item, size=10, wrap=True, indent=1)
        merge(ws, f'D{r}:E{r}')
        put(ws, f'D{r}', ln['qmemo'], size=10, wrap=True, indent=1)
        put(ws, f'F{r}', qtxt, size=10, wrap=True, h='center')
        nl = max(est_lines(item, W['C'] - 3, 10), est_lines(ln['qmemo'], W['D'] + W['E'] - 3, 10), est_lines(qtxt, W['F'], 10))
        ws.row_dimensions[r].height = max(18, row_h(nl, 10, pad=4))
        r += 1
    grid(ws, f'A{hdr}:F{r - 1}', outer='medium', vert='thin', horiz='thin')
    r += 1
    put(ws, f'A{r}', '【3】図面読み取りに関する確認事項', size=12, bold=True)
    ws.row_dimensions[r].height = 22
    r += 1
    for i, t in enumerate(job.NOTES['checks'], 1):
        put(ws, f'A{r}', f'{i}.', h='right', v='top')
        merge(ws, f'B{r}:F{r}')
        put(ws, f'B{r}', t, wrap=True, v='top')
        ws.row_dimensions[r].height = row_h(est_lines(t, BF, 11), 11, pad=8)
        r += 1
    page_setup(ws, f'A1:F{r - 1}', fit_h=0)
    return dict(last=r - 1, n_assump=k)


# =====================================================================
# 内部シート
# =====================================================================
def build_sources(ws, job, det):
    set_widths(ws, dict(A=12, B=30, C=34, D=7, E=6, F=10, G=12, H=16, I=14, J=60))
    put(ws, 'A1', '単価出典（内部用）　エスト見積 → エストマスタ → 推定 → 当社単価 → 要見積 の順で単価を決定', size=11, bold=True)
    hdr = ['区分', '品名', '規格', '数量', '単位', '単価', '金額', '単価区分', 'アラート', '単価の参照元']
    for i, h in enumerate(hdr):
        put(ws, f'{chr(65 + i)}3', h, size=10, bold=True, h='center', fill=FILL_G)
    r = 4
    for ln in det['lines']:
        if ln['kind'] not in ('item', 'tbd'):
            continue
        dr = ln['_row']
        put(ws, f'A{r}', {s['key']: s['title'] for s in det['secs']}[ln['_sec']], size=10)
        put(ws, f'B{r}', ln['name'], size=10)
        put(ws, f'C{r}', ln.get('spec') or None, size=9, shrink=True)
        put(ws, f'D{r}', f"={q(SH_DETAIL)}!E{dr}", size=10, h='right', fmt='#,##0.##')
        put(ws, f'E{r}', ln['unit'], size=10, h='center')
        put(ws, f'F{r}', f"={q(SH_DETAIL)}!G{dr}", size=10, h='right', fmt=FMT_AMT)
        put(ws, f'G{r}', f"={q(SH_DETAIL)}!H{dr}", size=10, h='right', fmt=FMT_TBD)
        src = R.SRC_TBD if ln['kind'] == 'tbd' else ln['src']
        alert = '要見積' if ln['kind'] == 'tbd' else ALERT_OF.get(src)
        put(ws, f'H{r}', src, size=9, h='center', shrink=True)
        put(ws, f'I{r}', alert, size=9, h='center', fill=(FILL_RED if alert == '要見積' else (FILL_AMB if alert else None)))
        put(ws, f'J{r}', ('単価情報なし。' + (ln.get('memo') or '')) if ln['kind'] == 'tbd' else ln.get('ref'), size=9)
        r += 1
    grid(ws, f'A3:J{r - 1}', outer='thin', vert='thin', horiz='hair')
    r += 1
    put(ws, f'A{r}', 'エストの計算規則', size=10, bold=True)
    for t in ['冷媒管: 継手類30%・消耗品15%・支持金物40%（管金額比）。塩ビ管: 20%・10%・25%',
              '配管工費＝国交省 公共建築工事標準単価積算基準 R8 の歩掛×33,750円/人。電線材料施工費＝0.017人/m×33,600円/人',
              '部門内経費: 空調 3%→4%→8%、配管 3%→5%→10% を順に積上げ（100円未満切上げ）、計を1,000円未満切上げ',
              '工事諸経費＝（空調＋配管）×5%。提出額は出精値引で調整。法定福利費は工事費に内含して表示',
              '当社単価＝原価（材料＋人工×労務原価日額）×1.25 を100円単位で切上げ'] + list(getattr(job, 'SRC_RULE_NOTES', [])):
        r += 1
        put(ws, f'A{r}', '・' + t, size=9)
    page_setup(ws, f'A1:J{r}', fit_h=0)
    ws.page_setup.orientation = 'landscape'


def build_master(ws):
    set_widths(ws, dict(A=30, B=34, C=80))
    put(ws, 'A1', 'エスト単価マスタ（参考・抜粋）　エスト見積 6330560-1 に単価がない種別はこのマスタの値を採用', size=11, bold=True)
    for i, h in enumerate(['区分', '種別', '単価（円）']):
        put(ws, f'{chr(65 + i)}3', h, size=10, bold=True, h='center', fill=FILL_G)
    r = 4
    for a, b, c in R.EST_MASTER:
        put(ws, f'A{r}', a, size=10)
        put(ws, f'B{r}', b, size=10)
        put(ws, f'C{r}', c, size=10)
        r += 1
    grid(ws, f'A3:C{r - 1}', outer='thin', vert='thin', horiz='hair')
    page_setup(ws, f'A1:C{r}', fit_h=0)
    ws.page_setup.orientation = 'landscape'


def build_basis(ws, job):
    set_widths(ws, dict(A=36, B=14, C=12, D=12, E=12, F=70))
    head = getattr(job, 'BASIS_NOTE', ('数量拾い根拠（内部用）', ''))
    put(ws, 'A1', head[0], size=10, bold=True)
    put(ws, 'A2', head[1], size=9)
    r = 4
    for title, headers, rows in job.BASIS_TABLES:
        put(ws, f'A{r}', title, bold=True, fill=FILL_L)
        r += 1
        for i, h in enumerate(headers):
            put(ws, f'{chr(65 + i)}{r}', h, size=10, bold=True, h='center', fill=FILL_G)
        hdr = r
        r += 1
        for row in rows:
            for i, v in enumerate(row):
                put(ws, f'{chr(65 + i)}{r}', v, size=10, h=('left' if isinstance(v, str) else 'right'),
                    fmt=('#,##0.00' if isinstance(v, float) else None))
            r += 1
        grid(ws, f'A{hdr}:{chr(64 + len(headers))}{r - 1}', outer='thin', vert='thin', horiz='hair')
        r += 1
    page_setup(ws, f'A1:F{r}', fit_h=0)
    ws.page_setup.orientation = 'landscape'


def build_cost(ws, job):
    set_widths(ws, dict(A=40, B=16, C=90))
    put(ws, 'A1', '原価根拠（内部用）　材料＝建設物価相当、労務＝人工×労務原価日額、法定福利費＝労務原価×事業主負担率', size=11, bold=True)
    r = 3
    put(ws, f'A{r}', '出典', bold=True, fill=FILL_L)
    r += 1
    for a, b in R.SOURCES:
        put(ws, f'A{r}', a, size=10, bold=True)
        merge(ws, f'B{r}:C{r}')
        put(ws, f'B{r}', b, size=9, wrap=True)
        ws.row_dimensions[r].height = 30
        r += 1
    r += 1
    put(ws, f'A{r}', '公共工事設計労務単価（令和8年3月適用・広島県、円/日）', bold=True, fill=FILL_L)
    r += 1
    for k, v in R.PUBLIC_LABOR_R8_HIROSHIMA.items():
        put(ws, f'A{r}', k, size=10)
        put(ws, f'B{r}', v, size=10, h='right', fmt=FMT_AMT)
        r += 1
    r += 1
    put(ws, f'A{r}', '国交省 公共建築工事標準単価積算基準（令和8年改定）配管工 歩掛（人/m）', bold=True, fill=FILL_L)
    r += 1
    for k, v in R.BUG.items():
        put(ws, f'A{r}', k, size=10)
        put(ws, f'B{r}', v, size=10, h='right', fmt='0.000')
        r += 1
    r += 1
    put(ws, f'A{r}', '原価の前提（内訳書 L・M 列の黄セル）', bold=True, fill=FILL_L)
    for t in job.COST_NOTES:
        r += 1
        merge(ws, f'A{r}:C{r}')
        put(ws, f'A{r}', '・' + t, size=9, wrap=True)
        ws.row_dimensions[r].height = 28
    page_setup(ws, f'A1:C{r}', fit_h=0)
    ws.page_setup.orientation = 'landscape'


def build_book(job, out_dir):
    wb = Workbook()
    base = fnt(11)
    wb._fonts = IndexedList([base])
    wb._named_styles['Normal'].font = base
    wb.properties.creator = R.COMPANY
    wb.properties.title = '御見積書 ' + job.PROJECT_FULL
    ws_cover = wb.active
    ws_cover.title = SH_COVER
    ws_det = wb.create_sheet(SH_DETAIL)
    ws_wf = wb.create_sheet(SH_WF)
    ws_cd = wb.create_sheet(SH_COND)
    ws_src = wb.create_sheet(SH_SRC)
    ws_mst = wb.create_sheet(SH_MST)
    ws_bs = wb.create_sheet(SH_BASIS)
    ws_cost = wb.create_sheet(SH_COST)
    det = build_detail(ws_det, job)
    cl = cover_rows(job, det['secs'])
    wf = build_welfare(ws_wf, job, det, cl)
    cv = build_cover(ws_cover, job, det, wf, cl)
    cd = build_conditions(ws_cd, job, det)
    build_sources(ws_src, job, det)
    build_master(ws_mst)
    build_basis(ws_bs, job)
    build_cost(ws_cost, job)
    names = {
        'TARGET_MARGIN': f"{q(SH_COVER)}!$N$4", 'ROUND_UNIT': f"{q(SH_COVER)}!$N$5",
        'LABOR_COST': f"{q(SH_COVER)}!$N$6", 'MAT_RATIO': f"{q(SH_COVER)}!$N$7",
        'EST_LABOR_PIPE': f"{q(SH_COVER)}!$N$8", 'EST_LABOR_ELEC': f"{q(SH_COVER)}!$N$9",
        'WELFARE_RATE': f"{q(SH_WF)}!{wf['rate_cell']}", 'INCLUDE_TBD': f"{q(SH_COVER)}!$N$11",
        'TARGET': f"{q(SH_COVER)}!$N$12", 'SUBMIT': f"{q(SH_COVER)}!$N$13",
        'COST_TOTAL': f"{q(SH_COVER)}!{cv['cost_cell']}",
    }
    for k, ref in names.items():
        wb.defined_names[k] = DefinedName(k, attr_text=ref)
    wb.active = 0
    for ws in wb.worksheets:
        ws.sheet_view.tabSelected = (ws.title == SH_COVER)
    wb.calculation.fullCalcOnLoad = True
    os.makedirs(out_dir, exist_ok=True)
    out = os.path.join(out_dir, job.FILE)
    wb.save(out)
    return dict(out=out, detail=dict(sec=det['sec'], oh=det['oh'], gt=det['gt']), conditions=cd,
                cover=dict(cl, rows=[(k, s['key'] if s else None) for k, s in cl['rows']]))


# =====================================================================
# 検算（ワークブックの数式と同じ計算）
# =====================================================================
def _ru(x, n):
    m = 10 ** n
    return math.ceil(x * m - 1e-9) / m


def simulate(job, include_tbd=R.INCLUDE_TBD, labor_cost=R.LABOR_COST, mat_ratio=R.MAT_RATIO,
             target_margin=R.TARGET_MARGIN):
    tot, cost_all, labor_all, md_all = {}, 0.0, 0.0, 0.0
    for sec in sections(job):
        H, C, QTY = {}, {}, {}
        base = cost = 0.0
        for ln in sec['lines']:
            k = ln['kind']
            if k == 'item':
                h = round(ln['qty'] * ln['price'])
                mat = ln['mat']
                if isinstance(mat, tuple):
                    mat = round(ln['price'] * mat_ratio * mat[1])
                lab = round(ln.get('md', 0) * labor_cost)
                c = ln['qty'] * (mat + lab)
                labor_all += ln['qty'] * lab
                md_all += ln['qty'] * ln.get('md', 0)
            elif k == 'tbd':
                price = _ru(ln['prov'] / (1 - target_margin), -3)
                h = round(ln['qty'] * price) if include_tbd else 0
                c = ln['qty'] * ln['prov'] if include_tbd else 0
            elif k == 'pct':
                h = round(sum(H[b] for b in ln['base']) * ln['rate'])
                c = round(sum(C[b] for b in ln['base']) * ln['rate'])
            elif k == 'labor':
                md = sum(QTY[b] * next(x for x in sec['lines'] if x.get('key') == b)['bug'] for b in ln['base'])
                rate = R.EST_LABOR_PIPE if ln['rate'] == 'EST_LABOR_PIPE' else R.EST_LABOR_ELEC
                h = _ru(md * rate, -1)
                c = md * round(labor_cost)
                labor_all += c
                md_all += md
            else:
                continue
            H[ln['key']], C[ln['key']], QTY[ln['key']] = h, c, ln.get('qty', 1)
            base += h
            cost += c
        if any(ln['kind'] == 'exp' for ln in sec['lines']):
            rates = [rt for _, rt in R.EST_EXP[sec.get('exp_key', sec['key'])]]
            e1 = _ru(base * rates[0], -2)
            e2 = _ru((base + e1) * rates[1], -2)
            s3 = base + e1 + e2
            total = _ru(s3 + _ru(s3 * rates[2], -2), -3)
            exp = (e1, e2, total - s3)
        else:
            total, exp = base, (0, 0, 0)
        tot[sec['key']] = dict(base=base, exp=exp, total=total, cost=cost, oh=sec.get('oh', True))
        cost_all += cost
    oh = _ru(sum(t['total'] for t in tot.values() if t['oh']) * R.EST_OVERHEAD, -2)
    sub = sum(t['total'] for t in tot.values()) + oh
    cost_total = cost_all + round(labor_all * sum(r for _, r, _ in R.WELFARE_RATES))
    submit = math.ceil(cost_total / (1 - target_margin) / R.ROUND_UNIT - 1e-9) * R.ROUND_UNIT
    disc = submit - sub if submit < sub else 0
    final = sub + disc
    return dict(**{k.lower(): t['total'] for k, t in tot.items()}, secs={k: t['total'] for k, t in tot.items()},
                oh=oh, sub=sub, disc=disc, final=final,
                cost=cost_total, md=md_all, margin=(final - cost_total) / final if final else 0)
