#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""提出用 PDF の作成（シート名に「(内部)」「(参考)」を含むシートを除いて LibreOffice で PDF 化）
使い方: python3 make_pdf.py <御見積書.xlsx> <出力ディレクトリ>"""
import os
import shutil
import subprocess
import sys
import tempfile

from openpyxl import load_workbook

src, out_dir = sys.argv[1], sys.argv[2]
tmpd = tempfile.mkdtemp()
wb = load_workbook(src)
for name in list(wb.sheetnames):
    if '(内部)' in name or '(参考)' in name:
        wb.remove(wb[name])
base = os.path.splitext(os.path.basename(src))[0]
tmp_xlsx = os.path.join(tmpd, base + '.xlsx')
wb.save(tmp_xlsx)
profile = 'file://' + os.path.join(tmpd, 'lo_profile')
subprocess.run(['soffice', f'-env:UserInstallation={profile}', '--headless', '--convert-to', 'pdf',
                '--outdir', tmpd, tmp_xlsx], check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
os.makedirs(out_dir, exist_ok=True)
shutil.move(os.path.join(tmpd, base + '.pdf'), os.path.join(out_dir, base + '.pdf'))
shutil.rmtree(tmpd, ignore_errors=True)
print(os.path.join(out_dir, base + '.pdf'))
