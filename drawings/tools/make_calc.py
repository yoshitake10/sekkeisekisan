"""24時間換気計算書（シックハウス対策 令20条の8）Excel 作成

  python3 make_calc.py <out.xlsx>
  → LibreOffice で再計算・PDF化:  soffice --headless --convert-to pdf <out.xlsx>

床面積・天井高・換気回数（黄色セル）を変更すると、気積・必要有効換気量・判定が自動再計算される。
"""
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.page import PageMargins

import project_data as P

FN = 'ＭＳ ゴシック'
thin = Side(style='thin'); med = Side(style='medium')
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
FILL_IN = PatternFill('solid', fgColor='FFF2CC')
FILL_H = PatternFill('solid', fgColor='D9D9D9')
FILL_OK = PatternFill('solid', fgColor='E2EFDA')
COLS = 'ABCDEFGHI'
WID = [5, 22, 8, 10, 9, 10, 9, 12, 38]


def build(out):
    wb = Workbook(); ws = wb.active; ws.title = '24時間換気計算書'
    for c, w in zip(COLS, WID):
        ws.column_dimensions[c].width = w
    r = [1]

    def put(cell, v, bold=False, size=10, fill=None, al='left', fmt=None, border=True, wrap=False, color=None):
        c = ws[cell]; c.value = v
        c.font = Font(name=FN, size=size, bold=bold, color=color)
        c.alignment = Alignment(horizontal=al, vertical='center', wrap_text=wrap)
        if fill: c.fill = fill
        if fmt: c.number_format = fmt
        if border: c.border = BOX
        return c

    def row(vals, **kw):
        for col, v in zip(COLS, vals):
            if v is not None:
                put('%s%d' % (col, r[0]), v, **kw)
        r[0] += 1

    def merge_line(text, size=10, bold=False, height=None, border=False, fill=None, wrap=True, color=None):
        put('A%d' % r[0], text, bold=bold, size=size, border=border, fill=fill, wrap=wrap, color=color)
        ws.merge_cells('A%d:I%d' % (r[0], r[0]))
        if height:
            ws.row_dimensions[r[0]].height = height
        r[0] += 1

    def head(text):
        r[0] += 1
        merge_line(text, size=11, bold=True)

    # ---- 表題 ----------------------------------------------------------------
    merge_line('シックハウス対策に係る 24時間換気計算書', size=16, bold=True, height=28)
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
    merge_line('（建築基準法 第28条の2 第三号、同法施行令 第20条の8 第1項第一号イ(1)）', size=9)
    ws['A2'].alignment = Alignment(horizontal='center')

    head('1. 建築物の概要')
    info = [('工事名称', '宗教法人 安楽寺 様（本堂）新築工事'), ('建築場所', '（確認申請書による）'),
            ('主要用途', '寺院（本堂）'), ('構造・階数', '木造 平家建（意匠図による）'),
            ('建築面積／延べ面積', '220.743 ㎡ ／ 198.56 ㎡（意匠図 平面図 令和8年7月20日版）'),
            ('意匠設計', '松森設計事務所'), ('設備設計（換気）', '')]
    for k, v in info:
        put('A%d' % r[0], k, fill=FILL_H); ws.merge_cells('A%d:C%d' % (r[0], r[0]))
        put('D%d' % r[0], v, fill=FILL_IN if not v else None); ws.merge_cells('D%d:I%d' % (r[0], r[0]))
        for col in 'BCEFGHI':
            ws['%s%d' % (col, r[0])].border = BOX
        r[0] += 1

    head('2. 換気方式')
    for s in ['第3種機械換気設備（排気機：天井埋込形換気扇 常時運転／給気：外気導入グリル）。',
              '全館の居室及び廊下・ホール・便所等を、建具アンダーカット等の通気経路により一体的に換気する一の換気経路とする。',
              '換気回数 n：本堂部（外陣・内陣・護摩堂・広縁）＝住宅等の居室以外の居室 0.3回/h、',
              '　　　　　　庫裏側（和室・談話室及び廊下等）＝安全側に住宅等の居室の値 0.5回/h を採用。',
              '局所換気（焼香・護摩の排煙、台所等：EF-2〜EF-4）は必要有効換気量の算定に算入しない。']:
        merge_line(s, size=9.5, height=15)

    # ---- 3. 必要有効換気量 ---------------------------------------------------------
    head('3. 必要有効換気量 Vr ＝ n × A × h')
    row(['No', '室  名', '区分', '床面積\nA (㎡)', '天井高\nh (m)', '気積\nV (m³)', '換気回数\nn (回/h)',
         '必要有効換気量\nVr (m³/h)', '備  考'], bold=True, fill=FILL_H, al='center', wrap=True, size=9)
    ws.row_dimensions[r[0] - 1].height = 30
    r0 = r[0]
    for rm in P.ROOMS:
        i = r[0]
        row([int(rm['no']), rm['name'], rm['kind'], P.area(rm), rm['h'], '=ROUND(D%d*E%d,2)' % (i, i), rm['n'],
             '=ROUND(F%d*G%d,2)' % (i, i), rm['note']], size=9)
        for col, fmt in (('D', '0.00'), ('E', '0.0'), ('F', '0.00'), ('G', '0.0'), ('H', '0.00')):
            ws['%s%d' % (col, i)].number_format = fmt
            ws['%s%d' % (col, i)].alignment = Alignment(horizontal='right', vertical='center')
        for col in 'DEG':
            ws['%s%d' % (col, i)].fill = FILL_IN
        ws['A%d' % i].alignment = Alignment(horizontal='center')
        ws['C%d' % i].alignment = Alignment(horizontal='center')
    r1 = r[0] - 1
    i = r[0]
    row(['', '合  計', '', '=SUM(D%d:D%d)' % (r0, r1), '', '=SUM(F%d:F%d)' % (r0, r1), '',
         '=ROUND(SUM(H%d:H%d),1)' % (r0, r1), ''], bold=True, size=9.5)
    for col, fmt in (('D', '0.00'), ('F', '0.00'), ('H', '0.0')):
        ws['%s%d' % (col, i)].number_format = fmt
        ws['%s%d' % (col, i)].alignment = Alignment(horizontal='right')
    ROW_VR = i
    merge_line('※床面積は意匠図（平面図 令和8年7月20日版）の壁芯寸法より算定。天井高は意匠未確定のため仮定値（黄色セル）。'
               '確定後に入力し直すと自動で再計算される。', size=8.5, height=24)
    merge_line('※換気対象外：' + '、'.join(P.EXCLUDED) + '（建具で区画された収納。内装はF☆☆☆☆材とする）。',
               size=8.5, height=24)

    # ---- 4. 有効換気量 --------------------------------------------------------------
    head('4. 有効換気量 Ve（機械換気設備：排気機）')
    row(['記号', '設置室', '台数', 'ダクト径', 'ダクト実長\nL (m)', '曲り数\n(個)', '直管相当長\n(m)',
         '有効換気量\n(m³/h･台)', '機種・備考'], bold=True, fill=FILL_H, al='center', wrap=True, size=9)
    ws.row_dimensions[r[0] - 1].height = 30
    f = P.FAN24
    r0 = r[0]
    for u in P.FAN24_UNITS:
        i = r[0]
        row([f['sym'], u['room'], 1, f['duct'], u['duct_len'], u['bends'],
             '=ROUND(E%d+F%d*$C$%d+$C$%d,1)' % (i, i, 0, 0),  # 後で置換
             '=IF(G%d<=%s,%s,"要P-Q確認")' % (i, f['eq_len_ref'], f['q_eff']),
             '%s %s（%s）' % (f['maker'], f['model'], f['power'])], size=9)
        for col in 'EF':
            ws['%s%d' % (col, i)].fill = FILL_IN
        for col in 'ACDEFGH':
            ws['%s%d' % (col, i)].alignment = Alignment(horizontal='center')
    r1 = r[0] - 1
    i = r[0]
    row(['', '合  計', '=SUM(C%d:C%d)' % (r0, r1), '', '', '', '', '=SUM(H%d:H%d)' % (r0, r1), ''], bold=True, size=9.5)
    ws['H%d' % i].number_format = '0.0'; ws['H%d' % i].alignment = Alignment(horizontal='right')
    ROW_VE = i
    # 相当長の係数セル
    i = r[0]
    put('A%d' % i, '直管相当長の算定', size=9, border=False); ws.merge_cells('A%d:B%d' % (i, i))
    put('C%d' % i, P.EQ_ELBOW, size=9, fill=FILL_IN, al='center')
    put('D%d' % i, 'm／90°曲り1個', size=9, border=False); ws.merge_cells('D%d:E%d' % (i, i))
    put('F%d' % i, P.EQ_HOOD, size=9, fill=FILL_IN, al='center')
    put('G%d' % i, 'm／外壁フード（深型・防虫網付）', size=9, border=False); ws.merge_cells('G%d:I%d' % (i, i))
    for k in range(r0, r1 + 1):
        ws['G%d' % k].value = '=ROUND(E%d+F%d*$C$%d+$F$%d,1)' % (k, k, i, i)
    r[0] += 1
    merge_line('※有効換気量はメーカー公表値（%s：静圧時 直管相当長20m で %.1f m³/h、30m で 76 m³/h、開放時 %.0f m³/h、'
               '消費電力 %.1fW）。直管相当長が20m以下であることを確認し、20m時の値を採用（安全側）。'
               % (f['model'], f['q_eff'], f['q0'], f['watt']), size=8.5, height=36)

    # ---- 5. 判定 ------------------------------------------------------------------
    head('5. 判  定')
    i = r[0]
    put('A%d' % i, '有効換気量 Ve', fill=FILL_H); ws.merge_cells('A%d:C%d' % (i, i))
    put('D%d' % i, '=H%d' % ROW_VE, fmt='0.0" m³/h"', al='right'); ws.merge_cells('D%d:E%d' % (i, i))
    put('F%d' % i, '≧', al='center')
    put('G%d' % i, '必要有効換気量 Vr', fill=FILL_H); ws.merge_cells('G%d:H%d' % (i, i))
    put('I%d' % i, '=H%d' % ROW_VR, fmt='0.0" m³/h"', al='right')
    r[0] += 1
    i = r[0]
    put('A%d' % i, '判定', fill=FILL_H, bold=True); ws.merge_cells('A%d:C%d' % (i, i))
    put('D%d' % i, '=IF(D%d>=I%d,"適合（OK）","不適合（NG）")' % (i - 1, i - 1), bold=True, fill=FILL_OK, al='center')
    ws.merge_cells('D%d:F%d' % (i, i))
    put('G%d' % i, '余裕率 Ve/Vr', fill=FILL_H); ws.merge_cells('G%d:H%d' % (i, i))
    put('I%d' % i, '=ROUND(D%d/I%d,2)' % (i - 1, i - 1), fmt='0.00" 倍"', al='right')
    r[0] += 1
    # 本堂部の天井高 上限（参考）
    hondo = [k for k, rm in enumerate(P.ROOMS) if rm['no'] in ('1', '3', '4')]
    a_cells = '+'.join('D%d' % (ROW_VR - len(P.ROOMS) + k) for k in hondo)
    vr_cells = '+'.join('H%d' % (ROW_VR - len(P.ROOMS) + k) for k in hondo)
    i = r[0]
    put('A%d' % i, '参考：外陣・内陣・護摩堂の天井高が一律 H(m) のとき適合する上限', size=9, border=False)
    ws.merge_cells('A%d:G%d' % (i, i))
    put('H%d' % i, '=ROUNDDOWN((D%d-(I%d-(%s)))/(0.3*(%s)),2)' % (i - 2, i - 2, vr_cells, a_cells), fmt='0.00" m"',
        al='right', size=9)
    r[0] += 1

    # ---- 6. 給気口 -------------------------------------------------------------------
    head('6. 給気口（外気導入）')
    row(['記号', '設置室', '台数', 'ダクト径', '型式', None, None, None, '仕様'], bold=True, fill=FILL_H, al='center', size=9)
    ws.merge_cells('E%d:H%d' % (r[0] - 1, r[0] - 1))
    for s in P.SUPPLY:
        rooms = {}
        for u in s['units']:
            rooms[u[0]] = rooms.get(u[0], 0) + 1
        for nm, n in rooms.items():
            i = r[0]
            row([s['sym'], nm, n, s['duct'], 'パナソニック %s' % s['model'], None, None, None,
                 '%s（%s）' % (s['spec'], s['color'])], size=9)
            ws.merge_cells('E%d:H%d' % (i, i))
            for col in 'ACD':
                ws['%s%d' % (col, i)].alignment = Alignment(horizontal='center')
    merge_line('※給気口は居室の天井付近に設け、外壁貫通部に深型フード（防虫網付）を設けて外気流・雨水の逆流を防止する。',
               size=8.5, height=24)

    # ---- 7. 換気経路・その他の措置 -----------------------------------------------------
    from openpyxl.worksheet.pagebreak import Break
    ws.row_breaks.append(Break(id=r[0]))
    head('7. 換気経路・その他の措置')
    for s in ['(1) 換気経路：給気口（各居室）→ 居室建具（アンダーカット）→ 廊下・ホール → 便所建具 → 排気機（EF-1）→ 屋外。',
              '(2) 建具：居室・廊下の建具はアンダーカット10mm以上（有効開口面積100cm²以上）、便所の建具はアンダーカット20mm以上又はガラリ。',
              '    障子・襖・格子戸等で常時通気が確保されるものはこれによる。',
              '(3) 排気機（EF-1）は24時間常時運転とし、スイッチに「24時間換気（常時運転）」の表示を行う。',
              '(4) 内装仕上げの制限（令20条の7）：居室の内装仕上げには第1種・第2種ホルムアルデヒド発散建築材料を使用しない'
              '（F☆☆☆☆ 又は規制対象外の材料とする）。',
              '(5) 天井裏等の措置（令20条の9・平15国交告第274号）：天井裏・小屋裏・床下・壁内・収納の下地材等に'
              'F☆☆☆以上の建築材料を用いる（建材による措置）。※意匠仕様と整合のこと。',
              '(6) 換気設備の構造（令129条の2の5）：給気口・排気口は外気流によって換気能力が低下しない構造（深型フード）とし、'
              '防虫網を設ける。']:
        merge_line(s, size=9, height=27)

    head('8. 前提条件・要確認事項')
    for s in ['・天井高（特に外陣・内陣・護摩堂の格天井・折上天井）は仮定値。矩計図・断面図確定後に本書を更新すること。',
              '・機器配置・ダクトルートは設計段階の計画。施工図段階でダクト直管相当長が20mを超える場合はP-Q曲線で有効換気量を確認する。',
              '・換気機器はパナソニック「換気システムご提案書」（当初プラン）を基に、現行平面図（令和8年7月20日版）に合わせて配置を修正した。',
              '・電源（EF-1 単相100V 常時通電回路）は電気設備工事による。']:
        merge_line(s, size=9, height=27)

    # ---- 印刷設定 -------------------------------------------------------------------
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = 'portrait'
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=0.5, right=0.4, top=0.6, bottom=0.6)
    ws.print_options.horizontalCentered = True
    ws.oddFooter.center.text = '&P / &N'
    ws.oddHeader.right.text = '安楽寺 本堂 新築工事　換気設備'
    wb.calculation.fullCalcOnLoad = True
    wb.save(out)


if __name__ == '__main__':
    build(sys.argv[1])
